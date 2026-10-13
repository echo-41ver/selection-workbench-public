from pathlib import Path
import sys

from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from selection_workflow.excel_verdict_style import COLORS, verdict_kind
from export_category_workbooks import export_stage_book


@pytest.mark.parametrize('state,kind', [
    ('不同款（Codex看图复核）', 'different'), ('不同款（文字硬冲突）', 'different'),
    ('确认同款（Codex看图复核）', 'same'), ('同款（存在规格变体）', 'same'),
    ('待复核', 'review'), ('待补充证据（Codex看图复核）', 'review'),
    ('保留复核（未纳入第三层）', 'review'), ('第三层待运行', 'unanalysed'),
    ('第二层未入围', 'unanalysed'), ('', 'unanalysed')])
def test_final_verdict_kind(state, kind):
    assert verdict_kind(state) == kind


def test_whole_row_uses_final_verdict_color_and_keeps_number_and_links(tmp_path):
    headers = ['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
        'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','第三层总分','最终判断','依据']
    outcomes = ['不同款（Codex看图复核）','待补充证据（Codex看图复核）','确认同款','第三层待运行']
    rows = [[f'A{i}','Amazon','https://amazon.com/a','https://amazon.com/a.jpg',f'T{i}',
        'TEMU','https://temu.com/a','https://temu.com/a.jpg','事实',77.25,'第二层：进入第三层',
        None,outcome,'依据'] for i,outcome in enumerate(outcomes)]
    destination = tmp_path / '第一二三层.xlsx'
    export_stage_book(destination, headers, rows, {'商品对总数':4}, '测试品类',
                      third_layer=True, low_asins={'A0','A1','A2','A3'})
    book = load_workbook(destination)
    sheet = book['全量综合结果']
    for row, kind in zip(range(5,9), ('different','review','same','unanalysed')):
        expected = COLORS[kind].lstrip('#')
        assert all(sheet.cell(row,col).fill.fgColor.rgb[-6:] == expected for col in range(1,16))
    assert sheet['I5'].number_format == '0.00' and sheet['I5'].value == 77.25
    assert sheet['K5'].value.startswith('=HYPERLINK(')
    assert sheet['N5'].value == outcomes[0]
    shortlist = book['同款不超过20_同款及待复核']
    assert [shortlist.cell(row,3).value for row in range(5,shortlist.max_row+1)] == ['A1','A2']
    assert shortlist['N4'].value == '最终判断'
    assert shortlist['N5'].fill.fgColor.rgb[-6:] == COLORS['review'][1:]
    book.close()
