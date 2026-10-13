"""Opt-in real snapshot plumbing acceptance; vectors and vision are explicit mocks."""
import copy,hashlib,json,os,sys
from pathlib import Path
from types import SimpleNamespace
import numpy as np
import pytest
from PIL import Image
from identity_fixtures import identity
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))


def test_real_snapshot_increment_then_manual_mock_visual(tmp_path,monkeypatch):
    source=os.getenv('OFFLINE_REAL_SNAPSHOT_RUN')
    if not source:pytest.skip('Requires an explicitly selected existing real snapshot')
    origin=Path(source);amazon_source=origin/'amazon/amazon_frontend_latest.jsonl';temu_source=origin/'temu/union/products.jsonl'
    original_hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in (amazon_source,temu_source)}
    amazon=[json.loads(s) for s in amazon_source.read_text('utf-8').splitlines() if s.strip()]
    temu=[json.loads(s) for s in temu_source.read_text('utf-8').splitlines() if s.strip()]
    assert len(amazon)==4 and len(temu)==649
    cat=copy.deepcopy(json.loads((origin/'category.json').read_text('utf-8')))
    cat['execution_mode']='full';cat['category_key']='offline_real_snapshot';cat['category_name_zh']='离线验收遥控器模拟结果';cat['stage3']['max_concurrency']=1
    base=tmp_path/'runs'/cat['category_key'];base.mkdir(parents=True)
    category=base/'category.json';category.write_text(json.dumps(cat),encoding='utf-8')
    image=base/'MOCK_IMAGE.png';Image.new('RGB',(16,16),'red').save(image,compress_level=0)
    # Preserve genuine titles/IDs/prices/sales and provenance. The local image is
    # a MOCK solely to test plumbing; no image accuracy is claimed.
    for row in amazon+temu:row['local_image_path']=str(image)
    a=base/'amazon.jsonl';t=base/'temu.jsonl'
    def write(path,rows):path.write_text(''.join(json.dumps(r)+'\n' for r in rows),encoding='utf-8')
    write(a,amazon);write(t,temu[:640])
    from selection_workflow.incremental_stage12 import IncrementalStage12,SourceSpec,PathSpec
    import run_category_layer3 as consumer
    import export_category_workbooks as exporter
    import openai
    monkeypatch.setattr(consumer,'ROOT',tmp_path);monkeypatch.setattr(exporter,'ROOT',tmp_path)
    class V:
        revision='OFFLINE_MOCK_VECTOR'
        def encode(self,path):return np.array([1.,0.])
        def encode_many(self,paths):return [self.encode(p) for p in paths]
    pipeline=IncrementalStage12(base/'stage12',vectorizer=V(),text_gate_name='declarative:'+str(category),rule_version=cat['stage3']['rule_version'])
    calls=[]
    class FakeClient:
        def __init__(self,**kw):self.responses=self
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def create(self,**kw):
            ids=[part['text'] for part in kw['input'][0]['content'] if part['type']=='input_text' and part['text'].startswith('C')]
            calls.append(len(ids))
            raw={'candidates':[{'i':i,'v':identity(),'c':.9,'m':{str(j):{'r':'S','a':'模拟事实：右上角三列圆形按钮','t':'模拟事实：右上角三列圆形按钮','h':'NONE'} for j in range(len(cat['stage3']['modules']))}} for i in ids]}
            return SimpleNamespace(output_text=json.dumps(raw),usage=SimpleNamespace(model_dump=lambda:{'total_tokens':0}),id='OFFLINE_MOCK')
    monkeypatch.setattr(openai,'OpenAI',FakeClient);monkeypatch.setenv('DEEPSEEK_API_KEY','TEST_ONLY_OFFLINE_NOT_A_REAL_KEY')
    frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True);candidate=frozen/f'layer3_candidates_{cat["category_key"]}.jsonl';result=base/'stage123/results.jsonl'
    def consume():
        monkeypatch.setattr(sys,'argv',['consumer','--category',str(category),'--candidates',str(candidate),'--results',str(result)])
        consumer.main()
    try:
        pipeline.configure([SourceSpec('a','AMAZON','amazon_generic',a,'replace'),SourceSpec('t','TEMU','temu_frontend',t,'replace')],[PathSpec(cat['category_key'],'a','t','OFFLINE_REAL_SNAPSHOT')])
        pipeline.ingest_source('a');pipeline.ingest_source('t');pipeline.process_batch(4096)
        pipeline.materialize_layer3_snapshot(cat['category_key'],frozen)
        before=list(consumer.rows(candidate));assert before
        with pytest.raises(SystemExit,match='API请求数为0'):consume()
        assert calls==[] and not result.exists()
        (base/'stage123/paid_enabled.json').write_text('{"scope":"OFFLINE_MOCK_ONLY"}')
        consume();first_requests=len(calls);assert first_requests>0
        consume();assert len(calls)==first_requests  # successful candidates are not charged again
        write(t,temu);pipeline.ingest_source('t');pipeline.process_batch(4096)
        pipeline.materialize_layer3_snapshot(cat['category_key'],frozen)
        after=list(consumer.rows(candidate));assert len(after)>=len(before)
        consume();assert sum(calls)==len(after)
        groups=pipeline.status()['groups'];assert groups[0]['actual_pairs']==4*649
        exporter.export(category,'layer123')
        from openpyxl import load_workbook
        workbook=base/'excel'/f'{cat["category_name_zh"]}_第一二三层.xlsx'
        book=load_workbook(workbook,read_only=True);sheet=book['全量综合结果'];assert sheet.max_row-4==2596;book.close()
        report={'source_hashes':original_hashes,'original_sources_unchanged':True,'amazon':4,'temu_initial':640,'temu_final':649,'pairs_initial':2560,'pairs_final':2596,'candidates_initial':len(before),'candidates_final':len(after),'mock_requests_before_manual_authorization':0,'mock_requests_after_authorization':first_requests,'mock_requests_total':len(calls),'mock_candidate_consumptions':sum(calls),'successful_candidates_reused':True,'excel_rows':2596,'real_api_requests':0,'vector_and_visual_accuracy_tested':False,'formal_history_modified':False}
        for path,expected in original_hashes.items():assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==expected
        destination=os.getenv('OFFLINE_ACCEPTANCE_REPORT')
        if destination:Path(destination).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    finally:pipeline.close()
