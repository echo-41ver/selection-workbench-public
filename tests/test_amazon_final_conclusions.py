from openpyxl import load_workbook
from export_category_workbooks import amazon_conclusions, export_stage_book


def test_same_only_presentation_preserves_legacy_verdicts_and_embeds_images(tmp_path):
    from PIL import Image
    image = tmp_path/'fixture.jpg'
    Image.new('RGB', (160, 120), 'white').save(image)
    headers = ['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
               'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    base = ['A','商品','https://amazon.com/a','https://amazon.com/a.jpg','T','商品',
            'https://temu.com/t','https://temu.com/same.jpg','事实',80,'候选']
    rows = [base+['确认同款'], [*base[:4], 'T2', *base[5:7], 'https://temu.com/review.jpg', *base[8:], '待补充证据'],
            ['B',*base[1:], '确认同款'], ['B',*base[1:], '确认同款']]
    conclusions = amazon_conclusions(headers, rows, 1, same_only=True)
    assert conclusions[0]['same'] == 1 and conclusions[0]['review'] == 1
    assert conclusions[0]['conclusion'].startswith('符合：当前采样同款')
    assert conclusions[0]['review_images'] == {base[3], base[7]}
    output = tmp_path/'accepted_legacy.xlsx'
    export_stage_book(output, headers, rows, {'商品对总数':4,'采集前设置的同款上限':1}, '测试',
        third_layer=True, low_asins={'A'}, conclusions=conclusions, same_only=True,
        thumbnails={base[3]:image, base[7]:image})
    book = load_workbook(output)
    assert book['全量综合结果']['M6'].value == '待补充证据'
    assert book['Amazon最终结论']['E5'].value == 1
    final = book['Amazon最终结论']
    assert final.max_row == 5 and final['A5'].value == 'A'
    assert final['M4'].value == 'Amazon商品链接'
    assert final['M5'].value == '=HYPERLINK("https://amazon.com/a","打开商品")'
    sheet = book['同款不超过1_同款明细']
    assert sheet.max_row == 5 and sheet['M5'].value == '确认同款'
    assert len(sheet._images) == 2
    assert [(picture.anchor._from.col, picture.anchor._from.row) for picture in sheet._images] == [(1,4),(4,4)]
    assert sheet['K5'].value == '=HYPERLINK("https://amazon.com/a","打开商品")'
    assert sheet['L5'].value == '=HYPERLINK("https://temu.com/t","打开商品")'
    book.close()


def test_shortlist_keeps_long_product_url_as_clickable_hyperlink(tmp_path):
    from PIL import Image
    image = tmp_path/'image.png'
    Image.new('RGB',(100,100),'white').save(image)
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
             'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    long_url = 'https://www.temu.com/' + 'portable-pump-'*24 + 'g-607153737801747.html'
    row=['A','商品','https://amazon.com/dp/A','https://amazon.com/a.jpg','T','商品',
         long_url,'https://temu.com/t.jpg','事实',90,'候选','确认同款']
    output=tmp_path/'long_url.xlsx'
    export_stage_book(output,headers,[row],{'商品对总数':1,'采集前设置的同款上限':5},'测试',
        third_layer=True,low_asins={'A'},same_only=True,thumbnails={row[3]:image,row[7]:image})
    book=load_workbook(output)
    assert book['同款不超过5_同款明细']['L5'].hyperlink.target==long_url
    assert book['同款不超过5_同款明细']['L5'].value=='打开商品'
    assert book['全量综合结果']['L5'].value==long_url
    book.close()


def test_binary_shortlist_contains_only_same_and_sampling_conclusion(tmp_path):
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
             'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    base=['A','商品','','','T','商品','','','事实',80,'候选']
    rows=[base+['确认同款'],[ *base[:4], 'T2', *base[5:], '不同款'],
          [*base[:4], 'T3', *base[5:], '第二层未入围第三层']]
    conclusions=amazon_conclusions(headers,rows,5,binary=True)
    assert conclusions[0]['same']==1 and conclusions[0]['review']==0
    assert conclusions[0]['conclusion'].startswith('符合：当前采样')
    path=tmp_path/'binary.xlsx'
    export_stage_book(path,headers,rows,{'商品对总数':3,'采集前设置的同款上限':5,
        '视觉结论口径':'visible_core_binary_v1'},'测试',third_layer=True,
        low_asins={'A'},conclusions=conclusions)
    book=load_workbook(path,read_only=True,data_only=True)
    sheet=book['同款不超过5_同款明细']
    assert sheet.max_row==5 and sheet['F5'].value=='T'
    assert book['Amazon最终结论']['K5'].value.startswith('符合：当前采样')
    book.close()


def test_failed_legacy_facts_are_unfinished_analysis_not_a_completed_product_conclusion(tmp_path):
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
             'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    base=['A','商品','','','T','商品','','','事实',80,'候选']
    rows=[base+['确认同款'], base+['分析失败（历史事实不完整）'],
          base+['第三层待运行'], base+['第二层未入围第三层']]
    conclusions=amazon_conclusions(headers,rows,5,binary=True)
    assert conclusions[0]['same']==1 and conclusions[0]['review']==0
    assert conclusions[0]['pending']==2 and conclusions[0]['unanalysed']==3
    assert conclusions[0]['conclusion'].startswith('分析未完成')
    path=tmp_path/'unfinished.xlsx'
    export_stage_book(path,headers,rows,{'商品对总数':4,'采集前设置的同款上限':5,
        '视觉结论口径':'visible_core_binary_v1'},'测试',third_layer=True,
        low_asins={'A'},conclusions=conclusions)
    book=load_workbook(path,read_only=True,data_only=True)
    assert book['Amazon最终结论']['E5'].value==2
    assert book['Amazon最终结论']['K5'].value.startswith('分析未完成')
    assert book['同款不超过5_同款明细'].max_row==5
    book.close()


def test_threshold_counts_same_and_preserves_review_without_different_rows(tmp_path):
    headers = ['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
               'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    def pair(asin, goods, state):
        return [asin,'商品','https://amazon.com/a','https://amazon.com/a.jpg',goods,'商品',
                'https://temu.com/a','https://temu.com/a.jpg','事实',80,'第二层状态',state]
    rows = [pair('A', '1','确认同款'), pair('A','2','同款（存在规格变体）'),
            pair('A','3','待补充证据（Codex看图复核）'),pair('A','4','不同款'),
            pair('A','5','第二层未入围'),pair('A','6','第三层待运行'),
            pair('B','1','确认同款'),pair('B','2','确认同款'),pair('B','3','确认同款'),
            pair('B','4','待复核'),pair('C','1','待复核')]
    conclusions = amazon_conclusions(headers, rows, 2)
    counts = {item['asin']:item for item in conclusions}
    assert counts['A']['selected'] and counts['A']['same'] == 2
    assert counts['A']['review'] == 1 and counts['A']['unanalysed'] == 2
    assert counts['A']['conclusion'].startswith('待复核')
    assert not counts['B']['selected'] and counts['B']['conclusion'].startswith('不符合')
    assert counts['C']['selected'] and counts['C']['same'] == 0
    path=tmp_path/'final.xlsx'
    export_stage_book(path,headers,rows,{'商品对总数':len(rows),'采集前设置的同款上限':2},'测试',
        third_layer=True, low_asins={item['asin'] for item in conclusions if item['selected']}, conclusions=conclusions)
    book=load_workbook(path,read_only=True,data_only=True)
    assert book.sheetnames[0] == 'Amazon最终结论'
    assert book['Amazon最终结论']['I5'].value == 2
    sheet=book['同款不超过2_同款及待复核']
    selected=list(sheet.iter_rows(min_row=5,values_only=True))
    assert {(row[2],row[5]) for row in selected} == {('A','1'),('A','2'),('A','3'),('C','1')}
    assert all('不同款' not in row[12] and '未入围' not in row[12] and '待运行' not in row[12] for row in selected)
    assert book['全量综合结果'].max_row == len(rows)+4
    book.close()


def test_new_summary_and_shortlist_embed_both_images_but_full_table_keeps_links(tmp_path):
    import zipfile
    import xml.etree.ElementTree as ET
    from PIL import Image
    image=tmp_path/'fixture.jpg'; Image.new('RGB',(160,120),'white').save(image)
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
             'TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态','最终判断']
    row=['A','商品','https://amazon.com/a','https://amazon.com/a.jpg','T','商品',
         'https://temu.com/t','https://temu.com/t.jpg','事实',90,'候选','待复核']
    conclusions=amazon_conclusions(headers,[row],5)
    output=tmp_path/'images.xlsx'
    export_stage_book(output,headers,[row],{'商品对总数':1,'采集前设置的同款上限':5},'测试',
        third_layer=True,low_asins={'A'},conclusions=conclusions,
        thumbnails={row[3]:image,row[7]:image})
    book=load_workbook(output,read_only=True,data_only=False)
    assert book['Amazon最终结论']['C4'].value=='Amazon主图'
    assert book['全量综合结果']['B5'].value.startswith('=HYPERLINK(')
    assert book['同款不超过5_同款及待复核']['B5'].value is None
    book.close()
    with zipfile.ZipFile(output) as package:
        drawings=[name for name in package.namelist() if name.startswith('xl/drawings/drawing') and name.endswith('.xml')]
        counts=sorted(len(ET.fromstring(package.read(name))) for name in drawings)
        assert counts==[1,2]
