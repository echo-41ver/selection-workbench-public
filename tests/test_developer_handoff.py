import hashlib
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from selection_automation.developer_handoff import build_bundle, describe_run

class HandoffTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.base=self.root/'runs'/'remote_test'
        (self.base/'amazon').mkdir(parents=True);(self.base/'temu'/'union').mkdir(parents=True)
    def tearDown(self): self.tmp.cleanup()
    def write(self,name,value): (self.base/name).write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')
    def bundle(self): return zipfile.ZipFile(io.BytesIO(build_bundle(self.root,'remote_test')))
    def test_traversal(self):
        for key in ['../remote_test','remote_test/../other','C:\\secret','remote_test%2f..','..','/tmp','REMOTE']:
            with self.assertRaises(ValueError): describe_run(self.root,key)
        with self.assertRaises(FileNotFoundError): describe_run(self.root,'absent')
    def test_empty(self):
        d=describe_run(self.root,'remote_test');self.assertEqual(d['counts'],{'amazon':0,'temu':0});self.assertTrue(d['snapshot']);self.assertFalse(d['platform_complete'])
        self.assertIsNone(d['collection']['is_collecting'])
        with self.bundle() as z:
            self.assertEqual(z.read('amazon.jsonl'),b'')
            for name in ('amazon.xlsx','temu.xlsx'):
                book=load_workbook(io.BytesIO(z.read(name)),read_only=True)
                self.assertEqual(book.sheetnames,['Products','Images','Specifications','Observations','Readme'])
                self.assertEqual(book['Products'].max_row,1)
    def test_fields_allowlist_hash_and_secrets(self):
        secret='NEVER_EXPORT_THIS_SECRET'
        self.write('category.json',{'category_name_zh':'遥控器','api_key':secret})
        self.write('status.json',{'状态':'采集中','proxy_user':secret,'日志':'C:\\private\\file'})
        self.write('amazon/amazon_frontend_latest.json',{'source':'VISIBLE_AMAZON_US_FRONTEND','complete':True,'records':[{'asin':'B000TEST','rank':2,'title':'Remote','product_url':'https://www.amazon.com/dp/B000TEST?api_key='+secret,'main_image_url':'https://m.media-amazon.com/images/I/example.jpg','specifications':[{'label':'Brand','value':'LG'}],'cookies':secret,'account':secret,'local_path':'C:\\private\\file'}]})
        row={'goods_id':'123','title':'Remote','source':'VISIBLE_TEMU_FRONTEND_KEYWORD_UNION','search_query':'remote','source_queries':['remote','tv remote'],'search_sort':{'requested':'sales','state':'UI_CONFIRMED','label':'Top sales','token':secret},'observation_sources':[{'search_query':'tv remote','snapshot':'C:\\private\\file','cookies':secret}],'product_url':'https://www.temu.com/goods.html?goods_id=123&token='+secret,'price_usd':2.34,'sales_text':'25K+ sold','main_image_url':'https://img.kwcdn.com/a.jpg','api_key':secret}
        (self.base/'temu/union/products.jsonl').write_text(json.dumps(row)+'\n'+json.dumps(row)+'\n',encoding='utf-8')
        with self.bundle() as z:
            raw=b'\n'.join(z.read(n) for n in z.namelist());self.assertNotIn(secret.encode(),raw);self.assertNotIn(b'C:\\\\private',raw)
            a=json.loads(z.read('amazon.jsonl'));t=json.loads(z.read('temu.jsonl'));m=json.loads(z.read('manifest.json'))
            self.assertEqual(a['product_id'],'B000TEST');self.assertEqual(a['sampling']['rank'],2);self.assertEqual(a['specifications'][0]['value'],'LG')
            self.assertEqual(t['sampling']['sort']['state'],'UI_CONFIRMED');self.assertEqual(t['sampling']['observations'][0]['keyword'],'tv remote');self.assertEqual(t['price_usd'],2.34)
            self.assertEqual(t['url'],'https://www.temu.com/goods.html?goods_id=123')
            self.assertEqual(t['images'],['https://img.kwcdn.com/a.jpg'])
            amazon=load_workbook(io.BytesIO(z.read('amazon.xlsx')),read_only=True)
            temu=load_workbook(io.BytesIO(z.read('temu.xlsx')),read_only=True)
            self.assertEqual(amazon['Products']['A2'].value,'B000TEST')
            self.assertEqual(amazon['Specifications']['C2'].value,'LG')
            self.assertEqual(temu['Products']['A2'].value,'123')
            self.assertEqual(temu['Images']['C2'].value,'https://img.kwcdn.com/a.jpg')
            self.assertEqual(temu['Observations']['B2'].value,'tv remote')
            self.assertEqual(m['counts'],{'amazon':1,'temu':1});self.assertTrue(m['collection']['is_collecting']);self.assertEqual(m['category_name'],'遥控器')
            for f in m['files']: self.assertEqual(hashlib.sha256(z.read(f['name'])).hexdigest(),f['sha256'])
    def test_corrupt_snapshot_is_not_silently_empty(self):
        (self.base/'amazon/amazon_frontend_latest.json').write_text('{"records": [',encoding='utf-8')
        with self.assertRaises(ValueError): build_bundle(self.root,'remote_test')
    def test_jsonl_fallback_and_missing_values(self):
        (self.base/'amazon/amazon_frontend_latest.jsonl').write_text('{"asin":"A"}\n',encoding='utf-8')
        with self.bundle() as z:
            a=json.loads(z.read('amazon.jsonl'));self.assertIsNone(a['title']);self.assertEqual(a['images'],[]);self.assertIsNone(a['source'])
    def test_absolute_path_and_credentials_in_business_fields(self):
        self.write('amazon/amazon_frontend_latest.json',{'records':[{'asin':'A','title':'API_KEY=secret','product_url':'https://user:secret@www.amazon.com/dp/A','specifications':[{'label':'debug','value':'/home/user/private'}]}]})
        with self.bundle() as z:
            a=json.loads(z.read('amazon.jsonl'));self.assertIsNone(a['title']);self.assertIsNone(a['url']);self.assertIsNone(a['specifications'][0]['value'])

if __name__=='__main__': unittest.main()
