import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys

import pytest
from selection_automation.rules import text_gate, score, generic_keywords, validate
from selection_automation.vision import protocol, decode, vision_revision
from selection_automation.process_lock import ProcessLock
from identity_fixtures import identity

ROOT=Path(__file__).resolve().parents[1]
CAT=json.loads((ROOT/'config/categories/power_drill.json').read_text(encoding='utf-8'))


def test_category_isolation_and_unknown_text():
    assert text_gate(CAT,'cordless drill driver','corded electric drill driver')['stage1_status']=='TEXT_HARD_EXCLUDE'
    assert text_gate(CAT,'power drill','corded drill')['stage1_status']!='TEXT_HARD_EXCLUDE'
    assert not text_gate(CAT,'not corded cordless drill','corded drill')['excluded']
    assert not text_gate(CAT,'cordless and corded drill combo','corded drill')['excluded']
    plan=generic_keywords([{'asin':'A','title':'Brushless cordless 20V drill'}],CAT)
    assert any(p['query']=='brushless power drill' for p in plan)
    assert not any('towel' in p['query'] for p in plan)


def test_vision_bilateral_evidence_and_revision(tmp_path):
    raw={'candidates':[{'i':'C00','c':.9,'v':identity(),'m':{str(i):{'r':'S','a':'左侧圆弧主体','t':'左侧圆弧主体','h':'NONE'} for i in range(4)}}]}
    facts=decode(CAT,raw,['C00'])['C00']
    assert score(CAT,facts)['state']=='确认同款'
    facts['modules']['body']['amazon_evidence']=''
    assert score(CAT,facts)['state']=='待补充证据'
    raw['candidates'].append(raw['candidates'][0])
    with pytest.raises(ValueError): decode(CAT,raw,['C00'])
    image=tmp_path/'image'; image.write_bytes(b'one')
    pair={'asin':'A','goods_id':'T','pair_revision':'p','amazon_image_path':str(image),'temu_image_path':str(image)}
    before=vision_revision(CAT,pair); image.write_bytes(b'two')
    assert before!=vision_revision(CAT,pair)
    modified=copy.deepcopy(CAT); modified['stage3']['model']='other'
    assert vision_revision(CAT,pair)!=vision_revision(modified,pair)


def test_process_lock_excludes_second_owner(tmp_path):
    with ProcessLock(tmp_path/'lock'):
        with pytest.raises(RuntimeError):
            with ProcessLock(tmp_path/'lock'): pass
    with ProcessLock(tmp_path/'lock'): pass


def test_successful_service_round_resets_failure_budget(tmp_path):
    from types import SimpleNamespace
    sys.path.insert(0,str(ROOT/'scripts'))
    from selection import service_step
    children={};finished=set();attempts={};next_at={};state={};launched=[]
    for code in [1,0,1,0,1]:
        children['stage3']=SimpleNamespace(poll=lambda:code,returncode=code)
        service_step('stage3',[],children=children,finished=finished,attempts=attempts,next_at=next_at,state=state,launch=lambda n,c:launched.append(n),log=tmp_path,repeat=True,interval=0)
    assert 'stage3' not in finished
    assert attempts['stage3']==1
    assert len(launched)==5


def test_mcp_stdio_has_only_protocol_stdout():
    data='\n'.join(json.dumps(m) for m in [{'jsonrpc':'2.0','id':1,'method':'initialize','params':{'protocolVersion':'2024-11-05'}},{'jsonrpc':'2.0','method':'notifications/initialized'},{'jsonrpc':'2.0','id':2,'method':'tools/list'},{'jsonrpc':'2.0','id':3,'method':'tools/call','params':{'name':'selection_status','arguments':{'category_key':'offline_nonexistent_fixture'}}}])+'\n'
    result=subprocess.run([sys.executable,str(ROOT/'scripts/mcp_server.py')],input=data,text=True,capture_output=True,encoding='utf-8')
    assert result.returncode==0,result.stderr
    lines=[json.loads(line) for line in result.stdout.splitlines()]
    assert len(lines)==3
    assert {tool['name'] for tool in lines[1]['result']['tools']} == {
        'selection_doctor', 'selection_prepare', 'selection_set_api_key', 'selection_start',
        'selection_status', 'selection_stop', 'selection_rotate', 'selection_resume_api', 'selection_category_contract'}
    assert not lines[2]['result'].get('isError')
    content=json.loads(lines[2]['result']['content'][0]['text'])
    assert json.loads(content['结果'])['状态']['状态']=='尚未启动'


@pytest.mark.skipif(sys.platform!='win32',reason='Windows installer bootstrap')
def test_installer_parses_in_windows_powershell5():
    import os
    shell=Path(os.environ['SystemRoot'])/'System32/WindowsPowerShell/v1.0/powershell.exe'
    command="$tok=$null; $err=$null; [System.Management.Automation.Language.Parser]::ParseFile((Join-Path (Get-Location) 'install.ps1'),[ref]$tok,[ref]$err) | Out-Null; if($err){$err; exit 1}"
    result=subprocess.run([str(shell),'-NoProfile','-Command',command],cwd=ROOT,capture_output=True)
    assert result.returncode==0,result.stdout


def test_reject_path_traversal():
    cat=copy.deepcopy(CAT); cat['category_key']='../escape'
    with pytest.raises(ValueError): validate(cat)


@pytest.mark.skipif(sys.platform!='win32',reason='Windows DPAPI')
def test_dpapi_roundtrip(tmp_path):
    from selection_automation.secrets import store,load
    path=tmp_path/'key.dpapi'; store(path,{'key':'test-secret-only'})
    assert b'test-secret-only' not in path.read_bytes()
    assert load(path)=={'key':'test-secret-only'}


def test_generic_excel_reopen_and_formula_injection(tmp_path):
    sys.path.insert(0,str(ROOT/'scripts'))
    from export_category_workbooks import write_book
    from openpyxl import load_workbook
    dest=tmp_path/'result.xlsx'
    write_book(dest,['ASIN','标题','主图链接'],[['A','=1+1','https://example.com/image.jpg']],{'商品对总数':1})
    book=load_workbook(dest)
    assert book['明细_1']['B2'].data_type=='s'
    assert book['明细_1']['C2'].hyperlink.target=='https://example.com/image.jpg'
    book.close()


def test_offline_end_to_end_incremental_export_and_zero_paid_calls(tmp_path,monkeypatch):
    """Synthetic fixtures exercise the wiring, not vision accuracy or real websites."""
    import numpy as np
    from PIL import Image
    from selection_workflow.incremental_stage12 import IncrementalStage12,SourceSpec,PathSpec
    sys.path.insert(0,str(ROOT/'scripts'))
    import export_category_workbooks as exporter
    import selection
    monkeypatch.setattr(exporter,'ROOT',tmp_path)
    monkeypatch.setattr(selection,'ROOT',tmp_path)
    cat=copy.deepcopy(CAT); key=cat['category_key']; base=tmp_path/'runs'/key
    base.mkdir(parents=True)
    category=base/'category.json'; category.write_text(json.dumps(cat),encoding='utf-8')
    image=base/'fixture.png'; Image.new('RGB',(16,16),'red').save(image,compress_level=0)
    amazon=base/'amazon.jsonl'; temu=base/'temu.jsonl'
    amazon.write_text(json.dumps({'asin':'A','title':'cordless drill driver','local_image_path':str(image)})+'\n',encoding='utf-8')
    temu.write_text('\n'.join(json.dumps({'goods_id':i,'title':title,'local_image_path':str(image)}) for i,title in [('T1','cordless drill driver'),('T2','corded electric drill driver')])+'\n',encoding='utf-8')
    class V:
        revision='fixture-only-vector'
        def encode(self,path): return np.array([1.,0.],dtype=np.float32)
        def encode_many(self,paths): return [self.encode(p) for p in paths]
    output=base/'stage12'
    pipeline=IncrementalStage12(output,vectorizer=V(),text_gate_name='declarative:'+str(category),rule_version=cat['stage3']['rule_version'])
    try:
        pipeline.configure([SourceSpec('a','AMAZON','amazon_generic',amazon,'replace'),SourceSpec('t','TEMU','temu_frontend',temu,'replace')],[PathSpec(key,'a','t','fixture')])
        pipeline.ingest_source('a');pipeline.ingest_source('t')
        pipeline.process_batch(10)
        frozen=base/'stage123'/'export_frozen'; frozen.mkdir(parents=True)
        pipeline.materialize_layer3_snapshot(key,frozen)
        candidates=list(exporter.jsonl(frozen/f'layer3_candidates_{key}.jsonl'))
        assert len(candidates)==1
        assert candidates[0]['goods_id']=='T1'
        config=base/'stage12_config.json';config.write_text(json.dumps({'output_dir':str(output)}),encoding='utf-8')
        assert selection.completion(config,False)['商品对']==2
        assert selection.completion(config,True) is None
        row=candidates[0]
        facts={'identity':identity(),'modules':{k:{'relation':'SAME','amazon_evidence':'左侧圆弧主体','temu_evidence':'左侧圆弧主体','hard_conflict_code':'NONE'} for k in cat['stage3']['modules']},'overall_confidence':.9}
        result={**row,'api_status':'success','visual_facts':facts,'vision_revision_key':vision_revision(cat,row),'model':'OFFLINE_FIXTURE','usage':{}}
        (base/'stage123'/'results.jsonl').write_text(json.dumps(result)+'\n',encoding='utf-8')
        exporter.export(category,'layer123')
        assert selection.completion(config,True)['缺失或失败']==0
        from openpyxl import load_workbook
        book=load_workbook(base/'excel'/'电钻_第一二三层.xlsx',read_only=True)
        sheet=book['全量综合结果']; values=list(sheet.values); final=values[3].index('最终判断')
        assert {r[final] for r in values[4:]}=={'确认同款','不同款（文字硬冲突）'}
        book.close()
        # Final collector snapshots must be merged before completion is allowed.
        snapshots=base/'temu/workers/worker-01/query';snapshots.mkdir(parents=True)
        (snapshots/'products.jsonl').write_text('new worker snapshot',encoding='utf-8')
        assert selection.completion(config,True) is None
        # Changed source bytes invalidate completion before the ingestion cycle.
        amazon.write_text(amazon.read_text(encoding='utf-8').replace('cordless','brushless cordless'),encoding='utf-8')
        assert selection.completion(config,False) is None
    finally: pipeline.close()


def test_validate_only_never_calls_provider(tmp_path):
    from PIL import Image
    image=tmp_path/'image.png'; Image.new('RGB',(8,8),'blue').save(image)
    candidate=tmp_path/'candidates.jsonl'
    candidate.write_text(json.dumps({'asin':'A','goods_id':'T','pair_revision':'p','amazon_image_path':str(image),'temu_image_path':str(image)})+'\n',encoding='utf-8')
    import os
    env={**os.environ,'PYTHONUTF8':'1','DEEPSEEK_API_KEY':''}
    result=subprocess.run([sys.executable,str(ROOT/'scripts/run_category_layer3.py'),'--category',str(ROOT/'config/categories/power_drill.json'),'--candidates',str(candidate),'--results',str(tmp_path/'results.jsonl'),'--validate-only'],env=env,capture_output=True,text=True,encoding='utf-8')
    assert result.returncode==0,result.stderr
    assert json.loads(result.stdout)['API请求数']==0
    assert not (tmp_path/'requests.jsonl').exists()


@pytest.mark.parametrize('mode,expected_requests,expected_success',[('split',3,2),('high',2,1),('auth',1,0)])
def test_paid_consumer_with_fake_provider_only(tmp_path,monkeypatch,capsys,mode,expected_requests,expected_success):
    from types import SimpleNamespace
    from PIL import Image
    import run_category_layer3 as consumer
    import openai
    cat=copy.deepcopy(CAT);cat['stage3']['max_concurrency']=1
    category=tmp_path/'category.json';category.write_text(json.dumps(cat),encoding='utf-8')
    monkeypatch.setattr(consumer,'ROOT',tmp_path)
    marker=tmp_path/'runs'/cat['category_key']/'stage123/paid_enabled.json'
    marker.parent.mkdir(parents=True,exist_ok=True)
    marker.write_text('{}',encoding='utf-8')
    image=tmp_path/'fixture.png';Image.new('RGB',(16,16),'blue').save(image)
    candidate=tmp_path/'candidates.jsonl';result=tmp_path/'results.jsonl'
    count=2 if mode=='split' else 1
    candidate.write_text(''.join(json.dumps({'asin':'A','goods_id':str(i),'pair_revision':str(i),'amazon_image_path':str(image),'temu_image_path':str(image)})+'\n' for i in range(count)),encoding='utf-8')
    calls=[]
    class ProviderError(Exception): status_code=401
    class FakeClient:
        def __init__(self,**kw): self.responses=self
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def create(self,**kw):
            calls.append(kw)
            if mode=='auth':raise ProviderError('fixture auth failure')
            ids=[part['text'] for part in kw['input'][0]['content'] if part['type']=='input_text' and part['text'].startswith('C')]
            low=mode=='high' and kw['reasoning']['effort']=='low'
            raw={'candidates':[{'i':i,'v':identity('DIFFERENT' if low else 'SAME'),'c':.5 if low else .9,'m':{str(j):{'r':'D' if low else 'S','a':'明确结构','t':'对应结构','h':cat['stage3']['identity_hard_conflicts'][0] if low else 'NONE'} for j in range(len(cat['stage3']['modules']))}} for i in ids]}
            return SimpleNamespace(output_text='invalid json' if mode=='split' and len(calls)==1 else json.dumps(raw),usage=SimpleNamespace(model_dump=lambda:{'total_tokens':101}),id=f'fixture-{len(calls)}')
    monkeypatch.setattr(openai,'OpenAI',FakeClient)
    monkeypatch.setenv('DEEPSEEK_API_KEY','fixture-only-not-a-real-key')
    monkeypatch.setattr(sys,'argv',['consumer','--category',str(category),'--candidates',str(candidate),'--results',str(result)])
    if mode=='auth':
        with pytest.raises(SystemExit,match='未结算视觉请求'):
            consumer.main()
    else:
        consumer.main()
    assert len(calls)==expected_requests
    rows=list(consumer.rows(result))
    assert sum(r['api_status']=='success' for r in rows)==expected_success
    ledger=list(consumer.rows(tmp_path/'requests.jsonl'))
    assert len(ledger)==expected_requests
    assert sum(r.get('usage',{}).get('total_tokens',0) for r in ledger)==(0 if mode=='auth' else expected_requests*101)
    if mode=='auth':
        with pytest.raises(SystemExit,match='未结算视觉请求|未结算的视觉请求'):consumer.main()
    else:
        consumer.main()
    assert len(calls)==expected_requests
