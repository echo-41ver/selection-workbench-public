"""Complete current Agent facts supplement only inconclusive decision-only API output."""
from copy import deepcopy
import json
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import export_category_workbooks as exporter
import run_codex_visual_pool as pool
from selection_automation import commercial
from selection_automation.codex_visual_review import import_reviews
from selection_automation.stage3_progress import snapshot
from selection_automation.stage3_provider import apply, settings
from selection_automation.vision import VisionRevisionContext, compatible_vision_keys
from prepare_codex_visual_batch import prepare_batch
from test_prepare_codex_visual_batch import setup
from test_set_inventory_identity import facts


def saved_evidence(tmp_path, monkeypatch, *, complete_api=False, decision_known=False, agent=True):
    cat, base, path, pairs = setup(tmp_path, monkeypatch, [{'index': 1, 'asin': 'A1'}])
    cat.update(execution_mode='full', commercial={})
    cat['temu']['base_queries'] = ['knives']
    cat['stage3']['codex_require_module_facts'] = True
    apply(cat, settings({'mode': 'liquid_d1', 'workers': 1, 'concurrency_mode': 'manual'}))
    pair = pairs[0]
    api_key = compatible_vision_keys(cat, pair, context=VisionRevisionContext(cat))[0]
    api = {**pair, 'vision_revision_key': api_key, 'api_status': 'success', 'model': 'existing-api',
           'visual_facts': facts(cat, 8, 6) if complete_api else {'decision_api': {
               'choice': 'same', 'probabilities': {'same': .96 if decision_known else .62,
                                                  'similar': .02 if decision_known else .18,
                                                  'different': .02 if decision_known else .20}}}}
    if agent:
        apply(cat, settings({'mode': 'local_agent', 'agent_ids': ['codex'], 'workers': 1,
                             'concurrency_mode': 'manual'}))
    path.write_text(json.dumps(cat), encoding='utf-8')
    revision = compatible_vision_keys(cat, pair, context=VisionRevisionContext(cat))[0]
    bound = {**pair, 'vision_revision_key': revision}
    folder = base / 'stage123'; folder.mkdir(exist_ok=True)
    (folder / 'results.jsonl').write_text(json.dumps(api) + '\n', encoding='utf-8')
    (base / 'stage12' / f'layer3_candidates_{base.name}.jsonl').write_text(json.dumps(pair) + '\n', encoding='utf-8')
    (base / 'amazon').mkdir(); (base / 'temu/union').mkdir(parents=True)
    (base / 'amazon/amazon_frontend_latest.json').write_text(json.dumps({'records': [
        {'asin': pair['asin'], 'title': 'TEST_ONLY_AMAZON', 'product_url': 'https://www.amazon.com/dp/A1'}]}), encoding='utf-8')
    (base / 'temu/union/products.jsonl').write_text(json.dumps({'goods_id': pair['goods_id'],
        'title': 'TEST_ONLY_TEMU', 'product_url': 'https://www.temu.com/goods.html?goods_id=G1'}) + '\n', encoding='utf-8')
    frozen = folder / 'export_frozen'; frozen.mkdir()
    for name in ('layer3_candidates', 'layer12_rows'):
        (frozen / f'{name}_{base.name}.jsonl').write_text(json.dumps({**pair, 'pair_status': 'CANDIDATE'}) + '\n', encoding='utf-8')
    (frozen / f'layer12_summary_{base.name}.json').write_text(json.dumps({
        'actual_pairs': 1, 'amazon_products': 1, 'temu_products': 1, 'states': {'CANDIDATE': 1}}), encoding='utf-8')
    review = {**bound, 'review_status': 'SAME', 'amazon_evidence': '前排六把牛排刀',
              'temu_evidence': '前排六把牛排刀', 'reason': '双方前排六把牛排刀与刀座结构逐项一致',
              'reviewer': 'saved-original-job', 'model': 'GPT-6.1-Sol', 'visual_facts': facts(cat)}
    output = folder / 'codex_reviews/accepted.jsonl'
    job = folder / 'saved-job'; job.mkdir()
    pool._write_jsonl(job / 'manifest.jsonl', [bound]); pool.bind_original_images(job, [bound])
    pool._write_jsonl(job / 'reviews.jsonl', [review])
    assert pool.import_job(cat, {'task_id': 'accepted', 'rows': [bound], 'output': output,
                                'raw_reviews': job / 'raw.jsonl'}, job) == (1, [])
    monkeypatch.setattr(exporter, 'ROOT', tmp_path)
    monkeypatch.setattr(exporter, 'load_category', lambda _: cat)
    return cat, base, path, pair, output


@pytest.mark.parametrize('agent', [False, True])
def test_full_agent_supplement_agrees_in_progress_report_and_excel(tmp_path, monkeypatch, agent):
    cat, base, path, pair, output = saved_evidence(tmp_path, monkeypatch, agent=agent)
    api_before = (base / 'stage123/results.jsonl').read_bytes()
    review_before = output.read_bytes()
    progress = snapshot(base, cat, history=False)
    assert progress['codex_same'] == progress['completed'] == 1
    assert progress['api_completed'] == progress['pending'] == 0
    assert prepare_batch(path, skip_empty=True)['selected'] == 0
    report = commercial.current_report(tmp_path, cat)
    entry = report['products'][0]
    assert entry['match']['confirmed'] == 1 and entry['match']['possible'] == 0
    assert entry['confirmed_pairs'][0]['model'] == 'GPT-6.1-Sol'
    delivered = exporter.export(path, 'layer123')
    book = load_workbook(delivered['固定查看副本'])
    try:
        sheet = book['采集分析进度']; headers = [c.value for c in sheet[4]]
        assert sheet.cell(5, headers.index('同款数') + 1).value == 1
        detail = book['同款商品明细']; headers = [c.value for c in detail[4]]
        assert 'GPT-6.1-Sol' in detail.cell(5, headers.index('分析模型') + 1).value
    finally:
        book.close()
    assert (base / 'stage123/results.jsonl').read_bytes() == api_before
    assert output.read_bytes() == review_before
    assert 'supersedes_api' not in json.loads(output.read_text('utf-8'))


@pytest.mark.parametrize('kind', ['complete-api', 'known-decision', 'stale-review', 'wrong-pair', 'incomplete-review', 'changed-image'])
def test_existing_valid_api_and_invalid_or_stale_review_keep_api_priority(tmp_path, monkeypatch, kind):
    cat, base, _, pair, output = saved_evidence(tmp_path, monkeypatch,
        complete_api=kind == 'complete-api', decision_known=kind == 'known-decision')
    review = json.loads(output.read_text('utf-8'))
    if kind == 'stale-review': review['vision_revision_key'] = 'stale'
    elif kind == 'wrong-pair': review['goods_id'] = 'OTHER'
    elif kind == 'incomplete-review': review['visual_facts'] = {}
    elif kind == 'changed-image':
        from PIL import Image
        Image.new('RGB', (80, 60), 'green').save(pair['temu_image_path'])
    output.write_text(json.dumps(review) + '\n', encoding='utf-8')
    report = commercial.current_report(tmp_path, cat)
    entry = report['products'][0]
    if kind == 'complete-api':
        assert entry['match']['different'] == 1 and entry['match']['confirmed'] == 0
    elif kind == 'known-decision':
        assert entry['confirmed_pairs'][0]['model'] == 'existing-api'
    elif kind == 'changed-image':
        assert entry['match']['pending'] == 1
    else:
        assert entry['match']['possible'] == 1 and entry['match']['confirmed'] == 0
        assert report['review_pairs'][0]['model'] == 'existing-api'


def test_full_api_precedence_is_unchanged_in_progress_and_batch(tmp_path, monkeypatch):
    cat, base, path, _, _ = saved_evidence(tmp_path, monkeypatch, complete_api=True)
    progress = snapshot(base, cat, history=False)
    assert progress['api_completed'] == progress['completed'] == 1
    assert progress['codex_completed'] == 0
    assert prepare_batch(path, skip_empty=True)['selected'] == 0


@pytest.mark.parametrize('complete_first', [False, True])
def test_complete_api_is_not_hidden_by_another_decision_only_success(tmp_path, monkeypatch, complete_first):
    cat, base, _, _, _ = saved_evidence(tmp_path, monkeypatch, complete_api=True)
    file = base / 'stage123/results.jsonl'
    complete = json.loads(file.read_text('utf-8'))
    decision = {**complete, 'visual_facts': {'decision_api': {'choice': 'same',
        'probabilities': {'same': .62, 'similar': .18, 'different': .20}}}}
    file.write_text(''.join(json.dumps(row) + '\n' for row in ([complete, decision] if complete_first else [decision, complete])), encoding='utf-8')
    current = snapshot(base, cat, history=False)
    assert current['api_completed'] == 1 and current['codex_completed'] == 0
    match = commercial.current_report(tmp_path, cat)['products'][0]['match']
    assert match['different'] == 1 and match['confirmed'] == 0


@pytest.mark.parametrize('entry', ['open', 'download'])
@pytest.mark.parametrize('existing', [False, True])
def test_normal_excel_entry_publishes_current_content_and_isolates_the_visible_task(tmp_path, monkeypatch, entry, existing):
    import io
    import threading
    import urllib.request
    from types import SimpleNamespace
    import web_console as web
    cat, base, path, _, output = saved_evidence(tmp_path, monkeypatch)
    review = output.read_bytes(); output.unlink()
    old = exporter.export(path, 'layer123')
    web.selection.save(base / 'publish_state.json', {'deliveries': {'layer123': old}})
    old_book = Path(old['固定查看副本']); original = old_book.read_bytes()
    with_copy = tmp_path / 'runs/other/excel/do-not-publish.xlsx'
    with_copy.parent.mkdir(parents=True); with_copy.write_bytes(original)
    output.write_bytes(review)
    if not existing:
        for book in old_book.parent.glob('*.xlsx'): book.unlink()
    console = web.Console(tmp_path); console.key = base.name
    monkeypatch.setattr(console, 'owner', lambda: None)
    opened = []
    monkeypatch.setattr(web, 'require_excel_open_association', lambda: None)
    monkeypatch.setattr(web.os, 'startfile', lambda p: opened.append(Path(p)), raising=False)
    def publish(args, **kwargs):
        assert Path(args[1]).name == 'publish_all.py'
        assert Path(args[args.index('--category') + 1]) == path
        assert args[args.index('--kind') + 1] == 'layer123'
        delivered = exporter.export(path, 'layer123')
        web.selection.save(base / 'publish_state.json', {'deliveries': {'layer123': delivered}})
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(web.subprocess, 'run', publish)
    if entry == 'open':
        assert console.open_excel('layer123')['ok']
        received = opened[0].read_bytes()
    else:
        server = web.ThreadingHTTPServer(('127.0.0.1', 0), web.Handler); server.console = console
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            received = opener.open('http://127.0.0.1:' + str(server.server_port) + '/api/excel/layer123', timeout=5).read()
        finally:
            server.shutdown(); server.server_close(); thread.join(3)
    assert received != original
    book = load_workbook(io.BytesIO(received))
    try:
        sheet = book['采集分析进度']; headers = [c.value for c in sheet[4]]
        assert sheet.cell(5, headers.index('同款数') + 1).value == 1
    finally: book.close()
    assert with_copy.read_bytes() == original


def test_failed_refresh_keeps_the_previous_workbook_and_reports_failure(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import web_console as web
    _, base, path, _, _ = saved_evidence(tmp_path, monkeypatch)
    old = exporter.export(path, 'layer123')
    web.selection.save(base / 'publish_state.json', {'deliveries': {'layer123': old}})
    file = Path(old['固定查看副本']); before = file.read_bytes()
    console = web.Console(tmp_path); console.key = base.name
    monkeypatch.setattr(console, 'owner', lambda: None)
    def failed(*a, **kw):
        web.selection.save(base / 'publish_state.json', {'errors': {'analysis': 'TEST_ONLY_FAILURE'}, 'deliveries': {'layer123': old}})
        return SimpleNamespace(returncode=1)
    monkeypatch.setattr(web.subprocess, 'run', failed)
    monkeypatch.setattr(web.os, 'startfile', lambda *_: pytest.fail('failed publish must not silently open a stale workbook'), raising=False)
    with pytest.raises(web.ConsoleError) as error:
        console.open_excel('layer123')
    assert error.value.code == 'EXCEL_UPDATE_FAILED'
    assert file.read_bytes() == before


@pytest.mark.parametrize('legacy_facts', [None, [], 'TEST_ONLY_NON_DICT', 42])
def test_legacy_non_dict_api_facts_keep_the_original_qualification_guard(legacy_facts):
    from selection_automation.codex_visual_review import qualified_api_facts
    from selection_automation.stage3_progress import _qualified_api_facts
    cat = {'stage3': {'modules': {}}}
    row = {'api_status': 'success', 'goods_id': 'T1', 'visual_facts': legacy_facts}
    assert qualified_api_facts(cat, row) is False
    assert _qualified_api_facts(cat, row) is False
