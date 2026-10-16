import json
from pathlib import Path
import sys
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from export_selection_result import export
from selection_automation.selection_delivery import build_delivery
from selection_automation.stage3_provider import settings, apply, VisualClient, request_key
from selection_automation.rules import score
from selection_automation.operations import price_range


def test_collected_parameters_are_readable_chinese_and_do_not_mutate_sources():
    from copy import deepcopy
    from selection_automation.selection_delivery import collected_parameters,display_parameter,SPECIFICATION_NAMES
    raw={'detail_price_text':'$11.11','brand':'Blue','monthly_sales_text':'10K+ bought in past month',
         'specifications':[{'label':'Color','value':'Black'},{'label':'Brand','value':'Blue'},
                           {'label':'Number of Keys','value':105}],
         'variants':[{'label':'Keyboard /* Temporary CSS */ .broken { color:red; }','asin':'A1'}],
         'image_candidates':[{'url':'https://example.com/raw.png'}],
         'primary_observation':{'snapshot':r'D:\internal\products.jsonl'},
         'merge_policy':'LATEST_COMPLETE_OBSERVATION_WITH_CONFLICT_PROVENANCE',
         'category_relevance':{'state':'RELEVANT'},'detail_quality':{'schema_version':'internal'}}
    original=deepcopy(raw);values=collected_parameters(raw)
    assert values['详情页价格']=='$11.11' and values['品牌']=='Blue'
    assert values['近一月页面销量']=='近一月购买：10,000+ 件'
    assert values['规格参数']=='颜色：黑色\n品牌：Blue\n按键数：105'
    assert values['已采集变体']=='键盘：A1'
    assert len(values)==5 and raw==original
    assert SPECIFICATION_NAMES['Water Resistance Level']=='防水等级'
    assert display_parameter({'min':0,'max':12,'currency':'USD','api_key':'TEST_ONLY_SECRET'})=='最低价：0\n最高价：12\n币种：美元'
    assert display_parameter('In Stock {"isInternal":false,"merchantId":"fixture"}')=='有货'
    assert display_parameter("4.4 4.4 out of 5 stars (3,828) var dpAcr; P.when('A').execute(function() {});")=='评分：4.4（满分5星），3,828条评价'


def test_export_does_not_expose_collector_columns_or_json_blobs(tmp_path):
    cat,report=setup(tmp_path)
    path=tmp_path/'runs/fixture/amazon/amazon_frontend_latest.json'
    source=json.loads(path.read_text('utf8'));source['records'][0].update(
        detail_price_text='$12.00',image_candidates=[{'url':'https://example.com/image.png'}],
        variants=[{'label':'Keyboard /* Temporary CSS */ bad{}','asin':'TEST_ASIN'}])
    path.write_text(json.dumps(source),encoding='utf8')
    result=export(tmp_path,cat,report);book=load_workbook(result['固定查看副本'])
    for sheet in book:
        headers=[c.value for c in sheet[4]]
        assert not any(isinstance(h,str) and any(k in h for k in ('detail_price_text','image_candidates','primary_observation')) for h in headers)
        strings=[str(c.value) for row in sheet.iter_rows(min_row=5) for c in row if c.value is not None]
        assert not any('Temporary CSS' in value or '"label"' in value or '"value"' in value or 'schema_version' in value for value in strings)
    book.close()


def test_current_report_refreshes_legacy_display_without_reanalysing(tmp_path,monkeypatch):
    from selection_automation import commercial
    cat,report=setup(tmp_path)
    report.update(evidence_signature='same-evidence')
    report['products'][0]['amazon_parameters']={'detail_quality':'legacy JSON'}
    path=tmp_path/'runs/fixture/commercial_report.json'
    path.write_text(json.dumps(report),encoding='utf8')
    monkeypatch.setattr(commercial,'report_signature',lambda *a:'same-evidence')
    monkeypatch.setattr(commercial,'build_report',lambda *a:pytest.fail('Cached analysis recomputed'))
    result=commercial.current_report(tmp_path,cat)
    assert result['products'][0]['match']==report['products'][0]['match']
    assert result['products'][0]['amazon_parameters']['规格参数']=='材质：棉'
    assert 'detail_quality' not in result['products'][0]['amazon_parameters']


def setup(root, n=2):
    base=root/'runs/fixture'
    (base/'amazon').mkdir(parents=True)
    (base/'temu/union').mkdir(parents=True)
    cat={'category_key':'fixture','category_name_zh':'测试商品','commercial':{'same_limit':n},'stage3':{'modules':{'design':100}}}
    amazon=[{'asin':a,'title':a+' title','price_usd':12,'product_url':'https://www.amazon.com/dp/'+a,
             'monthly_sales_text':'100+ bought in past month','rating':4.5,'review_count':10,
             'specifications':[{'label':'材质','value':'棉'}]} for a in ['A1','A2','A3']]
    (base/'amazon/amazon_frontend_latest.json').write_text(json.dumps({'records':amazon,'observed_at':'2026-10-14T01:00:00Z'}),encoding='utf-8')
    temu=[{'goods_id':str(g),'title':'same product','price_usd':4,'rating':4.7,'review_count':0,
           'product_url':'https://www.temu.com/goods.html?goods_id='+str(g),'sales_text':'1K+ sold',
           'price_range':{'min':3,'max':8,'currency':'USD','basis':'页面明确区间'},
           'specifications':{'材质':'棉'},'long_detail':'证据'*18000} for g in [1,2]]
    (base/'temu/union/products.jsonl').write_text('\n'.join(json.dumps(p,ensure_ascii=False) for p in temu),encoding='utf-8')
    products=[]
    for a,gids,state,pending in [('A1',['1'],'SAMPLE_PASS',0),('A2',['1','2'],'SAMPLE_PASS',0),('A3',[],'UNKNOWN',2)]:
        products.append({'asin':a,'confirmed_goods_ids':gids,'match':{'confirmed':len(gids),'state':state,'pending':pending,'possible':0},
                         'confirmed_pairs':[{'goods_id':g,'state':'确认同款','reason':'fixture visual fact','model':'fixture-model'} for g in gids]})
    return cat,{'products':products,'market':{},'as_of':'2026-10-14T02:00:00Z'}


def test_strict_threshold_and_unfinished_evidence(tmp_path):
    cat,report=setup(tmp_path)
    data=build_delivery(tmp_path,cat,report)
    assert [p['asin'] for p in data['products']]==['A1','A3']
    assert [(p['asin'],p['goods_id']) for p in data['pairs']]==[('A1','1'),('A2','1'),('A2','2')]
    assert data['products'][1]['advice']['conclusion']=='待完成同款核验'
    assert data['products'][1]['advice']['confidence']=='未知'
    assert data['products'][0]['amazon']['price']['min'] is None


def test_reopen_focused_excel_and_full_fields(tmp_path):
    cat,report=setup(tmp_path)
    delivered=export(tmp_path,cat,report)
    book=load_workbook(delivered['固定查看副本'],read_only=False,data_only=False)
    assert book.sheetnames==['选品结论','同款商品明细','商品完整字段']
    assert book['选品结论'].max_row==6 and book['同款商品明细'].max_row==7
    assert book['选品结论']['A5'].value=='A1' and book['选品结论']['A6'].value=='A3'
    detail=book['同款商品明细'];headers=[c.value for c in detail[4]]
    assert detail.cell(5,headers.index('TEMU最低价(USD)')+1).value==3
    assert detail.cell(5,headers.index('TEMU最高价(USD)')+1).value==8
    assert detail.cell(5,headers.index('TEMU评论数')+1).value==0
    link=detail.cell(5,headers.index('TEMU商品链接')+1)
    assert link.hyperlink and link.hyperlink.target.startswith('https://www.temu.com/') or link.data_type=='f' and 'https://www.temu.com/' in link.value
    text=''.join(row[3] for row in book['商品完整字段'].iter_rows(min_row=5,values_only=True) if row[0]=='TEMU' and row[1]=='1' and row[2]=='补充说明')
    assert text=='证据'*18000
    assert book['选品结论'].freeze_panes=='D5'
    book.close()


def test_empty_selection_has_no_cartesian_detail(tmp_path):
    cat,report=setup(tmp_path,n=0)
    delivery=export(tmp_path,cat,report)
    book=load_workbook(delivery['固定查看副本'],read_only=True)
    assert book['选品结论'].max_row==4 and book['同款商品明细'].max_row==7
    book.close()


def test_summary_has_five_counts_and_visible_collected_parameters(tmp_path):
    from openpyxl.utils import get_column_letter
    from selection_automation.selection_delivery import conclusion_label
    assert [conclusion_label(value) for value in
        ('确认同款','相似款（Agent看图核验）','DIFFERENT','待补充证据',None)]==['同款','类似款','不同款','无法判断','未分析']
    cat,report=setup(tmp_path,n=4)
    first,second,third=report['products']
    first['confirmed_pairs'][0]['reason']='简短事实。'*100
    second.update(confirmed_goods_ids=[],confirmed_pairs=[])
    second['match'].update(confirmed=0,different=2)
    delivered=export(tmp_path,cat,report)
    book=load_workbook(delivered['固定查看副本'])
    sheet=book['选品结论'];headers=[cell.value for cell in sheet[4]]
    assert [sheet.cell(row,headers.index('同款数')+1).value for row in (5,6,7)]==[1,0,0]
    assert sheet.cell(6,headers.index('不同款数')+1).value==2
    assert sheet.cell(7,headers.index('未分析数')+1).value==2
    assert all(h in headers for h in ['类似款数','无法判断数','Amazon规格参数','Amazon页面销量','Amazon评分'])
    assert not sheet.column_dimensions[get_column_letter(headers.index('Amazon评分')+1)].hidden
    assert len(book['同款商品明细'].cell(5,12-1).value)<=140
    assert book['商品完整字段'].sheet_state=='hidden'
    book.close()


def test_combined_same_and_similar_threshold_and_embedded_picture_order(tmp_path):
    from PIL import Image
    cat,report=setup(tmp_path,n=3)
    selected=report['products'][0]
    selected['match'].update(similar=1)
    selected['similar_goods_ids']=['2']
    selected['similar_pairs']=[{'goods_id':'2','state':'相似款','reason':'TEST_ONLY_SIMILAR'}]
    data=build_delivery(tmp_path,cat,report)
    item=next(p for p in data['products'] if p['asin']=='A1')
    assert item['same_count']==2 and item['confirmed_count']==item['similar_count']==1
    assert {(p['goods_id'],p['relation']) for p in data['pairs'] if p['asin']=='A1'}=={('1','确认同款'),('2','相似款')}
    picture=tmp_path/'TEST_ONLY_IMAGE.png';Image.new('RGB',(120,80),'blue').save(picture)
    base=tmp_path/'runs/fixture'
    am_file=base/'amazon/amazon_frontend_latest.json';am=json.loads(am_file.read_text('utf-8'))
    for row in am['records']:
        row.update(local_image_path=str(picture),main_image_url='https://images.example.invalid/TEST_ONLY_IMAGE.png')
    am_file.write_text(json.dumps(am),encoding='utf-8')
    temu_file=base/'temu/union/products.jsonl'
    rows=[json.loads(line) for line in temu_file.read_text('utf-8').splitlines()]
    for row in rows:row.update(local_image_path=str(picture),main_image_url='https://images.example.invalid/TEST_ONLY_IMAGE.png')
    temu_file.write_text('\n'.join(json.dumps(row) for row in rows),encoding='utf-8')
    delivery=export(tmp_path,cat,report)
    book=load_workbook(delivery['固定查看副本'])
    prefix=['Amazon ASIN','Amazon标题','Amazon图片','TEMU图片','同款+类似款数量','Amazon商品链接','TEMU商品链接']
    assert [c.value for c in book['同款商品明细'][4]][:7]==prefix
    assert [c.value for c in book['选品结论'][4]][:3]==prefix[:3]
    assert len(book['选品结论']._images)==3
    assert {(image.anchor._from.col) for image in book['选品结论']._images}=={2}
    assert len(book['同款商品明细']._images)==8
    assert {(image.anchor._from.col) for image in book['同款商品明细']._images}=={2,3}
    assert book['同款商品明细']['E5'].value==2
    book.close()
    from zipfile import ZipFile
    with ZipFile(delivery['固定查看副本']) as archive:
        drawings=[archive.read(n).decode() for n in archive.namelist() if n.startswith('xl/drawings/drawing') and n.endswith('.xml')]
        assert all('hlinkClick' in drawing for drawing in drawings)
        relations=''.join(archive.read(n).decode() for n in archive.namelist() if n.startswith('xl/drawings/_rels/'))
        assert 'https://images.example.invalid/TEST_ONLY_IMAGE.png' in relations
    cat['commercial']['same_limit']=2
    assert all(p['asin']!='A1' for p in build_delivery(tmp_path,cat,report)['products'])
    cat['commercial']['same_limit']=3
    selected['similar_goods_ids']=['1']
    with pytest.raises(ValueError,match='同时记为'):build_delivery(tmp_path,cat,report)


def test_native_similar_is_a_final_relation_not_an_unknown_or_same():
    cat={'stage3':{'modules':{'design':100}}};apply(cat,settings({'mode':'liquid_d1'}))
    facts={'decision_api':{'choice':'similar','probabilities':{'same':.02,'similar':.96,'different':.02},'protocol_version':'product-choice-v2'}}
    assert score(cat,facts)['state']=='相似款'
    facts['decision_api']['probabilities']={'same':.3,'similar':.5,'different':.2}
    assert score(cat,facts)['state']=='待补充证据'


@pytest.mark.parametrize('mode',['liquid_d1','multimodal_api','local_model','local_agent'])
def test_four_modes_have_no_stored_key(mode):
    config=settings({'mode':mode})
    assert config['mode']==mode and 'api_key' not in config


def test_key_cannot_follow_another_endpoint(monkeypatch):
    monkeypatch.setattr('selection_automation.stage3_credentials.saved_key',lambda root,provider:'')
    monkeypatch.setenv('DEEPSEEK_API_KEY','TEST_ONLY_COMPANY_KEY')
    with pytest.raises(ValueError):request_key(settings({'mode':'liquid_d1'}))
    monkeypatch.setenv('SELECTION_STAGE3_API_KEY','TEST_ONLY_CUSTOM_KEY')
    monkeypatch.setenv('SELECTION_STAGE3_ENDPOINT','https://example.invalid/v1')
    assert request_key(settings({'base_url':'https://example.invalid/v1','model':'vision'}))=='TEST_ONLY_CUSTOM_KEY'
    with pytest.raises(ValueError):request_key(settings({'base_url':'https://other.invalid/v1','model':'vision'}))


def test_native_decision_payload_and_uncertainty():
    provider=settings({'mode':'liquid_d1'})
    calls=[]
    probabilities={'same':.96,'different':.03,'similar':.01}
    response=SimpleNamespace(raise_for_status=lambda:None,json=lambda:{'answers':{'same_product':{'choice':'same','probabilities':probabilities}},'usage':{'input_tokens':42,'output_tokens':0}})
    transport=SimpleNamespace(post=lambda *args,**kwargs:calls.append((args,kwargs)) or response)
    client=VisualClient(provider,transport,'TEST_ONLY_KEY',30)
    result=client.create(model='d1',instructions='fixture identity rules',input=[{'content':[
        {'type':'input_image','image_url':'data:image/png;base64,TEST_ONLY_IMAGE'},
        {'type':'input_image','image_url':'data:image/png;base64,TEST_ONLY_IMAGE'}]}])
    payload=json.loads(calls[0][1]['content'])
    assert len(payload['images'])==2 and set(payload['questions']['same_product']['criteria'])=={'same','different','similar'}
    cat={'stage3':{'modules':{'design':100}}};apply(cat,provider)
    facts=json.loads(result.output_text)['decision_facts']
    assert score(cat,facts)['state']=='确认同款'
    assert score(cat,facts)['total_score'] is None
    facts['decision_api']['probabilities']={'same':.6,'different':.3,'similar':.1}
    assert score(cat,facts)['state']=='待补充证据'


def test_bad_config_and_partial_variant_band():
    with pytest.raises(ValueError):settings({'mode':'liquid_d1','model':'d1:free'})
    with pytest.raises(ValueError):settings({'mode':'local_model','base_url':'https://remote.invalid/v1'})
    with pytest.raises(ValueError):settings({'base_url':'https://TEST_ONLY_USER:TEST_ONLY_PASSWORD@proxy.example/v1'})
    band=price_range({'skus':[{'price_usd':3},{'price_usd':7},{'price_usd':None}],'all_variants_priced':True})
    assert (band['min'],band['max'])==(3,7) and '子集' in band['basis']


@pytest.mark.parametrize('mode',['liquid_d1','multimodal_api','local_model','local_agent'])
def test_console_dispatches_only_selected_provider_and_never_stops_collection(tmp_path,monkeypatch,mode):
    import web_console as web
    import run_codex_visual_pool as agent
    console=web.Console(tmp_path)
    base=tmp_path/'runs/fixture';base.mkdir(parents=True)
    category={'category_key':'fixture','execution_mode':'full','stage3':{'modules':{'design':100}}}
    (base/'category.json').write_text(json.dumps(category),encoding='utf-8')
    monkeypatch.setattr(console,'current_run_key',lambda:('fixture',False))
    monkeypatch.setattr(console,'stage3_snapshot',lambda key:{'total':3,'pending':3})
    monkeypatch.setattr(web,'codex_review_control',lambda base:{'running':False})
    monkeypatch.setattr(console,'saved_credential',lambda:'TEST_ONLY_COMPANY_KEY')
    calls=[]
    monkeypatch.setattr(console,'invoke',lambda *a,**kw:pytest.fail('Collection command invoked'))
    monkeypatch.setattr(console,'launch_stage3_controller',lambda key,path,token:calls.append(('api',key,token)))
    monkeypatch.setattr(agent,'codex_executable',lambda config:'fixture-codex.exe')
    monkeypatch.setattr('selection_automation.local_agents.ready_agents',lambda *a,**kw:[{'id':'codex','name':'Codex CLI','executable':'fixture-codex.exe','ready':True}])
    monkeypatch.setattr(agent,'ensure_pool',lambda path,root:calls.append(('agent',root)))
    provider=settings({'mode':mode})
    body={'provider':provider}
    if mode in {'multimodal_api','liquid_d1'}:body['api_key']='TEST_ONLY_SELECTED_KEY'
    result=console.start_stage3(body)
    assert result['ok'] and result['provider']['mode']==mode
    if mode=='local_agent':
        assert calls==[('agent',tmp_path)]
        config=json.loads((base/'stage123/codex_pool_config.json').read_text(encoding='utf-8'))
        assert config['enabled'] and 1<=config['workers']<=6 and config['agent_ids']==['codex']
    else:
        assert calls==[('api','fixture','TEST_ONLY_SELECTED_KEY' if 'api_key' in body else '')]
    assert 'TEST_ONLY_SELECTED_KEY' not in (base/'category.json').read_text()
    # Manual continuation can reuse only the same run and service's in-memory key.
    if mode=='liquid_d1':
        console.stage3_busy=False
        assert console.start_stage3({'provider':provider})['ok']
        assert calls[-1]==('api','fixture','TEST_ONLY_SELECTED_KEY')


def test_save_d1_provider_and_key_does_not_enable_or_start_paid_analysis(tmp_path,monkeypatch):
    import web_console as web
    console=web.Console(tmp_path);base=tmp_path/'runs/fixture';base.mkdir(parents=True)
    category={'category_key':'fixture','execution_mode':'full','stage3':{'modules':{'design':100}}}
    (base/'category.json').write_text(json.dumps(category),encoding='utf-8')
    monkeypatch.setattr(console,'current_run_key',lambda:('fixture',True))
    monkeypatch.setattr(console,'stage3_snapshot',lambda key:{'counting':True})
    monkeypatch.setattr(web,'codex_review_control',lambda base:{'running':False})
    monkeypatch.setattr(console,'invoke',lambda *a,**kw:pytest.fail('Collection touched'))
    monkeypatch.setattr(console,'launch_stage3_controller',lambda *a:pytest.fail('Paid analysis started'))
    result=console.start_stage3({'provider':settings({'mode':'liquid_d1'}),'api_key':'TEST_ONLY_D1_KEY'},configure_only=True)
    assert result['ok']
    assert console.stage3_keys[('fixture',settings({'mode':'liquid_d1'})['base_url'])]=='TEST_ONLY_D1_KEY'
    assert not (base/'stage123/paid_enabled.json').exists()
    assert not (base/'stage123/authorization.json').exists()
    assert 'TEST_ONLY_D1_KEY' not in (base/'category.json').read_text()


def test_daitu_gpt_uses_task_form_native_stream_and_records_usage(monkeypatch):
    import openai
    monkeypatch.setattr(openai,'OpenAI',lambda **kw:SimpleNamespace())
    provider=settings({'base_url':'https://api.daitu.cc','model':'gpt-6.1-sol','protocol':'chat_completions'})
    assert provider['base_url']=='https://api.daitu.cc/v1' and provider['protocol']=='responses'
    complete={'id':'response_fixture','status':'completed','usage':{'input_tokens':2,'output_tokens':1,'total_tokens':3},
              'output':[{'type':'message','role':'assistant','content':[{'type':'output_text','text':'{"candidates":[]}'}]}]}
    response=SimpleNamespace(headers={'content-type':'text/event-stream'},raise_for_status=lambda:None,
        iter_lines=lambda:iter(['data: '+json.dumps({'type':'response.created','response':{'id':'response_fixture','status':'in_progress'}}),'',
                               'data: '+json.dumps({'type':'response.completed','response':complete}),'']))
    calls=[]
    class Stream:
        def __enter__(self):return response
        def __exit__(self,*args):return False
    transport=SimpleNamespace(stream=lambda *a,**k:calls.append((a,k)) or Stream())
    client=VisualClient(provider,transport,'TEST_ONLY_DAITU_KEY',600)
    result=client.create(model=provider['model'],input=[],instructions='fixture',reasoning={'effort':'low'})
    assert calls[0][0]==('POST','https://api.daitu.cc/v1/responses')
    assert calls[0][1]['json']['stream'] is True and calls[0][1]['json']['store'] is False
    assert result.id=='response_fixture' and result.output_text=='{"candidates":[]}'
    assert result.usage.model_dump()['total_tokens']==3
