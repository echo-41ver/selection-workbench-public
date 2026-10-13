from pathlib import Path
import sys
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import web_console as web
from selection_automation.review_evidence import visual_evidence_text,review_explanation


def test_chinese_evidence_and_specific_review_preserve_facts():
    category={'stage3':{'identity_policy':{'features':[{'id':'layout','label':'部件排列','required_visual':True},{'id':'rear','label':'背部接口','required_visual':False}]}}}
    facts={'feature_comparisons':{'layout':{'relation':'UNKNOWN','amazon':'折叠侧边','temu':'侧边未展示'},'rear':{'relation':'UNKNOWN','amazon':'','temu':''}},
           'identity':{'reason':'主体清楚但边缘遮挡','differences':[{'field':'layout','kind':'UNRESOLVED','amazon':'折叠侧边','temu':'侧边未展示'}]}}
    text=visual_evidence_text(category,facts)
    assert '部件排列（未看清/未取得）' in text and 'UNKNOWN' not in text
    reason=review_explanation(category,facts,{'state':'待复核','reason':'尚未证实'})
    assert '需要核对：部件排列' in reason and '背部接口' not in reason
    assert facts['feature_comparisons']['layout']['relation']=='UNKNOWN'
    assert review_explanation(category,facts,{'state':'确认同款','reason':'具体双边设计一致'})=='具体双边设计一致'


def test_review_reads_only_published_snapshot_without_triggering_analysis(tmp_path,monkeypatch):
    path=tmp_path/'published.xlsx';book=Workbook();sheet=book.active;sheet.title='全量综合结果'
    sheet.append([]);sheet.append([]);sheet.append([])
    sheet.append(['Amazon ASIN','TEMU商品ID','最终判断','依据','视觉证据'])
    sheet.append(['A1','T1','待复核','需核对部件排列','Amazon：三按钮；TEMU：侧面未展示'])
    sheet.append(['A2','T2','确认同款','设计一致','清楚双边设计'])
    book.save(path);book.close();before=path.read_bytes()
    console=web.Console(tmp_path);calls=[]
    monkeypatch.setattr(console,'excel_workbook',lambda kind:(calls.append(kind) or path))
    monkeypatch.setattr(console,'invoke',lambda *a,**k: (_ for _ in ()).throw(AssertionError('no process launch')))
    result=console.visual_review()
    assert result['source']=='最近已发布综合 Excel' and result['sample_only']
    assert result['scanned_rows']==2 and len(result['examples'])==1
    assert result['examples'][0]['reason']=='需核对部件排列'
    assert result['filename']=='published.xlsx' and str(tmp_path) not in str(result)
    assert path.read_bytes()==before and calls==['layer123']


def test_no_published_review_is_not_reported_as_different(tmp_path,monkeypatch):
    console=web.Console(tmp_path);monkeypatch.setattr(console,'excel_workbook',lambda kind:None)
    result=console.visual_review()
    assert not result['available'] and '不代表不同款' in result['message']


def test_collection_history_composite_export_tolerates_no_commercial_report(tmp_path,monkeypatch):
    import json
    import export_category_workbooks as exporter
    root=Path(__file__).resolve().parents[1]
    cat=json.loads((root/'config/categories/power_drill.json').read_text('utf-8'))
    cat['execution_mode']='collection';base=tmp_path/'runs'/cat['category_key']
    frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True)
    category=base/'category.json';category.write_text(json.dumps(cat),encoding='utf-8')
    key=cat['category_key'];(frozen/f'layer12_summary_{key}.json').write_text(json.dumps({'amazon_products':1,'temu_products':1,'actual_pairs':1,'states':{'CANDIDATE':1}}))
    (frozen/f'layer12_rows_{key}.jsonl').write_text(json.dumps({'asin':'A','goods_id':'T','pair_status':'CANDIDATE','amazon_title':'drill','temu_title':'drill'})+'\n')
    monkeypatch.setattr(exporter,'ROOT',tmp_path)
    exporter.export(category,'layer123')
    from openpyxl import load_workbook
    book=load_workbook(base/'excel'/f'{cat["category_name_zh"]}_第一二三层.xlsx',read_only=True)
    sheet=book['全量综合结果'];assert sheet.max_row==5
    values=list(sheet.values);column=values[3].index('最终判断')
    assert '确认同款' not in str(values[4][column]);book.close()
