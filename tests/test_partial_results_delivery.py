import json
from pathlib import Path
import sys
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import export_category_workbooks as exporter
import web_console as web

@pytest.mark.parametrize('tail',[b'{"asin":"unfinished',b'{"title":"\xe4\xb8'])
def test_display_log_tail_keeps_complete_rows_and_source(tmp_path,tail):
    path=tmp_path/'results.jsonl';raw=b'{"asin":"A"}\n'+tail;path.write_bytes(raw);warnings=[]
    assert list(exporter.jsonl(path,allow_incomplete_tail=True,warnings=warnings))==[{'asin':'A'}]
    assert warnings==[{'file':'results.jsonl','line':2}]
    assert path.read_bytes()==raw
    with pytest.raises((ValueError,UnicodeDecodeError)):list(exporter.jsonl(path))

@pytest.mark.parametrize('raw',[b'{"asin":"A"}\nBROKEN\n',b'BROKEN\n{"asin":"A"}\n',b'[1]'])
def test_corrupt_records_are_not_silently_omitted(tmp_path,raw):
    path=tmp_path/'results.jsonl';path.write_bytes(raw)
    with pytest.raises(ValueError):list(exporter.jsonl(path,allow_incomplete_tail=True,warnings=[]))

@pytest.mark.parametrize('mode',['collection','full'])
def test_full_export_of_partial_logs_keeps_unknown_and_warning(tmp_path,monkeypatch,mode):
    root=Path(__file__).resolve().parents[1];cat=json.loads((root/'config/categories/power_drill.json').read_text('utf-8'));cat['execution_mode']=mode
    key=cat['category_key'];base=tmp_path/'runs'/key;frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True)
    category=base/'category.json';category.write_text(json.dumps(cat),encoding='utf-8')
    (frozen/f'layer12_summary_{key}.json').write_text(json.dumps({'amazon_products':1,'temu_products':2,'actual_pairs':2,'states':{'CANDIDATE':2}}))
    pairs=[{'asin':'A','goods_id':t,'pair_status':'CANDIDATE'} for t in ['T1','T2']]
    for name in ['layer12_rows','layer3_candidates']:
        (frozen/f'{name}_{key}.jsonl').write_text(''.join(json.dumps(r)+'\n' for r in pairs))
    log=base/'stage123/results.jsonl';raw=(json.dumps({'asin':'A','goods_id':'T1','vision_revision_key':'k','api_status':'failed','failure_reason':'TIMEOUT','attempt':1})+'\n'+ '{"asin":"A","goods_id":"T2"').encode();log.write_bytes(raw)
    ledger=base/'stage123/requests.jsonl';ledger.write_bytes(b'{"request_id":"r1","status":"received","usage":{"total_tokens":8}}\n{"status":')
    (base/'stage123/request_intents.jsonl').write_text(json.dumps({'request_id':'r1','status':'started','pairs':1})+'\n')
    monkeypatch.setattr(exporter,'ROOT',tmp_path);monkeypatch.setattr(exporter,'compatible_vision_keys',lambda *a,**kw:['k'])
    exporter.export(category,'layer123')
    book=load_workbook(base/'excel'/f'{cat["category_name_zh"]}_第一二三层.xlsx',read_only=True)
    try:
        sheet=book['全量综合结果'];data=list(sheet.values);assert sheet.max_row==6
        state=data[3].index('最终判断');assert 'API失败' in data[4][state]
        assert '确认同款' not in str(data[5][state]) and '不同款' not in str(data[5][state])
        audit=dict((row[0],row[1]) for row in book['汇总与规则'].iter_rows(min_row=2,values_only=True))
        assert '不完整' in audit['日志完整性提示'] or '未完成' in audit['日志完整性提示']
        assert len(json.loads(audit['未完成日志末行']))==2
        assert audit['视觉API请求数']==1 and audit['未结算视觉请求数']==1
        assert '费用未知' in audit['未结算请求提示']
        assert '被拒绝' in audit['未结算请求提示'] and '不代表零费用' in audit['未结算请求提示']
    finally:book.close()
    assert log.read_bytes()==raw

@pytest.mark.parametrize('publish_code',[0,1])
def test_failed_consumer_still_publishes_partial_results_without_key(tmp_path,monkeypatch,publish_code):
    console=web.Console(tmp_path);console.stage3_busy=True;calls=[]
    def invoke(argv,**kw):
        calls.append((argv,dict(kw['env'])))
        return SimpleNamespace(returncode=1 if len(calls)==1 else publish_code)
    monkeypatch.setattr(console,'invoke',invoke)
    console._run_stage3('fixture',tmp_path/'category.json',tmp_path/'candidates.jsonl','fixture-only-key')
    assert len(calls)==2 and 'publish_all.py' in calls[1][0][0]
    assert calls[0][1]['DEEPSEEK_API_KEY']=='fixture-only-key' and 'DEEPSEEK_API_KEY' not in calls[1][1]
    assert not console.stage3_busy
    if publish_code == 0: assert '被拒绝不代表零费用' in console.stage3_phase and '费用待核对' in console.stage3_phase
    assert ('部分结果已更新' in console.stage3_phase) if publish_code==0 else ('Excel 更新失败' in console.stage3_phase)

def test_failed_rebuild_preserves_previously_published_workbook(tmp_path):
    from selection_workflow.excel_delivery import data_workbook_path
    view=tmp_path/'published.xlsx'
    headers=['ASIN','title','link','image','goods_id','title','link','image','fact','vector','status']
    record=['A','drill',None,None,'T','drill',None,None,'visible',80,'CANDIDATE']
    exporter.export_stage_book(view,headers,[record],{'商品对总数':1},'fixture')
    prior_view=view.read_bytes();data=data_workbook_path(view);prior_data=data.read_bytes()
    def broken():
        yield record
        raise ValueError('Corrupt middle record')
    with pytest.raises(ValueError,match='Corrupt middle'):
        exporter.export_stage_book(view,headers,broken(),{'商品对总数':2},'fixture')
    assert view.read_bytes()==prior_view and data.read_bytes()==prior_data
    assert not list(tmp_path.glob('*.tmp.xlsx'))
