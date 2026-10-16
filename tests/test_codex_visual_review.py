from copy import deepcopy
import json
from pathlib import Path
import sys

import pytest

from selection_automation.codex_visual_review import checked_facts, import_reviews, load_reviews, normalize_review, supersedes_api
from test_set_inventory_identity import configured, facts


def fixture(status='DIFFERENT'):
    pair = {'asin': 'A', 'goods_id': 'T', 'vision_revision_key': 'revision',
            'amazon_image_path': '/amazon.jpg', 'temu_image_path': '/temu.jpg'}
    row = {**pair, 'review_status': status, 'amazon_evidence': '前排六把牛排刀',
           'temu_evidence': '前排八把牛排刀', 'reason': '刀型组成的牛排刀数量六把与八把不同',
           'reviewer': 'codex-01', 'model': 'Codex看图复核'}
    return pair, row


def test_only_manifest_can_authorize_complete_api_replacement():
    cat = configured()
    pair, row = fixture()
    row.update(visual_facts=facts(cat, a=6, t=8), supersedes_api=True)
    assert not supersedes_api(cat, normalize_review(cat, row, pair, 'unapproved'))
    pair['supersedes_api'] = True
    normalized = normalize_review(cat, row, pair, 'approved')
    assert supersedes_api(cat, normalized)
    assert normalized['review_status'] == 'DIFFERENT'
    row.pop('visual_facts')
    with pytest.raises(ValueError, match='complete original-image facts'):
        normalize_review(cat, row, pair, 'incomplete')


def test_missing_summaries_are_recovered_only_from_complete_visual_facts():
    cat = configured()
    pair, row = fixture('UNKNOWN')
    row.update(visual_facts=facts(cat,a=6,t=8),amazon_evidence='',temu_evidence='',reason='')
    normalized = normalize_review(cat,row,pair,'summaries')
    for side in ('amazon','temu'):
        module_evidence = [item[f'{side}_evidence'] for item in row['visual_facts']['modules'].values()]
        assert all(part in module_evidence for part in normalized[f'{side}_evidence'].split('；'))
        assert normalized[f'{side}_evidence']
    assert normalized['reason']
    assert row['amazon_evidence']==row['temu_evidence']==row['reason']==''
    row['visual_facts']['modules'].pop(next(iter(cat['stage3']['modules'])))
    with pytest.raises(ValueError,match='missing Codex review amazon_evidence'):
        normalize_review(cat,row,pair,'incomplete')


def test_new_import_beats_legacy_alphabetical_name_and_remains_idempotent(tmp_path):
    import os
    folder=tmp_path/'stage123/codex_reviews';folder.mkdir(parents=True)
    pair,row=fixture('DIFFERENT')
    legacy={**row,'review_source':'CODEX_VISUAL_REVIEW','review_status':'UNKNOWN'}
    path=folder/'zz-legacy.jsonl';path.write_text(json.dumps(legacy)+'\n',encoding='utf-8')
    os.utime(path,(1,1))
    destination=folder/'20261012-new.jsonl'
    import_reviews(configured(),{'pairs':[pair]},[row],destination,'new-review')
    saved=load_reviews(tmp_path)[('A','T')]
    assert saved['review_status']=='DIFFERENT' and saved['classified_at']>1
    before=destination.stat().st_mtime_ns
    assert not import_reviews(configured(),{'pairs':[pair]},[row],destination,'new-review')['changed']
    assert destination.stat().st_mtime_ns==before


def test_different_keeps_separate_source_no_api_metadata(tmp_path):
    pair, row = fixture()
    row.update(api_status='success', usage={'total_tokens': 100}, response_id='fake')
    destination = tmp_path/'stage123/codex_reviews/01.jsonl'
    result = import_reviews(configured(), {'pairs': [pair]}, [row], destination, 'codex-01')
    saved = load_reviews(tmp_path, {('A', 'T'): {'revision'}})[('A', 'T')]
    assert saved['review_status'] == 'DIFFERENT'
    assert not {'api_status', 'usage', 'response_id', 'billing_state'} & set(saved)
    assert result['api_requests_sent'] == 0
    assert 'visual_facts' not in saved
    assert load_reviews(tmp_path, {('A', 'T'): {'other_revision'}}) == {}


def test_review_import_retries_transient_windows_reader_locks(tmp_path,monkeypatch):
    from selection_automation import merge_temu_snapshots as atomic
    original=atomic.os.replace;attempts=[]
    def replace(source,target):
        attempts.append(target)
        if len(attempts)<3:raise PermissionError('TEST_ONLY_WINDOWS_READER_LOCK')
        return original(source,target)
    monkeypatch.setattr(atomic.os,'replace',replace)
    monkeypatch.setattr(atomic.time,'sleep',lambda seconds:None)
    pair,row=fixture();destination=tmp_path/'reviews.jsonl'
    result=import_reviews(configured(),[pair],[row],destination,'lock-retry')
    assert result['imported']==1 and len(attempts)==3
    assert json.loads(destination.read_text('utf-8'))['review_status']=='DIFFERENT'


def test_binary_completion_excludes_legacy_text_without_complete_visual_facts(tmp_path):
    category = configured()
    pair, row = fixture()
    destination = tmp_path/'stage123/codex_reviews/01.jsonl'
    import_reviews(category, {'pairs':[pair]}, [row], destination, 'legacy')
    assert len(load_reviews(tmp_path)) == 1
    category['visual_decision'] = 'visible_core_binary_v1'
    assert load_reviews(tmp_path, category=category) == {}
    row['visual_facts'] = facts(category)
    import_reviews(category, {'pairs':[pair]}, [row], destination, 'new')
    assert len(load_reviews(tmp_path, category=category)) == 1


def test_same_requires_complete_current_facts_and_confirmed_score():
    pair, row = fixture('SAME')
    assert normalize_review(configured(), row, pair, '01')['review_status'] == 'UNKNOWN'
    row['visual_facts'] = facts(configured())
    assert normalize_review(configured(), row, pair, '01')['review_status'] == 'SAME'
    row['visual_facts']['feature_comparisons']['body_type']['inventory']['temu']['groups'][0]['count'] = 8
    assert normalize_review(configured(), row, pair, '01')['review_status'] == 'UNKNOWN'
    del row['visual_facts']['modules']['body']
    assert normalize_review(configured(), row, pair, '01')['review_status'] == 'UNKNOWN'


def test_complete_facts_are_scored_and_retained_for_different_and_unknown():
    cat = configured()
    for status in ('DIFFERENT', 'UNKNOWN'):
        pair, row = fixture(status)
        row['visual_facts'] = facts(cat)
        if status == 'DIFFERENT':
            row['visual_facts']['identity']['verdict'] = 'DIFFERENT'
            row['visual_facts']['identity']['differences'] = [
                {'field': 'body_type', 'amazon': '直柄连接', 'temu': '弯柄连接',
                 'kind': 'IDENTITY', 'variant_field': 'NONE'}]
            row['visual_facts']['modules']['body']['relation'] = 'DIFFERENT'
            row['visual_facts']['modules']['body'].update(amazon_evidence='直柄连接', temu_evidence='弯柄连接')
        saved = normalize_review(cat, row, pair, '01')
        assert saved['review_status'] == status
        assert saved['module_result']['modules']['body']['score'] == (0 if status == 'DIFFERENT' else 80)
        assert saved['module_result']['total_score'] == (20 if status == 'DIFFERENT' else 100)


@pytest.mark.parametrize('status,expected', [('SAME','SAME'), ('DIFFERENT','UNKNOWN')])
def test_allowed_color_variant_uses_confirmed_identity_state(status, expected):
    cat = configured()
    pair, row = fixture(status)
    row['visual_facts'] = facts(cat)
    row['visual_facts']['modules']['appearance']['relation'] = 'VARIANT'
    row['visual_facts']['identity']['differences'] = [
        {'field':'表面颜色', 'amazon':'灰色外壳', 'temu':'红色外壳',
         'kind':'VARIANT', 'variant_field':'finish'}]
    saved = normalize_review(cat, row, pair, '01')
    assert saved['module_result']['state'] == '同款（存在规格变体）'
    assert saved['review_status'] == expected


def test_color_variant_cannot_override_an_identity_conflict():
    cat = configured()
    pair, row = fixture('SAME')
    row['visual_facts'] = facts(cat, a=6, t=8)
    row['visual_facts']['modules']['appearance']['relation'] = 'VARIANT'
    row['visual_facts']['identity']['differences'] = [
        {'field':'表面颜色', 'amazon':'灰色外壳', 'temu':'红色外壳',
         'kind':'VARIANT', 'variant_field':'finish'}]
    assert normalize_review(cat, row, pair, '01')['review_status'] == 'UNKNOWN'


def test_pump_facts_switch_rejects_missing_module_and_missing_fact():
    cat = configured()
    cat['stage3']['codex_require_module_facts'] = True
    pair, row = fixture('UNKNOWN')
    with pytest.raises(ValueError, match='require valid complete module facts'):
        normalize_review(cat, row, pair, '01')
    row['visual_facts'] = facts(cat)
    del row['visual_facts']['modules']['body']
    with pytest.raises(ValueError, match='require valid complete module facts'):
        normalize_review(cat, row, pair, '01')


def test_pump_facts_reject_reused_evidence_across_three_known_modules():
    cat = configured()
    cat['stage3']['codex_require_module_facts'] = True
    names = ['pump_body', 'control_layout', 'tube_connections', 'outlet_nozzle', 'power_compartment', 'visual_appearance']
    cat['stage3']['modules'] = dict(zip(names, [30, 15, 20, 20, 10, 5]))
    evidence = ('同一整机描述', '同一候选整机描述')
    module_facts = {name: {'relation': 'SAME', 'amazon_evidence': evidence[0],
                           'temu_evidence': evidence[1], 'hard_conflict_code': 'NONE'}
                    for name in names[:3]}
    module_facts.update({name: {'relation': 'UNKNOWN', 'amazon_evidence': '',
                                'temu_evidence': '', 'hard_conflict_code': 'NONE'}
                         for name in names[3:]})
    value = {'overall_confidence': 0.8, 'identity': {}, 'modules': module_facts,
             'feature_comparisons': {}}
    with pytest.raises(ValueError, match='reused evidence across visual modules'):
        checked_facts(cat, value, 'T')


@pytest.mark.parametrize('marker', ['\ufffd', '???'])
def test_pump_facts_reject_corrupted_nonunknown_evidence(marker):
    cat = configured()
    cat['stage3']['codex_require_module_facts'] = True
    pair, row = fixture('UNKNOWN')
    row['visual_facts'] = facts(cat)
    row['visual_facts']['modules']['body']['amazon_evidence'] = f'泵体轮廓{marker}'
    with pytest.raises(ValueError, match='require valid complete module facts'):
        normalize_review(cat, row, pair, '01')


def test_pump_facts_accept_utf8_chinese_evidence():
    cat = configured()
    cat['stage3']['codex_require_module_facts'] = True
    pair, row = fixture('UNKNOWN')
    row['visual_facts'] = facts(cat)
    row['visual_facts']['modules']['body']['amazon_evidence'] = '红色泵体与黑色接头'
    row['visual_facts']['modules']['body']['temu_evidence'] = '红色泵体与黑色接头'
    assert normalize_review(cat, row, pair, '01')['review_status'] == 'UNKNOWN'


def test_complete_decoded_facts_can_be_imported_without_changing_scores(tmp_path):
    cat = configured()
    cat['stage3']['codex_require_module_facts'] = True
    pair, row = fixture('UNKNOWN')
    row['visual_facts'] = facts(cat)
    normalized = normalize_review(cat, row, pair, '01')
    destination = tmp_path/'stage123/codex_reviews/01.jsonl'
    result = import_reviews(cat, {'pairs': [pair]}, [normalized], destination, '01')
    saved = load_reviews(tmp_path)[('A', 'T')]
    assert result['imported'] == 1
    assert saved['module_result'] == normalized['module_result']
    assert saved['visual_facts'] == normalized['visual_facts']


def test_unknown_module_has_null_score_when_complete_facts_are_present():
    cat = configured()
    pair, row = fixture('UNKNOWN')
    row['visual_facts'] = facts(cat)
    row['visual_facts']['modules']['body'].update(relation='UNKNOWN', amazon_evidence='', temu_evidence='')
    saved = normalize_review(cat, row, pair, '01')
    assert saved['module_result']['modules']['body']['score'] is None
    assert saved['module_result']['total_score'] is None


def test_generic_difference_does_not_become_exclusion():
    pair, row = fixture()
    row['amazon_evidence'] = '与基准一致'
    assert normalize_review(configured(), row, pair, '01')['review_status'] == 'UNKNOWN'


@pytest.mark.parametrize('as_list',[False,True])
def test_multiple_identity_features_are_accepted_when_each_is_identity(as_list):
    cat=configured()
    cat['stage3']['identity_policy']['features'].append({'id':'handle_construction','role':'identity'})
    pair,row=fixture()
    row['identity_feature']=['body_type','handle_construction'] if as_list else 'body_type,handle_construction'
    saved=normalize_review(cat,row,pair,'01')
    assert saved['review_status']=='DIFFERENT'
    assert saved['identity_feature']==['body_type','handle_construction']


def test_multiple_features_including_variant_cannot_be_identity_exclusion():
    pair,row=fixture()
    row['identity_feature']='body_type,finish'
    assert normalize_review(configured(),row,pair,'01')['review_status']=='UNKNOWN'


def knife_category():
    cat=configured()
    cat['stage3']['identity_policy']['features'] += [
        {'id':fid,'role':'identity','label':label} for fid,label in (
            ('knife_set_composition','套装身份组成'),('blade_shape_profile','刀刃轮廓'),
            ('handle_construction','柄部连接形式'),('bolster_ferrule_geometry','颈部几何'),
            ('storage_block_structure','收纳架结构'))]
    return cat


@pytest.mark.parametrize('raw', ['刀柄/刀座','刀柄，刀架','刀柄结构,刀座',['刀柄','刀座'],
                                '柄部连接形式/收纳架结构','handle_construction,storage_block_structure'])
def test_chinese_knife_aliases_and_exact_config_labels_bind_existing_ids(raw):
    pair,row=fixture()
    row['identity_feature']=raw
    saved=normalize_review(knife_category(),row,pair,'01')
    assert saved['review_status']=='DIFFERENT'
    assert saved['identity_feature']==['handle_construction','storage_block_structure']


@pytest.mark.parametrize('raw', ['刀型构成/刀刃/护指','套装构成，刀型，刀颈','构成,刀刃,护指'])
def test_remaining_bounded_knife_aliases_resolve(raw):
    pair,row=fixture()
    row['identity_feature']=raw
    saved=normalize_review(knife_category(),row,pair,'01')
    assert saved['review_status']=='DIFFERENT'
    assert saved['identity_feature']==['knife_set_composition','blade_shape_profile','bolster_ferrule_geometry']


@pytest.mark.parametrize('raw',['刀柄/颜色','刀座/材质','刀柄/尺寸','刀柄/未知部位','刀柄/finish'])
def test_unknown_or_nonidentity_chinese_feature_does_not_bind_identity(raw):
    pair,row=fixture()
    row['identity_feature']=raw
    assert normalize_review(knife_category(),row,pair,'01')['review_status']=='UNKNOWN'


def test_knife_alias_cannot_create_feature_in_unconfigured_category():
    pair,row=fixture()
    row['identity_feature']='刀柄/刀座'
    assert normalize_review(configured(),row,pair,'01')['review_status']=='UNKNOWN'


def test_import_is_idempotent_and_revision_mismatch_preserves_existing(tmp_path):
    pair, row = fixture()
    destination = tmp_path/'stage123/codex_reviews/01.jsonl'
    for _ in range(2):
        result = import_reviews(configured(), {'pairs': [pair]}, [row], destination, '01')
    assert result['changed'] is False
    before = destination.read_bytes()
    assert len(before.splitlines()) == 1
    row['vision_revision_key'] = 'stale'
    with pytest.raises(ValueError, match='does not match manifest'):
        import_reviews(configured(), {'pairs': [pair]}, [row], destination, '01')
    assert destination.read_bytes() == before


def test_publisher_source_signature_includes_independent_reviews(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
    from watch_stage3_excel import source_paths
    path = tmp_path/'stage123/codex_reviews/01.jsonl'
    path.parent.mkdir(parents=True)
    path.write_text('{}\n', encoding='utf-8')
    assert path in source_paths(tmp_path/'category.json', tmp_path)


@pytest.mark.parametrize('with_api,replace_api', [(False,False),(True,False),(True,True)])
def test_simple_excel_uses_api_success_unless_complete_recheck_authorized(tmp_path, monkeypatch, with_api, replace_api):
    from openpyxl import load_workbook
    from selection_automation import vision
    from selection_automation.commercial import current_report
    import export_category_workbooks as exporter
    cat=configured();cat['execution_mode']='full';key=cat['category_key']
    base=tmp_path/'runs'/key;frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True)
    (base/'amazon').mkdir();(base/'temu/union').mkdir(parents=True)
    (base/'amazon/amazon_frontend_latest.json').write_text(json.dumps({'records':[{'asin':'A','title':'TEST_ONLY_AMAZON'}]}),'utf-8')
    (base/'temu/union/products.jsonl').write_text(json.dumps({'goods_id':'T','title':'TEST_ONLY_TEMU'})+'\n','utf-8')
    pair,review=fixture()
    if replace_api:
        pair['supersedes_api']=True;review['visual_facts']=facts(cat,a=6,t=8)
    row={**pair,'pair_status':'CANDIDATE','vector_score':100}
    for name in ('layer3_candidates','layer12_rows'):
        (frozen/f'{name}_{key}.jsonl').write_text(json.dumps(row)+'\n','utf-8')
    import_reviews(cat,{'pairs':[pair]},[review],base/'stage123/codex_reviews/01.jsonl','01')
    if with_api:
        api={**pair,'api_status':'success','visual_facts':facts(cat),'model':'existing-api'}
        (base/'stage123/results.jsonl').write_text(json.dumps(api)+'\n','utf-8')
    monkeypatch.setattr(exporter,'ROOT',tmp_path)
    monkeypatch.setattr(exporter,'load_category',lambda _:cat)
    monkeypatch.setattr(vision,'compatible_vision_keys',lambda *a,**kw:('revision',))
    report=current_report(tmp_path,cat)
    assert report['products'][0]['match']['confirmed']==int(with_api and not replace_api)
    output=exporter.export(tmp_path/'category.json','layer123')
    book=load_workbook(output['固定查看副本']);sheet=book['选品结论']
    headers=[c.value for c in sheet[4]]
    assert sheet.cell(5,headers.index('Amazon ASIN')+1).value=='A'
    assert sheet.cell(5,headers.index('同款数')+1).value==int(with_api and not replace_api)
    assert sheet.cell(5,headers.index('不同款数')+1).value==int(not with_api or replace_api)
    if with_api and not replace_api:
        detail=book['同款商品明细'];cols=[c.value for c in detail[4]]
        assert detail.cell(5,cols.index('分析模型')+1).value=='existing-api'
    book.close()


def test_collection_only_task_does_not_export_a_fabricated_analysis(tmp_path,monkeypatch):
    import export_category_workbooks as exporter
    cat=configured();cat['execution_mode']='collection'
    monkeypatch.setattr(exporter,'ROOT',tmp_path)
    monkeypatch.setattr(exporter,'load_category',lambda _:cat)
    with pytest.raises(ValueError,match='尚未接入同款分析'):
        exporter.export(tmp_path/'category.json','layer123')


def test_progress_reports_codex_separately_and_ignores_stale_or_api_success(tmp_path, monkeypatch):
    from selection_automation import stage3_progress as progress
    pair, review = fixture()
    cat = configured()
    candidates = []
    for index, status in enumerate(('DIFFERENT','UNKNOWN','DIFFERENT')):
        bound = {**pair, 'goods_id':str(index), 'vision_revision_key':f'key{index}', 'vector_score':100}
        candidates.append(bound)
        raw = {**review, **bound, 'review_status':status}
        normalized = normalize_review(cat, raw, bound, '01')
        if index==2:
            normalized['vision_revision_key']='stale'
        destination = tmp_path/f'stage123/codex_reviews/{index}.jsonl'
        destination.parent.mkdir(parents=True,exist_ok=True)
        destination.write_text(json.dumps(normalized)+'\n',encoding='utf-8')
    (tmp_path/'stage123/results.jsonl').write_text(json.dumps({**candidates[0], 'api_status':'success'})+'\n',encoding='utf-8')
    monkeypatch.setattr(progress,'candidates',lambda _:candidates)
    monkeypatch.setattr(progress,'compatible_vision_keys',lambda _, row, **__:(row['vision_revision_key'],))
    result=progress.snapshot(tmp_path,cat,history=False)
    assert result['completed']==2 and result['api_completed']==1
    assert result['codex_completed']==1 and result['codex_unknown']==1
    assert result['codex_different']==0 and result['pending_after_codex']==1
    assert result['pending']==result['runnable_pending']==1
    assert result['processed']==2 and result['percent']==66.7


@pytest.mark.parametrize('status', ['DIFFERENT','UNKNOWN'])
def test_current_codex_review_prevents_paid_runner_repeating_pair(tmp_path,monkeypatch,status):
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
    import openai
    import run_category_layer3 as runner
    from types import SimpleNamespace
    cat=configured();pair,review=fixture(status)
    stage=tmp_path/'stage123';stage.mkdir()
    current={**pair,'vector_score':100}
    stale={**current,'goods_id':'stale','vision_revision_key':'new-key'}
    import_reviews(cat,{'pairs':[pair]},[review],stage/'codex_reviews/current.jsonl','current')
    old_pair={**stale,'vision_revision_key':'old-key'}
    import_reviews(cat,{'pairs':[old_pair]},[{**review,**old_pair}],stage/'codex_reviews/stale.jsonl','stale')
    path=tmp_path/'candidates.jsonl'
    path.write_text(''.join(json.dumps(row)+'\n' for row in (current,stale)),encoding='utf-8')
    monkeypatch.setattr(runner,'load_category',lambda _:cat)
    monkeypatch.setattr(runner,'compatible_vision_keys',lambda _,row, **__:(row['vision_revision_key'],))
    def no_paid_client(**kwargs):raise AssertionError('paid client must never be constructed')
    monkeypatch.setattr(openai,'OpenAI',no_paid_client)
    observed=[]
    monkeypatch.setattr(runner,'consume',lambda *args:observed.extend(args[-1]))
    args=SimpleNamespace(category=tmp_path/'category.json',results=stage/'results.jsonl',validate_only=True,max_pairs=0)
    runner.run(args,path)
    assert [row['goods_id'] for row in observed]==['stale']
    assert not args.results.exists()


def test_codex_completion_keeps_uncertain_billing_visible_but_no_runnable_work(tmp_path,monkeypatch):
    from selection_automation import stage3_progress as progress
    cat=configured();pair,review=fixture()
    import_reviews(cat,{'pairs':[pair]},[review],tmp_path/'stage123/codex_reviews/one.jsonl','one')
    monkeypatch.setattr(progress,'candidates',lambda _:[{**pair,'vector_score':100}])
    monkeypatch.setattr(progress,'compatible_vision_keys',lambda *_, **__:('revision',))
    (tmp_path/'stage123/request_intents.jsonl').write_text(json.dumps({'request_id':'unknown','status':'started',
        'pairs':1,'vision_revision_keys':['revision']})+'\n',encoding='utf-8')
    (tmp_path/'stage123/results.jsonl').write_text(json.dumps({**pair,'api_status':'failed',
        'billing_state':'UNKNOWN','attempt':1})+'\n',encoding='utf-8')
    before=(tmp_path/'stage123/request_intents.jsonl').read_bytes()
    result=progress.snapshot(tmp_path,cat,history=False)
    assert (result['completed'],result['pending'],result['runnable_pending'],result['pending_requests'])==(1,0,0,0)
    assert result['billing_unknown_pairs']==1 and result['billing_unknown_requests']==1
    assert result['percent']==100 and result['api_completed']==0
    assert (tmp_path/'stage123/request_intents.jsonl').read_bytes()==before


def test_stale_later_review_cannot_hide_valid_revision_and_trigger_repayment(tmp_path,monkeypatch):
    from selection_automation import stage3_progress as progress
    import run_category_layer3 as runner
    from types import SimpleNamespace
    cat=configured();pair,review=fixture()
    stage=tmp_path/'stage123'
    import_reviews(cat,{'pairs':[pair]},[review],stage/'codex_reviews/a-current.jsonl','current')
    old_pair={**pair,'vision_revision_key':'stale'}
    import_reviews(cat,{'pairs':[old_pair]},[{**review,**old_pair}],stage/'codex_reviews/z-stale.jsonl','stale')
    candidate={**pair,'vector_score':100}
    monkeypatch.setattr(progress,'candidates',lambda _:[candidate])
    monkeypatch.setattr(progress,'compatible_vision_keys',lambda *_, **__:('revision',))
    result=progress.snapshot(tmp_path,cat,history=False)
    assert result['completed']==1 and result['pending_requests']==result['runnable_pending']==0
    path=tmp_path/'candidates.jsonl';path.write_text(json.dumps(candidate)+'\n',encoding='utf-8')
    monkeypatch.setattr(runner,'load_category',lambda _:cat)
    monkeypatch.setattr(runner,'compatible_vision_keys',lambda *_, **__:('revision',))
    observed=[];monkeypatch.setattr(runner,'consume',lambda *args:observed.extend(args[-1]))
    args=SimpleNamespace(category=tmp_path/'category.json',results=stage/'results.jsonl',validate_only=True,max_pairs=0)
    runner.run(args,path)
    assert observed==[]


@pytest.mark.parametrize('status,bucket',[('DIFFERENT','different'),('UNKNOWN','possible')])
def test_business_match_counts_current_codex_reviews_with_same_api_priority(tmp_path,monkeypatch,status,bucket):
    from selection_automation import commercial
    cat=configured();cat.update(temu={'base_queries':['knives']},commercial={})
    key=cat['category_key'];base=tmp_path/'runs'/key
    (base/'amazon').mkdir(parents=True);(base/'temu/union').mkdir(parents=True)
    (base/'amazon/amazon_frontend_latest.json').write_text(json.dumps({'records':[{'asin':'A'}]}),encoding='utf-8')
    (base/'temu/union/products.jsonl').write_text(json.dumps({'goods_id':'T'})+'\n',encoding='utf-8')
    pair,review=fixture(status);row={**pair,'pair_status':'CANDIDATE'}
    frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True)
    for name in ('layer12_rows','layer3_candidates'):
        (frozen/f'{name}_{key}.jsonl').write_text(json.dumps(row)+'\n',encoding='utf-8')
    import_reviews(cat,{'pairs':[pair]},[review],base/'stage123/codex_reviews/one.jsonl','one')
    from selection_automation import vision
    monkeypatch.setattr(vision,'compatible_vision_keys',lambda *_, **__:('revision',))
    report=commercial.build_report(tmp_path,cat)
    match=report['products'][0]['match']
    assert match['pending']==0 and match['confirmed']==0
    assert match['possible']==(1 if bucket=='possible' else 0)
    assert match['state']==('UNKNOWN' if bucket=='possible' else 'SAMPLE_PASS')
    # Existing valid API facts always take precedence over independent reviews.
    (base/'stage123/results.jsonl').write_text(json.dumps({**pair,'api_status':'success','visual_facts':facts(cat)})+'\n',encoding='utf-8')
    match=commercial.build_report(tmp_path,cat)['products'][0]['match']
    assert match['confirmed']==1 and match['possible']==0
