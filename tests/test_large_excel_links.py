import sys,pathlib,zipfile
from openpyxl import load_workbook
sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]/'scripts'))
from export_category_workbooks import export_stage_book

def test_real_workbook_preserves_four_links_above_excel_relationship_limit(tmp_path):
    count=17000
    def records():
        for n in range(count):
            yield [f'A{n}','Amazon',f'https://amazon.test/p/{n}',f'https://amazon.test/img/{n}',
                str(n),'Temu',f'https://temu.test/p/{n}',f'https://temu.test/img/{n}', 'facts',.7,'候选']
    output=tmp_path/'large.xlsx'
    export_stage_book(output,[],records(),{'商品对数量':count},'TEST_ONLY')
    book=load_workbook(output,read_only=True,data_only=False)
    sheet=book['第一二层全量核验'];assert sheet.max_row==count+4
    checks={0,16381,16382,16383,count-1};seen=set();formula_count=0
    for n,row in enumerate(sheet.iter_rows(min_row=5,values_only=True)):
        values={1:f'https://amazon.test/img/{n}',4:f'https://temu.test/img/{n}',
                10:f'https://amazon.test/p/{n}',11:f'https://temu.test/p/{n}'}
        for col,url in values.items():
            assert row[col].startswith('=HYPERLINK(') and f'"{url}"' in row[col]
            formula_count+=1
        if n in checks:seen.add(n)
    assert seen==checks and formula_count==68000
    book.close()
    cached=load_workbook(output,read_only=True,data_only=True)
    label=next(cached['第一二层全量核验'].iter_rows(min_row=5,max_row=5,values_only=True))[10]
    assert label and not label.startswith('=')
    cached.close()
    with zipfile.ZipFile(output) as z:
        assert z.testzip() is None
        assert b'<hyperlinks' not in z.read('xl/worksheets/sheet2.xml')

def test_long_and_quoted_urls_keep_exact_source(tmp_path):
    long='https://temu.test/'+('x'*300);quoted='https://amazon.test/a"b'
    row=['A','Amazon',quoted,quoted,'1','Temu',long,long,'facts',.7,'候选']
    output=tmp_path/'edge.xlsx';export_stage_book(output,[],[row],{'商品对数量':1},'TEST_ONLY')
    book=load_workbook(output,read_only=True,data_only=False)
    values=next(book['第一二层全量核验'].iter_rows(min_row=5,max_row=5,values_only=True))
    assert values[4]==long and values[11]==long
    assert 'a""b' in values[1] and 'a""b' in values[10]
    book.close()
