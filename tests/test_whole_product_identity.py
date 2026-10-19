"""Regression of decisions and exports, using synthetic facts, never a real API."""
from category_fixtures import category_path
import json
from pathlib import Path
import sys

import pytest
from identity_fixtures import identity
from selection_automation.rules import score, validate
from selection_automation.vision import decode, protocol, vision_revision

ROOT = Path(__file__).resolve().parents[1]


def category(filename='bath_towel_set.json'):
    return json.loads(category_path(filename).read_text('utf-8'))


def facts(cat):
    return {'identity': identity(), 'overall_confidence': .95, 'modules': {
        k: {'relation': 'SAME', 'amazon_evidence': '右侧双条纹边带', 'temu_evidence': '右侧双条纹边带',
            'hard_conflict_code': 'NONE'} for k in cat['stage3']['modules']}}


@pytest.mark.parametrize('filename,field,left,right', [
    ('bath_towel_set.json', '套装构成', '六条整幅浴巾', '两浴巾两手巾两方巾'),
    ('bath_towel_set.json', '装饰条带', '短边单条平带', '短边多条宽窄条纹'),
    ('shower_curtain_local_union_20260922.json', '本体图案', '蓝底细链条白纹及流苏', '白底粗灰波纹无流苏'),
    ('power_drill.json', '出售主体', '带电池整台电钻', '单独钻夹头配件'),
])
def test_explicit_identity_difference_overrides_even_perfect_module_scores(filename, field, left, right):
    cat = category(filename); value = facts(cat)
    value['identity']['verdict'] = 'DIFFERENT'
    value['identity']['differences'] = [{'field': field, 'amazon': left, 'temu': right, 'kind': 'IDENTITY'}]
    result = score(cat, value)
    assert result['total_score'] == 100
    assert result['state'] == '不同款'


def test_critical_variant_cannot_be_paid_for_by_other_modules():
    cat = category(); value = facts(cat)
    value['modules']['set_composition']['relation'] = 'VARIANT'
    assert score(cat, value)['total_score'] > 88
    assert score(cat, value)['state'] == '待复核'


def test_specific_allowed_color_variant_preserves_base_identity():
    cat = category(); value = facts(cat)
    value['identity']['differences'] = [{'field': 'visible_color', 'amazon': '粉色本体', 'temu': '灰色本体', 'kind': 'VARIANT'}]
    value['modules']['color_pattern_variant']['relation'] = 'VARIANT'
    assert score(cat, value)['state'] == '同款（存在规格变体）'


def test_explicit_towel_pattern_variant_can_explain_bound_key_module():
    cat = category(); value = facts(cat)
    value['identity']['differences'] = [{'field': '装饰带内纹样', 'amazon': '短边同位置装饰带内蓝色花卉',
        'temu': '短边同位置装饰带内米色花卉', 'kind': 'VARIANT', 'variant_field': 'decorative_pattern'}]
    value['modules']['surface_design_weave_border']['relation'] = 'VARIANT'
    assert score(cat, value)['state'] == '同款（存在规格变体）'
    value['identity']['design_comparison']['relation'] = 'DIFFERENT'
    value['identity']['design_comparison'].update(amazon='短边单条窄带', temu='满幅宽条纹')
    assert score(cat, value)['state'] == '不同款'


def test_pattern_variant_without_same_structure_does_not_confirm():
    cat = category(); value = facts(cat)
    value['modules']['surface_design_weave_border']['relation'] = 'VARIANT'
    assert score(cat, value)['state'] == '待复核'
    value['identity']['differences'] = [{'field': '颜色', 'amazon': '白色', 'temu': '紫色',
        'kind': 'VARIANT', 'variant_field': 'visible_color'}]
    assert score(cat, value)['state'] == '待复核'


def test_pattern_rule_is_not_enabled_for_shower_curtains():
    cat = category('shower_curtain_local_union_20260922.json')
    assert 'decorative_pattern' not in cat['stage3']['variant_fields']
    assert '保持主体织法' in protocol(category(), 1)[0]


@pytest.mark.parametrize('binding', [[], {'absent':['decorative_pattern']},
    {'surface_design_weave_border':['surface_design_weave_border']},
    {'surface_design_weave_border':['unconfigured']}, {'surface_design_weave_border':'decorative_pattern'}])
def test_invalid_module_variant_bindings_are_rejected(binding):
    cat = category(); cat['stage3']['module_variant_fields'] = binding
    with pytest.raises(ValueError, match='module variants'):validate(cat)


def test_whole_module_name_is_not_a_permitted_variant():
    cat = category('shower_curtain_local_union_20260922.json'); value = facts(cat)
    value['identity']['differences'] = [{'field': 'design_identity', 'amazon': '细链条印花', 'temu': '石头印花', 'kind': 'VARIANT'}]
    assert score(cat, value)['state'] == '待复核'


@pytest.mark.parametrize('mutation', ['old', 'circular', 'empty', 'missing_side', 'unknown', 'zero'])
def test_unproven_identity_cannot_confirm(mutation):
    cat = category(); value = facts(cat)
    if mutation == 'old': del value['identity']
    elif mutation == 'circular':
        for r in value['identity']['matches']: r.update(amazon='与基准一致', temu='都是毛巾')
    elif mutation == 'empty': value['identity']['matches'] = []
    elif mutation == 'missing_side':
        for r in value['identity']['matches']: r['temu'] = ''
    elif mutation == 'unknown': value['modules']['set_composition']['relation'] = 'UNKNOWN'
    else: value['overall_confidence'] = 0
    assert score(cat, value)['state'] == '待补充证据'


def test_simple_clear_product_does_not_require_an_arbitrary_number_of_details():
    cat = category(); value = facts(cat)
    value['identity']['matches'] = value['identity']['matches'][:1]
    assert score(cat, value)['state'] == '确认同款'


def test_detailed_evidence_can_refer_to_reference_without_becoming_circular():
    cat = category(); value = facts(cat)
    value['identity']['design_comparison']['temu'] = '右侧两条窄平织带，中间宽毛圈带，包边转角与基准一致'
    assert score(cat, value)['state'] == '确认同款'
    value['identity']['design_comparison']['temu'] = '与基准一致'
    assert score(cat, value)['state'] == '待复核'


def test_explicit_identity_modules_do_not_require_hidden_nonidentity_details():
    cat = category(); value = facts(cat)
    cat['stage3']['identity_required_modules'] = ['surface_design_weave_border', 'towel_construction_pile_edge', 'set_composition']
    cat['stage3']['confirmation_minimums'].pop('known_modules')
    value['modules']['color_pattern_variant'].update(relation='UNKNOWN', amazon_evidence='', temu_evidence='')
    assert score(cat, value)['state'] == '确认同款'
    value['modules']['surface_design_weave_border'].update(relation='UNKNOWN', amazon_evidence='', temu_evidence='')
    assert score(cat, value)['state'] == '待补充证据'


def test_chinese_variant_label_uses_canonical_id_without_rejecting_true_same():
    cat = category(); value = facts(cat)
    value['identity']['differences'] = [{'field': '品牌吊牌', 'variant_field': 'packaging',
        'amazon': '无吊牌', 'temu': '纸质品牌吊牌', 'kind': 'VARIANT'}]
    assert score(cat, value)['state'] == '同款（存在规格变体）'


def test_unknown_variant_label_is_not_evidence_of_different_identity():
    cat = category(); value = facts(cat)
    value['identity']['differences'] = [{'field': '品牌吊牌', 'amazon': '无吊牌', 'temu': '纸质品牌吊牌', 'kind': 'VARIANT'}]
    assert score(cat, value)['state'] == '待复核'


def test_incidental_occlusion_does_not_override_visible_identity_agreement():
    cat = category(); value = facts(cat)
    value['identity']['uncertainties'] = ['营销横幅遮住右侧一小块边缘，主体设计与关键结构已清晰对应']
    assert score(cat, value)['state'] == '确认同款'
    value['modules']['surface_design_weave_border']['relation'] = 'UNKNOWN'
    assert score(cat, value)['state'] == '待补充证据'


def test_visible_feature_absence_is_a_difference_not_missing_evidence():
    cat = category(); value = facts(cat)
    value['identity']['verdict'] = 'DIFFERENT'
    value['identity']['differences'] = [{'field': '绣花', 'amazon': '中央花朵绣花清晰可见',
        'temu': '中央完整展开且无绣花', 'kind': 'IDENTITY'}]
    assert score(cat, value)['state'] == '不同款'


def test_clear_module_difference_does_not_require_a_hard_conflict_code():
    cat = category(); value = facts(cat)
    value['identity']['verdict'] = 'DIFFERENT'
    value['modules']['surface_design_weave_border'].update(relation='DIFFERENT',
        amazon_evidence='素色无绣花', temu_evidence='中央大幅花朵绣花', hard_conflict_code='NONE')
    assert score(cat, value)['state'] == '不同款'


def test_metadata_minimums_are_valid_but_typos_are_not():
    cat = category(); validate(cat)
    cat['stage3']['confirmation_minimums']['known_modules'] = True
    with pytest.raises(ValueError): validate(cat)


def test_schema_and_local_decoder_require_holistic_evidence():
    cat = category(); value = facts(cat)
    raw = {'candidates': [{'i': 'T', 'v': value['identity'], 'c': .95, 'm': {
        str(i): {'r': 'S', 'a': '右侧双条纹边带', 't': '右侧双条纹边带', 'h': 'NONE'}
        for i in range(len(cat['stage3']['modules']))}}]}
    assert score(cat, decode(cat, raw, ['T'])['T'])['state'] == '确认同款'
    _, schema = protocol(cat, 1)
    assert 'v' in schema['properties']['candidates']['items']['required']
    raw['candidates'][0]['v']['matches'][0]['score'] = 100
    with pytest.raises(ValueError): decode(cat, raw, ['T'])


def test_prior_results_without_new_protocol_are_not_relabelled_as_current():
    sys.path.insert(0, str(ROOT / 'scripts'))
    import import_prior_shower_layer3 as importer
    cat = category(); value = facts(cat); del value['identity']
    assert importer.normalize_facts(cat, {'visual_facts': value}) is None


@pytest.mark.parametrize('partial_log',[False,True])
def test_shower_excel_cannot_promote_unknown_identity_on_high_vector(tmp_path, monkeypatch, partial_log):
    sys.path.insert(0, str(ROOT / 'scripts'))
    import export_shower_layer123_excel as exporter
    from openpyxl import load_workbook
    cat = category('shower_curtain_local_union_20260922.json')
    key = cat['category_key']; base = tmp_path / 'runs' / key
    frozen = base / 'stage123/export_frozen'; frozen.mkdir(parents=True)
    path = base / 'category.json'; path.write_text(json.dumps(cat), encoding='utf-8')
    image = tmp_path / 'image'; image.write_bytes(b'synthetic-image-identity')
    pair = {'asin': 'A', 'goods_id': 'T', 'pair_revision': 'p', 'pair_status': 'CANDIDATE', 'vector_score': 99.9,
            'amazon_image_path': str(image), 'temu_image_path': str(image)}
    for name in ('layer12_rows', 'layer3_candidates'):
        (frozen / f'{name}_{key}.jsonl').write_text(json.dumps(pair) + '\n', encoding='utf-8')
    (frozen / f'layer12_summary_{key}.json').write_text(json.dumps({'amazon_products': 1, 'temu_products': 1}), encoding='utf-8')
    value = facts(cat); value['modules']['core_structure']['relation'] = 'UNKNOWN'
    (base / 'stage123/results.jsonl').write_text(json.dumps({**pair, 'api_status': 'success',
        'vision_revision_key': vision_revision(cat, pair), 'visual_facts': value}) + '\n', encoding='utf-8')
    if partial_log:
        with (base / 'stage123/results.jsonl').open('ab') as stream: stream.write(b'{\"unfinished\":')
        (base / 'stage123/requests.jsonl').write_bytes(b'{\"status\":')
        (base / 'stage123/request_intents.jsonl').write_text(json.dumps({'request_id':'r1','status':'started','pairs':1})+'\n')
    monkeypatch.setattr(exporter, 'ROOT', tmp_path)
    delivered = exporter.export(path)
    with_path = Path(delivered['固定查看副本'])
    book = load_workbook(with_path, read_only=True)
    try:
        rows = list(book['全量综合结果'].values)
        assert rows[4][24] == score(cat, value)['state'] == '待补充证据'
        if partial_log:
            audit={row[0]:row[1] for row in book.worksheets[1].iter_rows(min_row=2,values_only=True) if len(row)>1}
            assert '日志完整性提示' in audit
            assert audit['未结算视觉请求数']==1 and audit['第三层API请求批次']==1
    finally: book.close()

@pytest.mark.parametrize('field',['design_comparison','matches','differences'])
def test_whole_identity_opposite_source_fields_require_review(field):
    cat=category();value=facts(cat)
    reversed_pair={'amazon':'候选商品可见圆形连接件及双道包边','temu':'Amazon基准图可见圆形连接件及双道包边'}
    if field=='design_comparison':value['identity'][field].update(reversed_pair,relation='SAME')
    elif field=='matches':value['identity'][field][0].update(reversed_pair)
    else:value['identity'][field]=[{'field':'visible_color','kind':'VARIANT',**reversed_pair}]
    before=json.dumps(value,ensure_ascii=False)
    result=score(cat,value)
    assert result['state']=='待复核'
    assert json.dumps(value,ensure_ascii=False)==before


@pytest.mark.parametrize('filename', ['bath_towel_set.json', 'power_drill.json', 'shower_curtain_local_union_20260922.json'])
def test_unsupported_variant_reports_config_boundary_without_confirming(filename):
    cat = category(filename); value = facts(cat)
    value['identity']['differences'] = [{'field': 'unconfigured visible change',
        'amazon': 'visible narrow stripe', 'temu': 'visible wide stripe',
        'kind': 'VARIANT', 'variant_field': 'unconfigured_variant'}]
    result = score(cat, value)
    assert result['state'] == '待复核'
    assert result['unsupported_variant_fields'] == ['unconfigured_variant']
    assert '当前品类规则未允许变体字段：unconfigured_variant' in result['reason']
    value['identity']['differences'][0]['kind'] = 'IDENTITY'
    value['identity']['differences'][0]['variant_field'] = 'NONE'
    assert score(cat, value)['state'] == '不同款'
