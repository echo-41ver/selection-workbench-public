from __future__ import annotations

import json
import subprocess
import sys
import pytest
from pathlib import Path

from openpyxl import load_workbook

from selection_workflow.bath_towel_layer3_v1 import MODULES, score_facts
from selection_workflow.bath_towel_cache import refresh_candidate_identity

ROOT = Path(__file__).resolve().parents[1]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def visual_facts(relation: str) -> dict:
    return {
        "modules": {
            name: {
                "relation": relation,
                "amazon_evidence": "Amazon短证据" if relation != "UNKNOWN" else "",
                "temu_evidence": "TEMU短证据" if relation != "UNKNOWN" else "",
                "comparison_reason": "",
                "hard_conflict_code": "NONE",
            }
            for name in MODULES
        },
        "ignored_elements": [],
        "overall_confidence": 0.9 if relation != "UNKNOWN" else 0.0,
        "evidence_quality": "HIGH" if relation != "UNKNOWN" else "LOW",
        "variant_differences": [],
        "uncertainties": [] if relation != "UNKNOWN" else list(MODULES),
        "same_base_product": True if relation != "UNKNOWN" else None,
    }


@pytest.mark.parametrize("max_rows", [2, 1048572])
def test_layer123_export_adds_low_match_sheet_with_same_headers(tmp_path: Path, monkeypatch, max_rows) -> None:
    layer12_rows = tmp_path / "rows.jsonl"
    summary = tmp_path / "summary.json"
    candidates = tmp_path / "candidates.jsonl"
    results = tmp_path / "results.jsonl"
    output = tmp_path / "layer123.xlsx"
    pairs = []
    for asin in ("A1", "A2"):
        for goods in ("T1", "T2"):
            status = "CANDIDATE" if asin == "A1" else ("BELOW_VECTOR_THRESHOLD" if goods == "T1" else "STAGE1_EXCLUDED")
            pairs.append({
                "asin": asin, "goods_id": goods, "amazon_title": f"Amazon {asin}", "temu_title": f"TEMU {goods}",
                "amazon_image_url": "https://example.com/a.jpg", "temu_image_url": "https://example.com/t.jpg",
                "amazon_product_url": "https://example.com/a", "temu_product_url": "https://example.com/t",
                "stage1_text_fact": "未发现硬冲突", "vector_score": 80.0, "pair_status": status,
                "pair_revision": f"pair-{asin}-{goods}",
            })
    write_jsonl(layer12_rows, pairs)
    summary.write_text(json.dumps({
        "label": "fixture", "snapshot_at": "2026-09-22T00:00:00Z", "actual_pairs": 4,
        "amazon_products": 2, "temu_products": 2, "candidate_threshold": 60,
        "states": {"CANDIDATE": 2, "BELOW_VECTOR_THRESHOLD": 1, "STAGE1_EXCLUDED": 1},
    }, ensure_ascii=False), encoding="utf-8")
    image_path = tmp_path / "image.jpg"
    image_path.write_bytes(b"offline-cache-fixture")
    candidate_rows = [
        refresh_candidate_identity({**row, "amazon_image_path": str(image_path), "temu_image_path": str(image_path)}, expected_model="deepseek-flash", expected_endpoint="https://api.deepseek.com")
        for row in pairs if row["pair_status"] == "CANDIDATE"
    ]
    write_jsonl(candidates, candidate_rows)
    result_rows = []
    for index, row in enumerate(candidate_rows):
        facts = visual_facts("SAME" if index == 0 else "UNKNOWN")
        result_rows.append({
            "asin": row["asin"], "goods_id": row["goods_id"], "pair_revision": row["pair_revision"],
            "vision_revision_key": row["vision_revision_key"], "api_status": "success",
            "visual_facts": facts, "module_result": score_facts(facts),
            "usage": {"total_tokens": 100}, "elapsed_seconds": 1.2, "model": "deepseek-flash",
        })
    write_jsonl(results, result_rows)
    command = [
        sys.executable, str(ROOT / "scripts/export_incremental_layer123_excel.py"),
        "--layer12-rows", str(layer12_rows), "--layer12-summary", str(summary),
        "--candidates", str(candidates), "--results", str(results), "--output", str(output),
        "--max-data-rows-per-sheet", str(max_rows),
    ]
    # Export must use its own category identity, not another task's model env.
    monkeypatch.setenv("OPENAI_VISION_MODEL", "gpt-5.6-terra")
    subprocess.run(command, cwd=ROOT, check=True, capture_output=True, text=True, encoding="utf-8", errors="replace")
    subprocess.run([
        sys.executable, str(ROOT / "scripts/validate_incremental_layer123_excel.py"),
        "--book", str(output), "--layer12-summary", str(summary),
    ], cwd=ROOT, check=True, capture_output=True, text=True, encoding="utf-8", errors="replace")
    workbook = load_workbook(output, read_only=True, data_only=False)
    main_name = "全量综合结果" if max_rows > 4 else "全量综合结果_1"
    assert workbook.sheetnames[0] == "汇总与规则" and workbook.sheetnames[-1] == "已核验采样同款不超过20"
    main_headers = [cell.value for cell in workbook[main_name][4]]
    low_headers = [cell.value for cell in workbook["已核验采样同款不超过20"][4]]
    assert low_headers == main_headers
    assert {str(row[2]) for row in workbook["已核验采样同款不超过20"].iter_rows(min_row=5, values_only=True)} == {"A2"}
    summary_cells = {row[0]: row[1] for row in workbook["汇总与规则"].iter_rows(min_row=2, values_only=True)}
    assert summary_cells["第三层API成功"] == 2
    assert summary_cells["第三层确认同款"] == 1
    assert summary_cells["第三层待补充证据"] == 1
    values = list(workbook[main_name].iter_rows(min_row=5, values_only=True))
    assert values[0][21] == "确认同款" and values[0][20] == 100
    workbook.close()
