from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from export_category_workbooks import export_amazon_book, export_stage_book, export_temu_book


def test_operator_sheet_one_product_preserves_unknowns_and_real_review_sources(tmp_path):
    import xlsxwriter
    from export_category_workbooks import add_operator_sheet,_formats
    image=tmp_path/'fixture.png';Image.new('RGB',(100,100),'blue').save(image)
    amazon={'title':'=literal fixture','image':'https://example.com/a.jpg','url':'https://amazon.com/dp/A1',
            'price':{'observed_price':19},'sales':{'text':'100+ bought in past month','basis':'最近一月页面显示'},
            'negative_reviews':[{'rating':1,'text':'fixture defect','source_url':'https://amazon.com/review/fixture'}]}
    temu={'title':'fixture temu','url':'https://temu.com/fixture.html','price':{'min':9,'max':12,'basis':'页面显示区间'},
          'sales':{'text':'500+ sold','basis':'周期未知'},'rating':4.6,'listed_at':'2026-10-01','listed_at_source':'页面显示'}
    path=tmp_path/'operator-fixture.xlsx'
    book=xlsxwriter.Workbook(path,{'strings_to_formulas':False})
    add_operator_sheet(book,{'same_limit':5,'products':[{'asin':'A1','amazon':amazon,'temu':temu,'same_count':1,'reason':'fixture'}]},
                       _formats(book),{amazon['image']:image});book.close()
    reopened=load_workbook(path)
    sheet=reopened['运营候选']
    assert sheet.max_row==5 and len(sheet._images)==1
    assert sheet['B5'].data_type=='s' and sheet['B5'].value=='=literal fixture'
    assert '区间未取得' in sheet['E5'].value and '周期未知' in sheet['M5'].value
    assert 'https://amazon.com/review/fixture' in sheet['P5'].value
    assert sheet['Q5'].value=='尚未采集真实差评'
    reopened.close()


def test_temu_dual_sort_sheets_keep_separate_sales_observations(tmp_path):
    rows = [{'goods_id':'same','title':'drill','search_query':'drill','sales_text':f'{count}+ sold',
             'search_sort':{'requested':order,'state':'UI_CONFIRMED','label':order}}
            for order,count in [('sales',100),('recent',120)]]
    destination = tmp_path/'TEMU.xlsx'
    export_temu_book(destination, [rows[-1]], '电钻', 100, samples=rows)
    book = load_workbook(destination)
    try:
        assert book['商品明细'].max_row == 5
        assert book['销量最高样本']['D4'].value == '100+ sold'
        assert book['最近上架样本']['D4'].value == '120+ sold'
    finally:
        book.close()


def test_amazon_reference_headers_and_real_embedded_image(tmp_path):
    picture = tmp_path / 'source.png'
    Image.new('RGB', (400, 300), 'blue').save(picture)
    source = {'source_url': 'https://amazon.example/bestsellers', 'observed_at': '2026-09-25T00:00:00Z',
              'records': [{'rank': 1, 'asin': 'A1', 'title': 'electric drill', 'local_image_path': str(picture),
                           'main_image_url': 'https://images.example/a1.jpg', 'product_url': 'https://amazon.example/dp/A1',
                           'price_usd': 19.99, 'specifications': [{'label': 'voltage', 'value': '20V'}]}]}
    output = tmp_path / 'excel' / 'Amazon.xlsx'
    export_amazon_book(output, source, '电钻')
    workbook = load_workbook(output)
    sheet = workbook['商品列表']
    assert sheet['F4'].value == 'Amazon主图'
    assert len(sheet._images) == 1
    assert sheet._images[0].anchor._from.row == 4 and sheet._images[0].anchor._from.col == 5
    # Formula links retain clickability without Excel's relationship-count limit.
    assert sheet['H5'].value == '=HYPERLINK("https://images.example/a1.jpg","查看图片")'
    assert sheet['M5'].value == 19.99
    assert sheet.freeze_panes == 'E5'
    assert sheet.auto_filter.ref.startswith('A4:')
    assert workbook['采集说明']['B7'].value == '1/1件'
    workbook.close()


def test_temu_and_stage_workbooks_keep_reference_navigation(tmp_path):
    temu = tmp_path / 'TEMU.xlsx'
    export_temu_book(temu, [{'goods_id': 'T1', 'title': 'drill', 'price_usd': 12.5,
                             'product_url': 'https://temu.example/t1', 'source_queries': ['drill']}], '电钻', 100)
    workbook = load_workbook(temu)
    assert workbook.sheetnames == ['商品明细', '汇总与审计', '销量最高样本', '最近上架样本']
    assert workbook['商品明细']['B4'].value == 'TEMU商品ID'
    assert workbook['商品明细']['B5'].value == 'T1'
    assert workbook['商品明细'].freeze_panes == 'C5'
    workbook.close()

    stage = tmp_path / 'stage.xlsx'
    headers = ['ASIN', 'Amazon标题', 'Amazon商品链接', 'Amazon主图链接', 'TEMU商品ID',
               'TEMU标题', 'TEMU商品链接', 'TEMU主图链接', '文字判断依据', '向量分', '第二层状态', '最终判断']
    record = ['A1', 'Amazon drill', 'https://amazon.example/a1', 'https://images.example/a1.jpg',
              'T1', 'TEMU drill', 'https://temu.example/t1', 'https://images.example/t1.jpg',
              '未发现硬冲突', 70.5, '进入第三层：图片向量达到候选门槛', '待运行']
    export_stage_book(stage, headers, iter([record]), {'商品对总数': 1}, '电钻', third_layer=True, low_asins={'A1'})
    workbook = load_workbook(stage)
    assert workbook.sheetnames == ['汇总与规则', '全量综合结果', '同款不超过20_同款及待复核']
    assert workbook['全量综合结果']['C4'].value == 'Amazon ASIN'
    assert workbook['全量综合结果']['C5'].value == 'A1'
    assert workbook['全量综合结果']['M5'].value == '待运行'
    assert workbook['全量综合结果'].freeze_panes == 'A5'
    workbook.close()
