"""New-rule task preparation, real local L1/L2 processing and L3/export contract."""
import json
from pathlib import Path
from types import SimpleNamespace
import numpy as np
from PIL import Image
from openpyxl import load_workbook
import generate_category as generator
import selection
import pipeline
from rule_editor import draft_view,edit_rules
from extract_match_criteria import extract
from test_all_product_visual_criteria import response
from selection_automation.match_criteria import select,attach
from selection_automation.compact_visual import decode
from selection_automation.rules import score
from selection_automation.visual_decision import binary_mode
from selection_workflow.incremental_stage12 import IncrementalStage12,SourceSpec,PathSpec
from export_category_workbooks import amazon_conclusions,export_stage_book


def test_latest_rules_prepare_filter_three_classes_and_export(tmp_path,monkeypatch):
    receipt=json.loads((Path(__file__).parent/'fixtures/compact-monitor-model-response.json').read_text('utf-8'))
    evidence={'url':receipt['scope']['listing'],'title':'Best Sellers','heading':'Best Sellers in Computer Monitors','rank_count':50,'links':[]}
    monkeypatch.setenv('DEEPSEEK_API_KEY','TEST_ONLY_KEY')
    monkeypatch.setattr(generator,'approved_category',lambda *a:None)
    monkeypatch.setattr(generator,'read_listing',lambda _:evidence)
    monkeypatch.setattr(generator,'model_json',lambda *a,**k:receipt['proposal'])
    category=generator.generate_category('电脑显示器',2,10,bestsellers_url=evidence['url'])
    snapshot={'source_url':evidence['url'],'records':[{'asin':f'A{i:03}','title':f'Monitor {size} inches LCD',
        'main_image_url':f'https://m.media-amazon.com/{i}.jpg',
        'specifications':[{'label':'Screen Size','value':f'{size} inches'}]} for i,size in enumerate((7,8))]}
    extract(category,snapshot,response)
    size=next(row['id'] for row in category['match_criteria']['items'] if row.get('label')=='Screen Size')
    attach(category,select(category['match_criteria'],[size]))
    category=edit_rules(category,draft_view(category,'draft'),compile_model=lambda *a,**k:(_ for _ in ()).throw(AssertionError('paid repair')))
    source=tmp_path/'reviewed.json';source.write_text(json.dumps(category),encoding='utf-8')
    monkeypatch.setattr(selection,'ROOT',tmp_path);monkeypatch.setattr(pipeline,'ROOT',tmp_path)
    frozen,category,base=selection.prepare(source)
    config=json.loads((base/'stage12_config.json').read_text('utf-8'))
    assert config['text_gate']=='declarative:'+str(frozen) and not binary_mode(category)
    assert config['candidate_threshold']==category['stage2']['candidate_threshold']
    image=tmp_path/'same.png';low=tmp_path/'low.png'
    Image.new('RGB',(256,256),'red').save(image,compress_level=0);Image.new('RGB',(256,256),'blue').save(low,compress_level=0)
    ap=tmp_path/'amazon.jsonl';tp=tmp_path/'temu.jsonl'
    ap.write_text(json.dumps({'asin':'A000','title':'Monitor 7 inches LCD','local_image_path':str(image)})+'\n')
    rows=[{'goods_id':gid,'title':title,'local_image_path':str(path)} for gid,title,path in
          [('size_conflict','Monitor 8 inches LCD',image),('equal','Monitor 7 inches LCD',image),
           ('missing','Monitor LCD',image),('below','Monitor LCD',low)]]
    tp.write_text(''.join(json.dumps(row)+'\n' for row in rows))
    vector=SimpleNamespace(revision='LOCAL_CONTRACT_ONLY',encode=lambda path:np.array([0.,1.]) if Image.open(path).getpixel((0,0))[2]>200 else np.array([1.,0.]))
    engine=IncrementalStage12(tmp_path/'db',vectorizer=vector,text_gate_name=config['text_gate'],threshold=config['candidate_threshold'])
    try:
        engine.configure([SourceSpec('a','AMAZON','amazon_generic',ap,'replace'),SourceSpec('t','TEMU','temu_frontend',tp,'replace')],[PathSpec('task','a','t','NEW_RULES')])
        engine.ingest_source('a');engine.ingest_source('t');engine.process_batch(20)
        states=dict(engine.con.execute('SELECT goods_id,status FROM pairs'))
        assert states=={'size_conflict':'STAGE1_EXCLUDED','equal':'CANDIDATE','missing':'CANDIDATE','below':'BELOW_VECTOR_THRESHOLD'}
        assert engine.materialize_candidates(tmp_path/'candidate-output')
    finally:engine.close()
    outcomes=[]
    for relation,expected in [('S','确认同款'),('L','相似款'),('D','不同款')]:
        raw={'candidates':[{'i':'C00','r':relation,'a':'矩形面板与中置支架','t':'矩形面板与中置支架',
                            'why':'主体造型及对应结构已比较','q':{},'c':.95}]}
        facts=decode(category,raw,['C00'])['C00']
        assert score(category,facts)['state']==expected
        outcomes.append(expected)
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题',
             'TEMU商品链接','TEMU主图链接','文字依据','向量分','第二层状态','最终判断']
    records=[['A000','monitor','','',str(i),'monitor','','','',90,'候选',state] for i,state in enumerate(outcomes)]
    conclusions=amazon_conclusions(headers,records,3,same_only=True,operator='lt')
    assert conclusions[0]['same']==1 and conclusions[0]['similar']==1 and conclusions[0]['different']==1
    assert conclusions[0]['selected']
    assert not amazon_conclusions(headers,records,2,same_only=True,operator='lt')[0]['selected']
    destination=tmp_path/'three-classes.xlsx'
    export_stage_book(destination,headers,records,{'视觉结论口径':'automatic-visual-v1','采集前设置的同款上限':3,'同款数量比较方式':'lt'},
        '监视器',third_layer=True,low_asins={'A000'},conclusions=conclusions,same_only=True)
    book=load_workbook(destination)
    assert book['Amazon最终结论']['N4'].value=='类似款数'
    assert book['Amazon最终结论']['N5'].value==1
    detail=book['同款及类似款少于3_明细']
    assert detail.max_row==6
    book.close()
