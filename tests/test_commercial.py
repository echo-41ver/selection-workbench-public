"""Synthetic observations test arithmetic; they are never live market evidence."""
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

from selection_automation.commercial import (History, build_report, changes, market_key,
    market_metrics, match_metric, sales_evidence, settings)

NOW = datetime.now(timezone.utc).replace(microsecond=0)


def category(**policy):
    return {'category_key': 'fixture', 'category_name_zh': '电钻',
            'amazon': {'bestsellers_url': 'https://www.amazon.com/zgbs/tools/123'},
            'temu': {'base_queries': ['drill']}, 'commercial': policy}


def observation(gid, days, amount, **extra):
    return {'goods_id': gid, 'observed_at': (NOW-timedelta(days=days)).isoformat(),
            'sales_count': amount, 'sales_precision': 'exact', 'sales_basis': 'lifetime',
            'product_url': 'https://example.invalid/product/'+gid, 'source_queries': ['drill'],
            'listed_at': (NOW-timedelta(days=80)).isoformat(), 'listed_at_source': 'fixture-detail', **extra}


@pytest.mark.parametrize('counts,expected,state', [
    ({'confirmed':20},20,'SAMPLE_PASS'), ({'confirmed':21,'pending':3},24,'FAIL'),
    ({},0,'UNKNOWN'), ({},20,'UNKNOWN'), ({'confirmed':1,'possible':1},2,'UNKNOWN'),
    ({'confirmed':2,'screened_out':3,'text_excluded':4},9,'SAMPLE_PASS'),
])
def test_match_does_not_turn_incomplete_analysis_into_low_competition(counts,expected,state):
    result=match_metric(counts,expected,settings())
    assert result['state']==state
    assert result['confirmed']==counts.get('confirmed',0)


def test_same_and_similar_share_the_strict_final_limit():
    policy=settings({'same_limit':5,'same_operator':'lt'})
    result=match_metric({'confirmed':2,'similar':3},5,policy)
    assert result['state']=='FAIL' and result['same_or_similar']==5
    assert match_metric({'confirmed':1,'similar':3},4,policy)['state']=='SAMPLE_PASS'
    assert match_metric({'confirmed':1,'similar':3,'possible':1},5,policy)['state']=='UNKNOWN'


@pytest.mark.parametrize('value',[[], {'same_limit':True}, {'new_days':0},
                                   {'min_sales_coverage':0}, {'growth_min_percent':float('nan')}])
def test_invalid_settings_fail(value):
    with pytest.raises(ValueError): settings(value)


def test_displayed_sales_are_not_precise_lifetime_totals():
    for raw in ('14K+ sold','1.3K+ sold','123 sold'):
        evidence=sales_evidence({'sales_text':raw})
        assert evidence['value'] is None and evidence['basis']=='displayed_unspecified'
    assert sales_evidence({'sales_text':'14K+ sold'})['lower']==14000
    assert sales_evidence(observation('1',0,123))['value']==123
    assert sales_evidence({**observation('1',0,123),'product_url':None})['value'] is None


def test_history_is_idempotent_shared_across_runs_and_isolated_by_category(tmp_path):
    store=History(tmp_path/'history.db');cat=category()
    row=observation('1',0,100)
    assert store.ingest(cat,[row])==1
    assert store.ingest({**cat,'category_key':'next_run'},[row])==0
    assert store.ingest(cat,[{**row,'sales_count':101}])==1
    assert len(store.rows(cat))==2
    other=category();other['amazon']['bestsellers_url']='https://www.amazon.com/zgbs/tools/456'
    assert store.rows(other)==[]
    store.close()


def test_missing_listed_date_never_uses_first_seen_as_listing():
    row=observation('1',0,100,listed_at=None,first_seen_at=NOW.isoformat())
    report=market_metrics(category(),[row],[row],NOW)
    assert report['dated_new_products']==0
    assert report['growth']['state']==report['concentration']['state']=='UNKNOWN'


@pytest.mark.parametrize('last,state',[(220,'FAIL'),(221,'SAMPLE_PASS')])
def test_week_growth_strict_boundary(last,state):
    history=[observation('1',14,0),observation('1',7,100),observation('1',0,last)]
    report=market_metrics(category(min_new_products=1,growth_period='week'),[history[-1]],history,NOW)
    assert report['growth']['state']==state
    assert report['growth']['periods']['week']['valid_products']==1


def test_zero_baseline_counter_reset_and_conflicting_snapshots_are_unknown():
    for values in [(0,0,10),(100,200,150)]:
        rows=[observation('1',days,count) for days,count in zip([14,7,0],values)]
        report=market_metrics(category(min_new_products=1,growth_period='week'),rows[-1:],rows,NOW)
        assert report['growth']['state']=='UNKNOWN'
    rows=[observation('1',14,0),observation('1',7,100),observation('1',0,230),observation('1',0,250)]
    assert changes(rows,NOW,7) is None


@pytest.mark.parametrize('amounts,state,share', [([90,10],'FAIL',90),([89,11],'SAMPLE_PASS',89)])
def test_concentration_is_recent_sales_not_search_position(amounts,state,share):
    rows=[];current=[]
    for index,amount in enumerate(amounts):
        rows.extend([observation(str(index),30,5000),observation(str(index),0,5000+amount)])
        current.append(rows[-1])
    report=market_metrics(category(top_n=1,min_market_products=2),list(reversed(current)),rows,NOW)
    assert report['concentration']['state']==state
    assert report['concentration']['value_percent']==share
    assert market_metrics(category(),current,rows,NOW)['concentration']['state']=='UNKNOWN'


def test_keyword_scope_and_insufficient_sales_coverage():
    rows=[observation('1',0,10,source_queries=['cordless drill'])]
    assert market_metrics(category(),rows,rows,NOW)['market_products']==0
    history=[observation('1',14,0),observation('1',7,100),observation('1',0,230)]
    current=[history[-1],observation('2',0,100)]
    report=market_metrics(category(min_new_products=1,growth_period='week'),current,[*history,current[-1]],NOW)
    assert report['growth']['periods']['week']['coverage']==.5
    assert report['growth']['state']=='UNKNOWN'


def test_generated_feature_queries_cannot_stand_in_for_category_sample():
    cat=category(top_n=1,min_market_products=2)
    cat['temu']['base_queries']=['drill','cordless drill','drill driver']
    history=[];current=[]
    for gid,amount in [('1',90),('2',10)]:
        history.extend([observation(gid,30,100,source_queries=['cordless drill']),
                        observation(gid,0,100+amount,source_queries=['cordless drill'])])
        current.append(history[-1])
    report=market_metrics(cat,current,history,NOW)
    assert report['market_products']==0 and report['all_frontend_products']==2
    assert report['primary_category_query']=='drill'
    assert report['concentration']['state']=='UNKNOWN'
    # A union row observed under both queries does belong to the category sample.
    for row in current: row['source_queries']=['cordless drill','drill']
    report=market_metrics(cat,current,history,NOW)
    assert report['market_products']==2
    assert report['concentration']['state']=='FAIL'
    assert report['concentration']['value_percent']==90


def test_report_pending_pairs_and_all_four_excel_variants(tmp_path):
    from selection_workflow.commercial_excel import add_report
    import xlsxwriter
    cat=category();base=tmp_path/'runs/fixture'
    (base/'amazon').mkdir(parents=True);(base/'temu/union').mkdir(parents=True)
    (base/'amazon/amazon_frontend_latest.json').write_text(json.dumps({'records':[{'asin':'A1','title':'=fixture'}]}),'utf-8')
    (base/'temu/union/products.jsonl').write_text(json.dumps(observation('1',0,100))+'\n','utf-8')
    report=build_report(tmp_path,cat,as_of=NOW)
    assert report['products'][0]['match']['pending']==1
    assert report['products'][0]['signals']==0
    assert report['market']['frontend_sort']['state']=='UNKNOWN'
    for kind in ('amazon','temu','layer12','layer123'):
        path=tmp_path/(kind+'.xlsx')
        book=xlsxwriter.Workbook(path,{'strings_to_formulas':False})
        book.add_worksheet('原有数据');add_report(book,report);book.close()
        wb=load_workbook(path)
        assert wb.sheetnames==['原有数据','商业选品','市场证据']
        assert wb['商业选品']['B2'].value=='=fixture' and wb['商业选品']['B2'].data_type=='s'
        assert wb['商业选品']['H2'].value=='证据不足'
        evidence=dict(wb['市场证据'].values)
        assert '尚未确认' in evidence['TEMU网页销量排序']
        wb.close()


def test_sort_evidence_is_scoped_to_primary_query_and_not_sales_concentration(tmp_path):
    cat=category();base=tmp_path/'runs/fixture/temu/workers'
    def progress(name,query,state):
        path=base/name/'query/progress.json';path.parent.mkdir(parents=True,exist_ok=True)
        path.write_text(json.dumps({'search_query':query,'sales_sort':{'state':state,'label':'Top sales',
            'observed_at':NOW.isoformat()}}),'utf-8')
    progress('w1','cordless drill','UI_CONFIRMED')
    report=build_report(tmp_path,cat,as_of=NOW)
    assert report['market']['frontend_sort']['state']=='UNKNOWN'
    progress('w2','drill','UI_CONFIRMED')
    report=build_report(tmp_path,cat,as_of=NOW)
    assert report['market']['frontend_sort']['state']=='UI_CONFIRMED'
    assert report['market']['concentration']['state']=='UNKNOWN'
    progress('w3','drill','UNCONFIRMED')
    assert build_report(tmp_path,cat,as_of=NOW)['market']['frontend_sort']['state']=='UNKNOWN'


def test_selection_report_uses_collected_global_top5_not_category_only_denominator(tmp_path):
    cat=category();base=tmp_path/'runs/fixture/temu/union'
    base.mkdir(parents=True)
    rows=[{'goods_id':str(i),'sales_text':f'{i*10} sold',
           'source_queries':['drill'] if i<=6 else ['feature']} for i in range(1,11)]
    (base/'products.jsonl').write_text('\n'.join(json.dumps(row) for row in rows),encoding='utf-8')
    report=build_report(tmp_path,cat,as_of=NOW)
    concentration=report['market']['concentration']
    assert concentration['value_percent']==72.7
    assert concentration['sample_products']==10
    assert concentration['denominator_displayed_sales']==550
    assert concentration['numerator_displayed_sales']==400
    assert concentration['basis']=='DISPLAYED_SALES_APPROXIMATE'
    assert report['market']['primary_category_sales']['top5_displayed_share_approx']==36.4
    assert report['market']['period_concentration']['state']=='UNKNOWN'


def test_commercial_threshold_does_not_change_paid_vision_revision():
    sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
    from rule_editor import draft_view,edit_rules
    cat=json.loads((Path(__file__).resolve().parents[1]/'config/categories/power_drill.json').read_text('utf-8'))
    edits=draft_view(cat,'x');first=edit_rules(cat,edits)
    edits['commercial']['same_limit']=35
    second=edit_rules(cat,edits)
    assert first['stage3']==second['stage3']
    assert second['commercial']['same_limit']==35


def test_cached_report_is_current_for_codex_revisions_and_does_not_scan_frozen_files(tmp_path,monkeypatch):
    from selection_automation import commercial
    cat=category();base=tmp_path/'runs/fixture';base.mkdir(parents=True)
    path=base/'commercial_report.json'
    path.write_text(json.dumps({'legacy':True}),encoding='utf-8')
    frozen=base/'stage123/export_frozen';frozen.mkdir(parents=True)
    huge=frozen/'layer12_rows_fixture.jsonl';huge.write_text('never read in signature',encoding='utf-8')
    original_open=Path.open
    def guarded_open(self,*args,**kwargs):
        if self==huge:raise AssertionError('signature must only stat huge snapshots')
        return original_open(self,*args,**kwargs)
    monkeypatch.setattr(Path,'open',guarded_open)
    calls=[]
    def build(root,category):
        calls.append(True)
        return {'evidence_signature':commercial.report_signature(root,category),'build':len(calls)}
    monkeypatch.setattr(commercial,'build_report',build)
    first=commercial.current_report(tmp_path,cat)
    assert first['build']==1 and not first.get('legacy')
    assert commercial.current_report(tmp_path,cat)==first
    review=base/'stage123/codex_reviews/new.jsonl';review.parent.mkdir(parents=True)
    review.write_text('{}\n',encoding='utf-8')
    assert commercial.current_report(tmp_path,cat)['build']==2
    assert commercial.current_report(tmp_path,cat)['build']==2
    review.write_text('{"changed_revision":true}\n',encoding='utf-8')
    assert commercial.current_report(tmp_path,cat)['build']==3
    changed={**cat,'commercial':{'same_limit':5}}
    assert commercial.current_report(tmp_path,changed)['build']==4
    review.unlink()
    assert commercial.current_report(tmp_path,changed)['build']==5


def test_signature_ignores_heartbeat_and_identical_small_product_rewrites(tmp_path):
    from selection_automation.commercial import report_signature
    cat=category();folder=tmp_path/'runs/fixture/temu/workers/w1/drill'
    folder.mkdir(parents=True)
    progress=folder/'progress.json';products=folder/'products.jsonl'
    data={'search_query':'drill','sales_sort':{'state':'PENDING'},'updated_at':'first'}
    progress.write_text(json.dumps(data),'utf-8');products.write_text('','utf-8')
    first=report_signature(tmp_path,cat)
    data.update(updated_at='later',round=10,image_policy={'blocked_requests':12})
    progress.write_text(json.dumps(data),'utf-8');products.write_text('','utf-8')
    assert report_signature(tmp_path,cat)==first
    products.write_text('{"goods_id":"new"}\n','utf-8')
    second=report_signature(tmp_path,cat)
    assert second!=first
    products.write_text('{"goods_id":"new"}\n','utf-8')
    assert report_signature(tmp_path,cat)==second
    data['sales_sort']={'state':'UI_CONFIRMED','label':'Top sales','observed_at':'now'}
    progress.write_text(json.dumps(data),'utf-8')
    assert report_signature(tmp_path,cat)!=second
