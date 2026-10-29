"""Portable, allowlisted business snapshots for a separate Linux/Docker consumer."""
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

VERSION = 'developer-handoff-v2'
CONTRACT = {
    'schema_version': VERSION, 'encoding': 'UTF-8', 'format': 'JSONL and XLSX',
    'identity': ['platform', 'product_id'], 'unknown': None,
    'files': ['amazon.jsonl', 'temu.jsonl', 'amazon.xlsx', 'temu.xlsx'],
    'excel_sheets': {'Products':'one row per product', 'Images':'one row per public image URL',
                     'Specifications':'one row per product specification',
                     'Observations':'one row per keyword/sort observation',
                     'Readme':'scope and unknown-value guidance'},
    'fields': {'product_id':'string|null', 'platform':'amazon|temu', 'title':'string|null',
        'url':'https public product URL|null', 'images':'array of public HTTPS URLs',
        'specifications':'array of {label,value}', 'source':'string|null', 'observed_at':'ISO timestamp|null',
        'price_usd':'number|null', 'rating':'number|null', 'review_count':'number|null',
        'sales_text':'observed text|null; never exact sales', 'brand':'string|null',
        'sampling':'rank, keyword, keywords, sort, observations, first_seen_round, first_seen_ordinal',
        'quality':'detail_status,image_status,missing_fields'},
    'limitations':['Snapshot only; never a platform-wide inventory.',
        'Empty arrays and null mean unavailable, not negative evidence.',
        'Search order is not a sales rank; only explicit sort UI evidence may establish sorting.',
        'An image URL alone is not an image-quality or same-product verification.',
        'No image bytes, credentials, worker configuration or local paths are included.'],
}
README = '''# 第二部分 Linux / Docker 数据交接 v2

这是第一部分当前任务的只读业务快照，不是平台完整全集。查看 manifest.json 的 counts、collection、provenance、warnings 和 files 校验值。采集中导出跨文件不保证同一事务时刻；不要把样本量当市场规模。

输入文件：amazon.xlsx、temu.xlsx 供 Excel 读取；amazon.jsonl、temu.jsonl 供程序读取。每份 Excel 的 Products、Images、Specifications、Observations 工作表用 product_id 关联。Excel 空白格与 JSONL 的 null/空数组均表示未知。图片是公开原始地址，不包含图片文件。保留榜单 rank、关键词、前端排序证据和多次观察来源；无销量周期/准确销量不得自行补造。

交付要求：提供 README、Dockerfile、requirements.txt 及精确锁定依赖文件（或等效锁文件），清晰 CLI 入口、最小可运行示例、测试、输出字段契约及失败重试说明。实现内容为独立第二部分，不改第一部分采集流程。

建议 CLI 契约（待后续实现，不是本包已有算法）：
python -m part2 --input-dir /data/input --output-dir /data/output

docker run --rm --env-file .env -v "$PWD/input:/data/input:ro" -v "$PWD/output:/data/output" part2:version --input-dir /data/input --output-dir /data/output

API 密钥只通过运行环境变量注入（如 MODEL_API_KEY）；.env 由部署方单独创建，禁止提交源码、镜像或交接包。输入只读，输出写独立目录。不要依赖 Windows 路径、浏览器 profile、第一部分账户或代理。

建议输出 matches.jsonl（amazon_product_id、temu_product_id、结论、证据、原因、版本）、analysis.json、errors.jsonl。判定为 CONFIRMED_SAME / NOT_SAME / UNKNOWN（可另设候选）；UNKNOWN 不得记为不同款。第三层付费处理独立显式启动，默认关闭。第二部分可读取 Excel 或 JSONL，但应以 product_id 关联；内部数据引用使用相对目录。缺证据输出 UNKNOWN；印花/结构等硬冲突不可当同款。商业分析保留样本来源和统计边界。算法、模型与价格配置由第二部分明确说明，本模块不实施第二部分算法。
'''


def _text(value):
    if not isinstance(value, str) or not value.strip(): return None
    if re.search(r'(?i)((?<![a-z])[a-z]:[\\/]|\\\\|/(?:home|users|tmp|etc|mnt|opt)/|(?:api[_ -]?key|password|passwd|cookie|authorization|proxy_user|access_token)\s*[:=]|bearer\s+|sk-[A-Za-z0-9_-]{12,})', value): return None
    return value.strip()


def _url(value):
    value = _text(value)
    if not value: return None
    try:
        u = urlsplit(value)
        host = (u.hostname or '').lower()
        allowed = ('amazon.com','m.media-amazon.com','images-na.ssl-images-amazon.com','images-eu.ssl-images-amazon.com','temu.com','kwcdn.com')
        if u.scheme != 'https' or u.username or u.password or u.port or not any(host == h or host.endswith('.'+h) for h in allowed): return None
        query = urlencode([(k,v) for k,v in parse_qsl(u.query) if k in ('goods_id','search_key','node') and _text(v)])
        return urlunsplit(('https',host,u.path,query,''))
    except ValueError: return None


def _number(value):
    return value if type(value) in (int,float) and math.isfinite(value) else None


def _sort(value):
    value = value if isinstance(value,dict) else {}
    return {k:_text(value.get(k)) for k in ('requested','state','label','confirmed_by','observed_at')}


def _row(row, platform, meta):
    urls = [row.get('main_image_url'),row.get('thumbnail_url')]
    for key in ('all_image_urls','observed_image_urls'):
        if isinstance(row.get(key),list): urls.extend(row[key])
    images = list(dict.fromkeys(u for value in urls if (u := _url(value))))
    specs = [{'label':_text(s.get('label')),'value':_text(s.get('value'))} for s in row.get('specifications',[]) if isinstance(s,dict)] if isinstance(row.get('specifications'),list) else []
    quality = row.get('detail_quality') if isinstance(row.get('detail_quality'),dict) else {}
    observations = []
    for ob in row.get('observation_sources',[]) if isinstance(row.get('observation_sources'),list) else []:
        if isinstance(ob,dict): observations.append({'keyword':_text(ob.get('search_query')),'observed_at':_text(ob.get('observed_at')),'sort':_sort(ob.get('search_sort'))})
    pid = row.get('asin' if platform=='amazon' else 'goods_id')
    return {'product_id':_text(str(pid)) if isinstance(pid,(str,int)) else None,'platform':platform,
        'title':_text(row.get('title')),'url':_url(row.get('product_url')),'images':images,'specifications':specs,
        'source':_text(row.get('source')) or _text(meta.get('source')),'source_url':_url(row.get('source_url')) or _url(meta.get('source_url')),
        'observed_at':_text(row.get('observed_at')) or _text(meta.get('observed_at')),
        'price_usd':_number(row.get('price_usd')),'rating':_number(row.get('rating')),'review_count':_number(row.get('review_count')),
        'sales_text':_text(row.get('sales_text')),'brand':_text(row.get('brand')),
        'sampling':{'rank':_number(row.get('rank')),'keyword':_text(row.get('search_query')),
            'keywords':[t for v in row.get('source_queries',[]) if (t:=_text(v))] if isinstance(row.get('source_queries'),list) else [],
            'sort':_sort(row.get('search_sort')),'observations':observations,
            'first_seen_round':_number(row.get('first_seen_round')),'first_seen_ordinal':_number(row.get('first_seen_ordinal'))},
        'quality':{'detail_status':_text(row.get('detail_status')),'image_status':_text(row.get('image_status')),
            'missing_fields':[t for v in quality.get('missing_fields',[]) if (t:=_text(v))] if isinstance(quality.get('missing_fields'),list) else []}}


def _base(root,key):
    if not isinstance(key,str) or not re.fullmatch(r'[a-z][a-z0-9_]{0,127}',key): raise ValueError('Invalid run key')
    runs = (Path(root)/'runs').resolve()
    base = (runs/key).resolve()
    if base.parent != runs: raise ValueError('Run path escapes runs directory')
    if not base.is_dir(): raise FileNotFoundError('Run not found')
    return base


def _read(base,relative):
    p=(base/relative).resolve()
    if not p.is_relative_to(base): raise ValueError('Data file escapes run directory')
    return p.read_text('utf-8-sig') if p.is_file() else None


def _json(text):
    return json.loads(text) if text else {}


def _encode(value):
    return (json.dumps(value,ensure_ascii=False,allow_nan=False,indent=2)+'\n').encode('utf-8')


def _cell(value):
    if value is None: return None
    if isinstance(value,str):
        value=value[:32767]
        return "'"+value if value.startswith(('=','+','-','@')) else value
    return value


def _workbook(rows, platform):
    book=Workbook(); products=book.active;products.title='Products'
    fields=('product_id','platform','title','url','main_image_url','image_count','price_usd',
            'rating','review_count','sales_text','brand','rank','keyword','keywords',
            'sort_requested','sort_state','sort_confirmed_by','observed_at','source',
            'source_url','detail_status','image_status','missing_fields')
    products.append(fields)
    images=book.create_sheet('Images');images.append(('product_id','image_number','image_url'))
    specs=book.create_sheet('Specifications');specs.append(('product_id','label','value'))
    observations=book.create_sheet('Observations');observations.append(('product_id','keyword','observed_at','sort_requested','sort_state'))
    for row in rows:
        sampling=row['sampling'];quality=row['quality'];image_urls=row['images']
        values=(row['product_id'],platform,row['title'],row['url'],image_urls[0] if image_urls else None,
                len(image_urls),row['price_usd'],row['rating'],row['review_count'],row['sales_text'],
                row['brand'],sampling['rank'],sampling['keyword'],' | '.join(sampling['keywords']),
                sampling['sort']['requested'],sampling['sort']['state'],sampling['sort']['confirmed_by'],
                row['observed_at'],row['source'],row['source_url'],quality['detail_status'],
                quality['image_status'],' | '.join(quality['missing_fields']))
        products.append(tuple(_cell(value) for value in values))
        for number,url in enumerate(image_urls,1): images.append((_cell(row['product_id']),number,url))
        for spec in row['specifications']: specs.append((_cell(row['product_id']),_cell(spec['label']),_cell(spec['value'])))
        for observation in sampling['observations']:
            observations.append((_cell(row['product_id']),_cell(observation['keyword']),observation['observed_at'],
                                 observation['sort']['requested'],observation['sort']['state']))
    readme=book.create_sheet('Readme')
    for line in ('采集快照；只代表本次已采商品，不代表平台完整全集。',
                 'Products 每行一件商品；其他工作表通过 product_id 关联。',
                 '图片工作表仅含公开 HTTPS 地址，不包含图片文件。',
                 '空白格表示未知；页面展示销量的周期不明。',
                 'sort_state=UI_CONFIRMED 才能认定前端已选择目标排序。'):
        readme.append((line,))
    for sheet in (products,images,specs,observations):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(color='FFFFFF',bold=True)
            cell.fill=PatternFill('solid',fgColor='185B79')
            cell.alignment=Alignment(wrap_text=True)
        sheet.column_dimensions['A'].width=22
        for col in range(2,sheet.max_column+1): sheet.column_dimensions[sheet.cell(1,col).column_letter].width=24
    readme.column_dimensions['A'].width=85
    output=io.BytesIO();book.save(output);return output.getvalue()


def _prepare(root,key):
    base=_base(root,key); warnings=[]; data={}; provenance={}
    category=_json(_read(base,'category.json'))
    for platform,relative in [('amazon','amazon/amazon_frontend_latest.json'),('temu','temu/union/products.jsonl')]:
        raw=_read(base,relative); meta={}; rows=[]
        if platform=='amazon' and raw:
            meta=_json(raw); rows=meta.get('records',[])
        else:
            if platform=='amazon': relative='amazon/amazon_frontend_latest.jsonl';raw=_read(base,relative)
            if raw:
                for line in raw.splitlines():
                    if line.strip(): rows.append(json.loads(line))
        if not raw: warnings.append(platform+': no data available')
        if not isinstance(rows,list): raise ValueError('Invalid product records')
        normalized=[_row(r,platform,meta) for r in rows if isinstance(r,dict)]
        # Preserve unknown-ID rows; deduplicate known IDs within this platform only.
        seen=set(); output=[]
        for r in normalized:
            pid=r['product_id']
            if pid and pid in seen: continue
            if pid: seen.add(pid)
            output.append(r)
        data[platform]=output
        provenance[platform]={'input':relative,'source':_text(meta.get('source')) or next((r['source'] for r in output if r['source']),None),
            'status':_text(meta.get('status')),'complete':meta.get('complete') is True if platform=='amazon' else False}
    state=_json(_read(base,'status.json')); status=_text(state.get('状态') or state.get('status'))
    collecting=False if status in ('已停止','已完成','STOPPED','COMPLETE') else True if status in ('运行中','采集中','RUNNING') else None
    files={'amazon.jsonl':b''.join((json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n').encode('utf-8') for r in data['amazon']),
           'temu.jsonl':b''.join((json.dumps(r,ensure_ascii=False,allow_nan=False)+'\n').encode('utf-8') for r in data['temu']),
           'amazon.xlsx':_workbook(data['amazon'],'amazon'),
           'temu.xlsx':_workbook(data['temu'],'temu'),
           'contract.json':_encode(CONTRACT),'README.md':README.encode('utf-8')}
    manifest={'schema_version':VERSION,'run_key':key,'category_name':_text(category.get('category_name_zh')),'exported_at':datetime.now(timezone.utc).isoformat(),
        'contract':CONTRACT,'counts':{p:len(rows) for p,rows in data.items()},'snapshot':True,'platform_complete':False,
        'collection':{'recorded_status':status,'is_collecting':collecting,'status_basis':'persisted_status_only'},
        'provenance':provenance,'warnings':warnings+['Snapshot; not a complete platform inventory. Collection status is recorded state, not a live process check.', 'Files are read sequentially; collection can change while exporting.'],
        'files':[{'name':name,'bytes':len(body),'sha256':hashlib.sha256(body).hexdigest()} for name,body in files.items()]}
    return manifest,files


def describe_run(root:Path,key:str)->dict:
    return _prepare(root,key)[0]


def build_bundle(root:Path,key:str)->bytes:
    manifest,files=_prepare(root,key)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,body in files.items(): archive.writestr(name,body)
        archive.writestr('manifest.json',_encode(manifest))
    return out.getvalue()
