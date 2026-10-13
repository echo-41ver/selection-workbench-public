"""Validate the reader-facing TEMU keyword-union workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--book", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    args = parser.parse_args()
    expected = int(json.loads(args.summary.read_text(encoding="utf-8"))["unique_goods_id"])
    book = load_workbook(args.book, read_only=False, data_only=False)
    assert book.sheetnames == ["商品明细", "汇总与审计"], book.sheetnames
    sheet = book["商品明细"]
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref == f"A4:L{expected + 4}"
    assert sheet.max_row - 4 == expected
    ids: set[str] = set()
    links = 0
    for row in sheet.iter_rows(min_row=5, values_only=True):
        goods_id = str(row[1] or "")
        assert goods_id and goods_id not in ids
        ids.add(goods_id)
        for value in (row[5], row[6]):
            text = str(value or "")
            assert text == "链接缺失" or text.startswith('=HYPERLINK("https://')
            links += text.startswith('=HYPERLINK("https://')
    print(json.dumps({"状态": "验证通过", "商品数": len(ids), "公式链接": links}, ensure_ascii=False))


if __name__ == "__main__":
    main()
