"""Exercise the installed browser, image model and Excel writer without paid analysis."""
from __future__ import annotations

from contextlib import contextmanager
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import ProxyHandler, Request, build_opener
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from selection_automation.runtime_environment import activate
from selection_automation.merge_temu_snapshots import atomic_write
from selection_automation.adspower_diagnostics import AdsPowerAPIError


# Only standard WebDriver identifiers may leave a local driver's error body.
WEBDRIVER_ERRORS = frozenset({
    'detached shadow root', 'element click intercepted', 'element not interactable',
    'insecure certificate', 'invalid argument', 'invalid cookie domain',
    'invalid element state', 'invalid selector', 'invalid session id', 'javascript error',
    'move target out of bounds', 'no such alert', 'no such cookie', 'no such element',
    'no such frame', 'no such shadow root', 'no such window', 'script timeout',
    'session not created', 'stale element reference', 'timeout', 'unable to capture screen',
    'unable to set cookie', 'unexpected alert open', 'unknown command', 'unknown error',
    'unknown method', 'unsupported operation',
})


class WebDriverSmokeError(RuntimeError):
    def __init__(self, command, http_status, error):
        self.diagnostic = {'command': command, 'http_status': http_status, 'error': error}
        super().__init__('Local WebDriver command rejected')


def _webdriver_request(transport, address, suffix, body):
    command = {'/url': 'NAVIGATE', '/execute/sync': 'DOM_READ'}[suffix]
    call = Request(address + suffix, data=json.dumps(body).encode('utf-8'),
        headers={'Content-Type': 'application/json'}, method='POST')
    try:
        with transport.open(call, timeout=15) as response:
            return json.load(response).get('value')
    except HTTPError as error:
        identifier = 'UNKNOWN_DRIVER_ERROR'
        try:
            payload = json.loads(error.read(16 * 1024))
            value = payload.get('value') if isinstance(payload, dict) else None
            candidate = value.get('error') if isinstance(value, dict) else None
            if isinstance(candidate, str) and candidate in WEBDRIVER_ERRORS:
                identifier = candidate
        except (OSError, ValueError, TypeError):
            pass
        finally:
            error.close()
        status = error.code if isinstance(error.code, int) and 100 <= error.code <= 599 else None
        raise WebDriverSmokeError(command, status, identifier) from None


@contextmanager
def _local_marker_page(marker):
    route = '/' + uuid.uuid4().hex
    page = ('<!doctype html><meta charset="utf-8"><title>Deployment</title>'
            '<div id="proof">' + marker + '</div>').encode('utf-8')

    class Handler(BaseHTTPRequestHandler):
        def setup(self):
            super().setup()
            self.connection.settimeout(5)

        def do_GET(self):
            found = self.path == route
            payload = page if found else b'Not found'
            self.send_response(200 if found else 404)
            self.send_header('Content-Type', 'text/html; charset=utf-8' if found else 'text/plain')
            self.send_header('Content-Length', str(len(payload)))
            self.send_header('Cache-Control', 'no-store')
            self.end_headers()
            self.wfile.write(payload)

        def log_message(self, format, *args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    server.daemon_threads = True
    server.block_on_close = False
    worker = threading.Thread(target=lambda: server.serve_forever(poll_interval=0.05),
                              name='deployment-smoke-http', daemon=True)
    started = False
    try:
        worker.start()
        started = True
        yield 'http://127.0.0.1:' + str(server.server_address[1]) + route
    finally:
        original_error = sys.exc_info()[1]
        cleanup_failed = False
        stopper = None
        try:
            if started:
                stopper = threading.Thread(target=server.shutdown, name='deployment-smoke-http-stop', daemon=True)
                stopper.start()
                stopper.join(timeout=2)
        except Exception:
            cleanup_failed = True
        finally:
            try:
                server.server_close()
            except OSError:
                cleanup_failed = True
        if started:
            worker.join(timeout=2)
            cleanup_failed |= worker.is_alive() or stopper is None or stopper.is_alive()
        if cleanup_failed:
            if original_error is None:
                raise RuntimeError('LOCAL_SMOKE_SERVER_SHUTDOWN_FAILED')
            original_error.add_note('LOCAL_SMOKE_SERVER_SHUTDOWN_FAILED')


def _safe_adspower_diagnostic(error):
    if isinstance(error, AdsPowerAPIError):
        return {key: error.diagnostic[key] for key in
                ('operation', 'category', 'provider_code', 'http_status', 'retryable')
                if key in error.diagnostic}
    return None


class FlowerCleanupError(RuntimeError):
    def __init__(self, diagnostic):
        self.flower_cleanup = diagnostic
        super().__init__('DEPLOYMENT_PROFILE_CLEANUP_INCOMPLETE')


def _release_smoke_driver(folder, profile, driver, port, psutil):
    try:
        record = json.loads((folder / 'driver.json').read_text('utf-8-sig'))
    except FileNotFoundError:
        if driver is not None and port is not None:
            raise RuntimeError('DEPLOYMENT_DRIVER_OWNERSHIP_UNCONFIRMED') from None
        return 'NOT_RECORDED'
    if (str(record.get('profile_id')) != str(profile) or record.get('port') != port or not driver
            or type(record.get('pid')) is not int or record['pid'] < 1
            or type(record.get('started_ticks')) is not int):
        raise RuntimeError('DEPLOYMENT_DRIVER_OWNERSHIP_UNCONFIRMED')
    try:
        process = psutil.Process(record['pid'])
        ticks = round(process.create_time() * 10_000_000) + 621355968000000000
        if (abs(ticks - record['started_ticks']) > 100_000
                or Path(process.exe()).resolve() != Path(driver).resolve()
                or Path(driver).name.lower() != 'geckodriver.exe'):
            raise RuntimeError('DEPLOYMENT_DRIVER_OWNERSHIP_UNCONFIRMED')
        process.terminate()
        process.wait(timeout=5)
    except psutil.NoSuchProcess:
        return 'ALREADY_EXITED'
    return 'RELEASED'


def _readback_smoke_profile(base, profile, api, diagnostic):
    diagnostic['readback_api_calls'] += 1
    try:
        response = api(base, '/api/v1/user/list?' + urlencode(
            {'user_id': profile, 'page': 1, 'page_size': 100}), timeout=12)
        items = (response.get('data') or {}).get('list')
        state = 'UNKNOWN'
        if isinstance(items, list):
            if not items:
                state = 'ABSENT'
            elif all(isinstance(item, dict) and
                     {str(item[key]) for key in ('user_id', 'profile_id', 'id') if key in item}
                     == {str(profile)} for item in items):
                state = 'PRESENT'
        diagnostic['events'].append({'phase': 'DELETE_READBACK', 'state': state})
        return state
    except Exception as error:
        event = {'phase': 'DELETE_READBACK', 'state': 'UNKNOWN', 'error_type': type(error).__name__}
        safe = _safe_adspower_diagnostic(error)
        if safe:
            event['adspower_diagnostic'] = safe
        diagnostic['events'].append(event)
        raise


def _delete_smoke_profile(base, profile, api, diagnostic):
    for attempt in range(1, 4):
        diagnostic['delete_api_calls'] = attempt
        try:
            api(base, '/api/v2/browser-profile/delete', method='POST', body={'profile_id': [profile]}, timeout=15)
        except AdsPowerAPIError as error:
            diagnostic['events'].append({'phase': 'DELETE_PROFILE', 'attempt': attempt,
                                          'adspower_diagnostic': _safe_adspower_diagnostic(error)})
            category = error.diagnostic.get('category')
            if category == 'CONNECTIVITY':
                # Reconcile this exact id before any idempotent retry. An
                # unfiltered, malformed or ambiguous listing proves nothing.
                try:
                    state = _readback_smoke_profile(base, profile, api, diagnostic)
                except Exception:
                    raise error from None
                if state == 'ABSENT':
                    return 'READBACK_ABSENT'
                if state != 'PRESENT':
                    raise
            elif category != 'PROFILE_IN_USE':
                raise
            if attempt == 3:
                raise
            time.sleep(1)
        else:
            diagnostic['events'].append({'phase': 'DELETE_ACKNOWLEDGED', 'attempt': attempt})
            if _readback_smoke_profile(base, profile, api, diagnostic) != 'ABSENT':
                raise RuntimeError('DEPLOYMENT_PROFILE_DELETE_UNCONFIRMED')
            return 'READBACK_ABSENT'


def _cleanup_flower_profile(folder, profile, driver, port, base, api, stop_and_wait, psutil):
    started = time.monotonic()
    diagnostic = {'ready': False, 'events': [], 'delete_api_calls': 0, 'readback_api_calls': 0,
                  'max_delete_api_calls': 3, 'stop_wait_timeout_seconds': 25}
    phase = 'RELEASE_OWNED_DRIVER'
    try:
        outcome = _release_smoke_driver(folder, profile, driver, port, psutil)
        diagnostic['events'].append({'phase': phase, 'outcome': outcome})
        phase = 'STOP_AND_WAIT'
        stop_and_wait(base, profile, api=api, timeout=25)
        diagnostic['events'].append({'phase': phase, 'api_inactive': True, 'profile_process_exited': True})
        phase = 'DELETE_PROFILE'
        confirmation = _delete_smoke_profile(base, profile, api, diagnostic)
        diagnostic['events'].append({'phase': phase, 'confirmation': confirmation})
        diagnostic['ready'] = True
    except Exception as error:
        diagnostic.update(failed_phase=phase, error_type=type(error).__name__)
        safe = _safe_adspower_diagnostic(error)
        if safe:
            diagnostic['adspower_diagnostic'] = safe
    diagnostic['elapsed_seconds'] = round(time.monotonic() - started, 3)
    return diagnostic


def installation_token(root):
    root = Path(root).resolve()
    inputs = [str(root).casefold(), os.environ.get('COMPUTERNAME', ''), os.environ.get('USERNAME', ''), sys.version.split()[0]]
    for name in ('BUILD_INFO.json', 'requirements.lock.txt', 'package-lock.json'):
        path = root / name
        if path.is_file():
            inputs.append(hashlib.sha256(path.read_bytes()).hexdigest())
    return hashlib.sha256('\n'.join(inputs).encode('utf-8')).hexdigest()


def verified(root):
    try:
        report = json.loads((Path(root) / 'runtime/deployment-smoke.json').read_text('utf-8'))
        return report.get('ready') is True and report.get('installation') == installation_token(root)
    except (OSError, ValueError):
        return False


def check_model(root):
    import torch
    from torchvision.models import resnet50
    model_path = Path(root) / 'runtime/models/torch/hub/checkpoints/resnet50-11ad3fa6.pth'
    if not model_path.is_file():
        raise RuntimeError('IMAGE_MODEL_FILE_MISSING')
    torch.set_num_threads(2)
    model = resnet50(weights=None).eval()
    model.load_state_dict(torch.load(model_path, map_location='cpu', weights_only=True))
    with torch.inference_mode():
        result = model(torch.zeros(1, 3, 64, 64))
    if tuple(result.shape) != (1, 1000) or not torch.isfinite(result).all().item():
        raise RuntimeError('IMAGE_MODEL_INFERENCE_FAILED')
    return {'ready': True, 'device': 'cpu', 'output_shape': [1, 1000]}


def check_excel(folder):
    from PIL import Image
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as ExcelImage
    photo = Path(folder) / 'image.png'
    Image.new('RGB', (32, 32), (30, 90, 150)).save(photo)
    path = Path(folder) / '部署验证.xlsx'
    book = Workbook()
    book.active['A1'] = '部署验证'
    book.active['B1'] = 12.5
    book.active.add_image(ExcelImage(photo), 'C1')
    book.save(path)
    reopened = load_workbook(path)
    ready = reopened.active['A1'].value == '部署验证' and reopened.active['B1'].value == 12.5 and len(reopened.active._images) == 1
    reopened.close()
    if not ready:
        raise RuntimeError('EXCEL_ROUNDTRIP_FAILED')
    return {'ready': True, 'chinese_text': True, 'embedded_image': True}


def check_chromium(root):
    script = r"""const {chromium}=require('playwright');
(async()=>{let browser;try{browser=await chromium.launch({headless:true});
const page=await browser.newPage();await page.setContent('<!doctype html><meta charset="utf-8"><div id="proof">Deployment OK</div>');
if(await page.locator('#proof').innerText()!=='Deployment OK')throw Error('DOM_READ_FAILED');
console.log(JSON.stringify({ready:true,version:browser.version(),dom_read:true}));
}finally{if(browser)await browser.close();}})().catch(()=>process.exit(1));"""
    result = subprocess.run([shutil.which('node') or 'node', '-e', script], cwd=root,
        capture_output=True, text=True, encoding='utf-8', timeout=60, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    if result.returncode:
        raise RuntimeError('CHROMIUM_LAUNCH_FAILED')
    return json.loads(result.stdout)


def check_flower(root):
    from company_setup import ensure_adspower_api, free_block
    from replace_proxyshare_worker_environment import api_call, stop_profile_and_wait
    from selection_automation.secrets import load
    from selection_automation.adspower_kernel import REQUIRED_KERNEL
    from selection_automation.collection_environment import profile_browser_settings
    from selection_automation.background_launch import run_background_launcher
    import psutil
    root = Path(root)
    company = load(root / 'runtime/secrets/company.dpapi')
    ensure_adspower_api(company)
    base = company.get('adspower_api_url', 'http://127.0.0.1:50325')
    profile = None
    worker = 'deployment-smoke-' + uuid.uuid4().hex[:12]
    state_folder = root / 'runtime' / worker
    result = None
    driver = None
    port = None
    try:
        created = api_call(base, '/api/v1/user/create', method='POST', body={
            'name': worker, 'group_id': '0', 'user_proxy_config': {'proxy_soft': 'no_proxy'},
            **profile_browser_settings()}, timeout=30)
        profile = (created.get('data') or {}).get('id')
        if not profile:
            raise RuntimeError('BROWSER_PROFILE_CREATE_UNCONFIRMED')
        opened = api_call(base, '/api/v1/browser/start?' + urlencode({'user_id': profile,
            'ip_tab': '0', 'proxy_detection': '0', 'last_opened_tabs': '0', 'headless': '0'}), timeout=90).get('data') or {}
        driver = opened.get('webdriver')
        if not driver or not Path(driver).is_file():
            raise RuntimeError('FIREFOX_WEBDRIVER_MISSING')
        port = free_block(41000, count=1)[0]
        binding = run_background_launcher([shutil.which('pwsh') or 'powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass',
            '-File', str(root / 'scripts/bind_flower_worker.ps1'), '-WorkerName', worker,
            '-ProfileId', str(profile), '-WebDriverPort', str(port), '-GeckoDriverPath', str(driver), '-WaitSeconds', '120'],
            cwd=root, timeout=260)
        if binding.returncode:
            raise RuntimeError('FIREFOX_BINDING_FAILED')
        sid = (state_folder / 'flower_session_id.txt').read_text('utf-8').strip()
        transport = build_opener(ProxyHandler({}))
        address = f'http://127.0.0.1:{port}/session/{sid}'
        def request(suffix, body):
            return _webdriver_request(transport, address, suffix, body)
        marker = uuid.uuid4().hex
        with _local_marker_page(marker) as page_url:
            request('/url', {'url': page_url})
            deadline=time.monotonic()+15
            observed=None
            while time.monotonic()<deadline:
                observed = request('/execute/sync', {'script': "return document.getElementById('proof')?.textContent", 'args': []})
                if observed==marker:break
                time.sleep(.25)
            if observed != marker:
                raise RuntimeError('FIREFOX_DOM_READ_FAILED')
        record = json.loads((state_folder / 'binding.json').read_text('utf-8-sig'))
        result = {'ready': True, 'kernel': dict(REQUIRED_KERNEL), 'webdriver': True,
            'bidi': bool((state_folder / 'flower_bidi_url.txt').read_text('utf-8').strip()),
            'foreground': record.get('foreground', {}).get('activated') is True,
            'window_visible': record.get('foreground', {}).get('visible') is True, 'dom_read': True}
        if not result['bidi'] or not result['window_visible']:
            raise RuntimeError('FIREFOX_BINDING_PROOF_INCOMPLETE')
    finally:
        if profile:
            original_error = sys.exc_info()[1]
            cleanup = _cleanup_flower_profile(state_folder, profile, driver, port, base,
                                             api_call, stop_profile_and_wait, psutil)
            if original_error is not None:
                original_error.flower_cleanup = cleanup
            elif not cleanup['ready']:
                raise FlowerCleanupError(cleanup)
            elif result is not None:
                result['cleanup'] = cleanup
    return result


def run(root=ROOT):
    root = Path(root).resolve()
    activate(root)
    report = {'schema': 1, 'ready': False, 'installation': installation_token(root), 'checked_at': time.time(),
              'paid_analysis_requests': 0, 'checks': {}}
    destination = root / 'runtime/deployment-smoke.json'
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = 'image_model'
    try:
        report['checks'][stage] = check_model(root)
        with tempfile.TemporaryDirectory(prefix='deployment-', dir=destination.parent) as folder:
            stage = 'excel'
            report['checks'][stage] = check_excel(folder)
        stage = 'chromium'
        report['checks'][stage] = check_chromium(root)
        stage = 'flower_binding'
        report['checks'][stage] = check_flower(root)
        report['ready'] = True
    except Exception as error:
        report.update(failed_stage=stage, error_type=type(error).__name__)
        if 'LOCAL_SMOKE_SERVER_SHUTDOWN_FAILED' in getattr(error, '__notes__', ()):
            report['cleanup_error_code'] = 'LOCAL_SMOKE_SERVER_SHUTDOWN_FAILED'
        if isinstance(error, WebDriverSmokeError):
            report['webdriver_diagnostic'] = {key: error.diagnostic[key] for key in
                ('command', 'http_status', 'error')}
        if hasattr(error, 'flower_cleanup'):
            report['flower_cleanup'] = error.flower_cleanup
        safe = _safe_adspower_diagnostic(error)
        if safe:
            report['adspower_diagnostic'] = safe
        if isinstance(error, RuntimeError) and str(error).isupper() and str(error).replace('_', '').isalnum():
            report['error_code'] = str(error)
    atomic_write(destination, json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    report = run()
    print(json.dumps(report, ensure_ascii=False))
    raise SystemExit(0 if report['ready'] else 1)
