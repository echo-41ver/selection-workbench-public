"""Portable exports. Data is frozen by the existing SQLite snapshot command."""
from __future__ import annotations
import argparse
import json
import os
from pathlib import Path
import sys
from collections import Counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from selection_automation.category_config import load_category
from selection_automation.rules import score
from selection_automation.vision import compatible_vision_keys, VisionRevisionContext
from selection_automation.stage3_admission import admitted, threshold as admission_threshold
from selection_workflow.excel_delivery import data_workbook_path, publish_view_copies
from selection_workflow.excel_images import cache_amazon_images, cached_thumbnail
from selection_workflow.commercial_excel import add_report
from selection_workflow.excel_verdict_style import LEGEND, verdict_formats, verdict_kind
from selection_automation.selection_delivery import display_parameter


from selection_automation.jsonl_records import read_jsonl as jsonl


def _write_link(sheet,row,col,url,fmt,label):
    """Formula links do not consume worksheet hyperlink relationships."""
    if len(url)>32767:
        raise ValueError('URL exceeds Excel cell capacity; previous workbook preserved')
    if len(url)>255:
        sheet.write_string(row,col,url,fmt)
        return
    destination=url.replace('"','""');display=label.replace('"','""')
    sheet.write_formula(row,col,f'=HYPERLINK("{destination}","{display}")',fmt,label)


def write_book(destination, headers, records, summary, low_asins=None):
    import xlsxwriter
    from openpyxl import load_workbook
    destination.parent.mkdir(parents=True, exist_ok=True)
    data = data_workbook_path(destination)
    temp = data.with_name(f'.{data.stem}.{os.getpid()}.tmp.xlsx')
    book = xlsxwriter.Workbook(temp, {'constant_memory':True, 'strings_to_urls':False, 'strings_to_formulas':False})
    fmt = book.add_format({'bold':True, 'bg_color':'#24476B','font_color':'white','text_wrap':True})
    sm = book.add_worksheet('汇总'); sm.set_column(0,0,32); sm.set_column(1,1,90)
    for i,(k,v) in enumerate(summary.items()):
        sm.write(i,0,k); sm.write(i,1,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v)
    def sheet(name):
        s = book.add_worksheet(name); s.write_row(0,0,headers,fmt); s.freeze_panes(1,0)
        s.set_column(0,len(headers)-1,22); return s
    ws, lw = sheet('明细_1'), None
    count, low_count = 0, 0
    link = book.add_format({'font_color':'blue','underline':1})
    for record in records:
        if count and count % 1048575 == 0:
            ws.autofilter(0,0,1048575,len(headers)-1); ws = sheet(f'明细_{count//1048575+1}')
        def emit(target,index):
            for col,value in enumerate(record):
                if isinstance(value,(list,dict)): value=json.dumps(value,ensure_ascii=False)
                if isinstance(value,str) and value.startswith('https://') and ('链接' in headers[col] or '主图' in headers[col]):
                    _write_link(target,index,col,value,link,'查看')
                else: target.write(index,col,value[:32760] if isinstance(value,str) else value)
        emit(ws,count%1048575+1); count+=1
        if low_asins is not None and record[0] in low_asins:
            if low_count%1048575 == 0: lw=sheet(f'已核验采样同款不超过20_{low_count//1048575+1}')
            emit(lw,low_count%1048575+1); low_count+=1
    ws.autofilter(0,0,(count-1)%1048575+1 if count else 1,len(headers)-1)
    book.close()
    try:
        check=load_workbook(temp,read_only=True,data_only=False)
        actual=sum(s.max_row-1 for s in check if s.title.startswith('明细_'))
        check.close()
        if actual!=count: raise ValueError('Excel reopen row count mismatch')
        if '商品对总数' in summary and count!=summary['商品对总数']: raise ValueError('snapshot row count mismatch')
        os.replace(temp,data)
        return publish_view_copies(data,destination,keep_snapshots=3)
    finally:
        temp.unlink(missing_ok=True)


def _formats(book):
    common = {'font_name': 'Arial', 'font_size': 10, 'valign': 'vcenter'}
    return {
        'title': book.add_format({**common, 'font_size': 15, 'bold': True, 'font_color': 'white', 'bg_color': '#1F4E78', 'align': 'center'}),
        'note': book.add_format({**common, 'italic': True, 'text_wrap': True, 'bg_color': '#D9EAF7'}),
        'header': book.add_format({**common, 'bold': True, 'font_color': 'white', 'bg_color': '#4472C4', 'align': 'center', 'text_wrap': True}),
        'text': book.add_format({**common, 'text_wrap': True, 'valign': 'top'}),
        'center': book.add_format({**common, 'align': 'center'}),
        'money': book.add_format({**common, 'num_format': '$0.00'}),
        'score': book.add_format({**common, 'num_format': '0.00', 'align': 'center'}),
        'link': book.add_format({**common, 'font_color': 'blue', 'underline': 1, 'align': 'center'}),
        'hard': book.add_format({**common, 'text_wrap': True, 'bg_color': '#FCE4D6'}),
        'pending': book.add_format({**common, 'text_wrap': True, 'bg_color': '#FFF2CC'}),
        'candidate': book.add_format({**common, 'text_wrap': True, 'bg_color': '#E2F0D9'}),
        'missing': book.add_format({**common, 'text_wrap': True, 'bg_color': '#F4CCCC'}),
    }


def _display_details(value, fallback='未展示'):
    return display_parameter(value) or fallback


def _write_value(sheet, row, col, value, fmt, formats, header, *, native_links=False):
    from urllib.parse import urlparse
    if isinstance(value, (list, dict)) or header in (
            '页面月销量文本','销量文本','展示销量原文','来源','上架时间来源','销量口径','排序确认',
            '前台库存提示','原始榜单卡片文本'):
        value = display_parameter(value)
    if isinstance(value, str):
        parsed = urlparse(value)
        if parsed.scheme == 'https' and parsed.netloc and ('链接' in header or 'URL' in header or header in ('Amazon主图', 'TEMU主图')):
            if native_links and len(value) > 255:
                sheet.write_url(row, col, value, formats['link'], '打开商品' if '商品链接' in header else '查看图片')
                return
            _write_link(sheet,row, col, value, formats['link'], '打开商品' if '商品链接' in header else '查看图片')
            return
    sheet.write(row, col, value[:32760] if isinstance(value,str) else value, fmt)


def _publish_formatted(destination, build):
    import xlsxwriter
    from openpyxl import load_workbook

    destination.parent.mkdir(parents=True, exist_ok=True)
    data = data_workbook_path(destination)
    temporary = data.with_name(f'.{data.stem}.{os.getpid()}.tmp.xlsx')
    book = xlsxwriter.Workbook(temporary, {'constant_memory': True, 'strings_to_urls': False, 'strings_to_formulas': False})
    try:
        expected = build(book)
        book.close()
        check = load_workbook(temporary, read_only=True)
        try:
            sheet_names, header_row, count = expected
            if isinstance(sheet_names, str): sheet_names = [sheet_names]
            actual = sum(check[name].max_row - header_row for name in sheet_names)
            if actual != count:
                raise ValueError(f'Excel行数与数据不一致: {actual} != {count}')
        finally:
            check.close()
        os.replace(temporary, data)
        return publish_view_copies(data, destination, keep_snapshots=3)
    finally:
        if not book.fileclosed:
            book.close()
        temporary.unlink(missing_ok=True)


def export_amazon_book(destination, source, category_name, target_count=None, report=None):
    records = sorted(source['records'], key=lambda row: int(row.get('rank') or 0))
    target_count = int(target_count or len(records))
    images = cache_amazon_images(records, destination.parent.parent / 'amazon' / 'excel_image_cache')
    headers = ['序号', '榜单排名', '榜单页', 'ASIN', 'Amazon标题', 'Amazon主图', 'Amazon商品链接',
               'Amazon主图原链接', '榜单缩略图链接', '主图来源', '图片尺寸提示', '图片状态',
               '当前价格(USD)', '评分', '评论数', '类目', 'Best Seller证据', 'Amazon变体/ASIN',
               'Amazon规格', '品牌/卖家', '详情页BSR', '前台库存提示', '详情页字段状态',
               '采集时间(UTC)', '原始榜单卡片文本', '榜单来源URL']

    headers += ['规格最低价(USD)','规格最高价(USD)','价格区间口径','页面月销量文本','销量口径','真实差评原文','差评来源链接']
    def build(book):
        f = _formats(book)
        sheet = book.add_worksheet('商品列表'); sheet.hide_gridlines(2)
        notes = book.add_worksheet('采集说明'); notes.hide_gridlines(2)
        sheet.merge_range(0, 0, 0, len(headers)-1, f'Amazon US｜{category_name} Best Sellers｜目标前{target_count}名，已采{len(records)}件', f['title'])
        sheet.merge_range(1, 0, 1, len(headers)-1,
                          f"来源：{source.get('source_url') or '未记录'}；采集时间：{source.get('observed_at') or '未记录'}；榜单每页通常50件，以实际排名为准。", f['note'])
        sheet.set_row(0, 28); sheet.set_row(1, 28); sheet.set_row(3, 34)
        sheet.write_row(3, 0, headers, f['header'])
        widths = [7, 10, 8, 15, 52, 20, 18, 24, 24, 18, 15, 20, 16, 10, 12, 30, 22, 42, 48, 32, 36, 26, 28, 24, 58, 42]
        for col, width in enumerate(widths): sheet.set_column(col, col, width)
        sheet.set_column(26,32,25)
        sheet.freeze_panes(4, 4)
        for index, (item, image) in enumerate(zip(records, images), 4):
            rank = item.get('rank')
            values = [index-3, f'#{rank}' if rank is not None else '', item.get('page_no') or (int(rank)-1)//50+1 if rank else None,
                      item.get('asin'), item.get('title'), None, item.get('product_url'), item.get('main_image_url'),
                      item.get('listing_image_url'), item.get('image_source'),
                      f"{item['image_width_hint']}×{item['image_height_hint']}" if item.get('image_width_hint') and item.get('image_height_hint') else None,
                      item.get('image_status'), item.get('price_usd'), item.get('rating'), item.get('review_count'),
                      item.get('category') or category_name, f'榜单排名 #{rank}' if rank else None,
                      _display_details(item.get('variants'), '详情页未展示可选变体'),
                      _display_details(item.get('specifications'), '详情页未展示规格'),
                      '\n'.join(s for s in [f"品牌：{item['brand']}" if item.get('brand') else '', f"卖家：{item['seller']}" if item.get('seller') else ''] if s),
                      item.get('best_sellers_rank'), item.get('availability'), item.get('detail_status'),
                      item.get('observed_at') or source.get('observed_at'), item.get('raw_card_text'), item.get('source_url') or source.get('source_url')]
            from selection_automation.operations import product_card
            card=product_card(item,'amazon');price=card['price']
            values += [price['min'],price['max'],price['basis'],card['sales']['text'],card['sales']['basis'],
                       '\n'.join(r['text'] for r in card['negative_reviews']) or '尚未采集真实差评',
                       '\n'.join(r['source_url'] for r in card['negative_reviews'])]
            sheet.set_row(index, 92)
            for col, value in enumerate(values):
                _write_value(sheet, index, col, value, f['money'] if col == 12 else f['text'], f, headers[col])
            if image:
                from PIL import Image
                with Image.open(image) as thumbnail:
                    scale = min(112 / thumbnail.width, 112 / thumbnail.height)
                sheet.insert_image(index, 5, str(image), {'x_scale': scale, 'y_scale': scale,
                                   'x_offset': 4, 'y_offset': 4, 'object_position': 1})
            else:
                sheet.write(index, 5, '图片获取失败；原链接见右侧', f['missing'])
        sheet.autofilter(3, 0, max(4, len(records)+3), len(headers)-1)
        notes.merge_range('A1:D1', '采集口径与字段边界', f['title'])
        notes.write_row(2, 0, ['项目', '内容', '状态', '来源/说明'], f['header'])
        entries = [('榜单范围', f'目标前{target_count}名；当前已采{len(records)}件', '完整' if len(records) == target_count and source.get('complete') else '进行中/待核对', source.get('source_url')),
                   ('榜单分页', '按实际榜单排名及页码记录；通常每页50件', '已记录', source.get('source_url')),
                   ('商品主键', 'ASIN', '已记录', 'Amazon榜单及详情页'),
                   ('主图内嵌', f'{sum(bool(p) for p in images)}/{len(records)}件', '完成' if all(images) else '部分失败', '图片固定在商品行；原链接保留供追溯'),
                   ('详情页字段', '品牌、卖家、变体、规格、BSR、库存提示', '以实际采集为准', '未展示字段不推断'),
                   ('月销量', '不在榜单前台明确提供时不填', '未推断', '不使用估算值')]
        for row, entry in enumerate(entries, 3):
            for col, value in enumerate(entry): _write_value(notes, row, col, value, f['text'], f, '来源/说明' if col == 3 else '')
        for col, width in enumerate([18, 48, 16, 85]): notes.set_column(col, col, width)
        notes.freeze_panes(3, 0)
        add_report(book, report)
        return '商品列表', 4, len(records)
    return _publish_formatted(destination, build)


def export_temu_book(destination, rows, category_name, max_per_keyword, union_summary=None, report=None, samples=None):
    union_summary = union_summary or {}
    ids = [str(row.get('goods_id') or '').strip() for row in rows]
    if rows and (any(not identity for identity in ids) or len(ids) != len(set(ids))):
        raise ValueError('TEMU并集明细存在空商品ID或重复商品ID')
    if union_summary.get('unique_goods_id') is not None and int(union_summary['unique_goods_id']) != len(rows):
        raise ValueError('TEMU并集汇总与明细行数不一致')
    headers = ['序号', 'TEMU商品ID', '标题', '价格(USD)', '销量文本', '商品链接', '主图链接',
               '全部图片链接', '命中搜索词', '首次发现轮次', '观测时间(UTC)', '来源']
    headers += ['规格最低价(USD)','规格最高价(USD)','价格区间口径','商品评分','上架时间','上架时间来源','销量口径','真实差评原文','差评来源链接']
    def build(book):
        f = _formats(book)
        sheet = book.add_worksheet('商品明细'); sheet.hide_gridlines(2)
        audit = book.add_worksheet('汇总与审计'); audit.hide_gridlines(2)
        sheet.merge_range(0, 0, 0, len(headers)-1, f'TEMU 美国站{category_name}商品明细（网页优先，极鲸云补充）', f['title'])
        sheet.set_row(0, 28); sheet.set_row(3, 34)
        sheet.write_row(3, 0, headers, f['header'])
        for col, width in enumerate([9, 20, 60, 14, 18, 18, 18, 62, 48, 16, 25, 30]): sheet.set_column(col, col, width)
        sheet.set_column(12,20,25)
        for index, item in enumerate(rows, 4):
            values = [index-3, str(item.get('goods_id') or ''), item.get('title'), item.get('price_usd'),
                      item.get('sales_text'), item.get('product_url'), item.get('main_image_url'),
                      item.get('all_image_urls'), '；'.join(item.get('source_queries') or [str(item.get('search_query') or '')]),
                      item.get('first_seen_round'), item.get('observed_at'), item.get('source') or 'VISIBLE_TEMU_FRONTEND']
            from selection_automation.operations import product_card
            card=product_card(item,'temu');price=card['price']
            values += [price['min'],price['max'],price['basis'],card['rating'],card['listed_at'],card['listed_at_source'],card['sales']['basis'],
                       '\n'.join(r['text'] for r in card['negative_reviews']) or '尚未采集真实差评',
                       '\n'.join(r['source_url'] for r in card['negative_reviews'])]
            for col, value in enumerate(values): _write_value(sheet, index, col, value, f['money'] if col == 3 else f['text'], f, headers[col])
        sheet.freeze_panes(4, 2); sheet.autofilter(3, 0, max(4, len(rows)+3), len(headers)-1)
        audit.merge_range('A1:B1', 'TEMU关键词采集汇总与goods_id审计', f['title'])
        audit.set_column('A:A', 38); audit.set_column('B:B', 90)
        entries = [('数据来源', 'TEMU网页采集优先；极鲸云只补充网页未采到的商品ID，来源见原有来源列'),
                   ('网页商品数',union_summary.get('frontend_unique_goods',len(rows))),
                   ('极鲸云补充数',union_summary.get('provider_supplement_goods',0)),
                   ('源快照数量', len(union_summary.get('input_snapshots') or [])),
                   ('源记录总数', union_summary.get('source_rows') or len(rows)),
                   ('唯一goods_id数', len(rows)), ('每关键词每种排序上限', max_per_keyword),
                   ('采集顺序', '每词先销量最高，再最近上架；两轮各自达到上限后继续下一词；汇总去重'),
                   ('合并重复记录数', union_summary.get('duplicate_rows_merged') or 0),
                   ('有商品链接', sum(bool(r.get('product_url')) for r in rows)),
                   ('有主图链接', sum(bool(r.get('main_image_url')) for r in rows)),
                   ('内容SHA-256', union_summary.get('content_sha256') or '未记录'),
                   ('去重口径', '按goods_id去重；同一商品命中多个关键词时合并关键词')]
        for index, (key, value) in enumerate(entries, 1):
            audit.write(index, 0, key, f['header']); audit.write(index, 1, value, f['text'])
        for order, label in [('sales', '销量最高样本'), ('recent', '最近上架样本')]:
            detail = book.add_worksheet(label)
            detail.hide_gridlines(2)
            cols = ['关键词', 'TEMU商品ID', '标题', '展示销量原文', '价格(USD)', '商品链接', '观测时间', '排序确认']
            detail.merge_range(0, 0, 0, len(cols)-1,
                label+'：展示销量用于粗略选品参考，周期不明；最近上架排序不等于确认90天内上架，也不是环比。', f['note'])
            detail.write_row(2, 0, cols, f['header'])
            detail.set_column(0, 0, 28); detail.set_column(1, 1, 22); detail.set_column(2, 2, 55)
            detail.set_column(3, 7, 25); detail.freeze_panes(3, 2)
            unique = {}
            for item in samples if samples is not None else rows:
                evidence = item.get('search_sort') or {}
                if evidence.get('requested') == order and evidence.get('state') == 'UI_CONFIRMED':
                    identity = (item.get('search_query'), str(item.get('goods_id') or ''))
                    if identity not in unique or str(item.get('observed_at') or '') > str(unique[identity].get('observed_at') or ''):
                        unique[identity] = item
            for index, item in enumerate(unique.values(), 3):
                values = [item.get('search_query'), str(item.get('goods_id') or ''), item.get('title'),
                          item.get('sales_text'), item.get('price_usd'), item.get('product_url'),
                          item.get('observed_at'), (item.get('search_sort') or {}).get('label')]
                for col, value in enumerate(values): _write_value(detail, index, col, value, f['text'], f, cols[col])
            detail.autofilter(2, 0, max(2, len(unique)+2), len(cols)-1)
        add_report(book, report)
        return '商品明细', 4, len(rows)
    return _publish_formatted(destination, build)


def amazon_conclusions(headers, rows, limit, *, binary=False, same_only=False, operator='lte'):
    """Count the actual exported verdicts, including unresolved Amazon products."""
    final_index = headers.index('最终判断')
    products = {}
    for row in rows:
        asin = str(row[0])
        item = products.setdefault(asin, {'asin': asin, 'title': row[1], 'url': row[2], 'image': row[3], 'review_images': set(),
            'same': 0, 'review': 0, 'different': 0, 'unanalysed': 0, 'pending': 0, 'total': 0})
        kind = verdict_kind(row[final_index])
        item[kind] += 1
        if binary and ('分析失败' in str(row[final_index]) or '第三层待运行' in str(row[final_index])):
            item['pending'] += 1
        if kind == 'same' or (kind == 'review' and not (same_only or binary)):
            item['review_images'].update(filter(None, (row[3], row[7])))
        item['total'] += 1
    for item in products.values():
        item['selected'] = item['same'] < limit if operator=='lt' else item['same'] <= limit
        if not item['selected']:
            item['conclusion'] = '不符合：确认同款超过门槛'
            item['reason'] = f"已确认{item['same']}个同款，不满足本任务{('严格少于' if operator=='lt' else '不超过')}{limit}个的条件"
        elif binary or same_only:
            item['conclusion'] = ('分析未完成：已有同款在门槛内' if item['pending'] else '符合：当前采样同款在门槛内')
            item['reason'] = f"当前采样确认同款{item['same']}个，{('严格少于' if operator=='lt' else '不超过')}{limit}个；未分析或未入围{item['unanalysed']}对，不代表TEMU全站数量"
        elif item['review'] or item['unanalysed']:
            item['conclusion'] = '待复核：当前确认同款在门槛内'
            item['reason'] = f"确认同款{item['same']}个；待复核{item['review']}对、未分析或未入围{item['unanalysed']}对，不把未知当不同款"
        else:
            item['conclusion'] = '符合：采样确认同款在门槛内'
            item['reason'] = f"本次采样已确认{item['same']}个同款，{('严格少于' if operator=='lt' else '不超过')}{limit}个；不代表TEMU全站数量"
    return list(products.values())


def add_operator_sheet(book, snapshot, formats, thumbnails=None):
    sheet = book.add_worksheet('运营候选'); sheet.hide_gridlines(2)
    headers = ['Amazon ASIN','Amazon标题','Amazon图片','Amazon链接','Amazon价格区间 USD','Amazon销量与口径','Amazon评分',
               '已确认同款数','TEMU标题','TEMU图片','TEMU链接','TEMU价格区间 USD','TEMU销量与口径','TEMU评分',
               'TEMU上架时间及来源','Amazon差评及来源','TEMU差评及来源','入选依据']
    sheet.merge_range(0,0,0,len(headers)-1,'运营候选｜一行一个Amazon商品',formats['title'])
    note = (f"同款数严格少于{snapshot.get('same_limit','N')}；最低显示销量{snapshot.get('min_sales',100)}；"
            f"新品窗口{snapshot.get('new_days',90)}天。当前采样范围，非TEMU全站。" + snapshot.get('message',''))
    sheet.merge_range(1,0,1,len(headers)-1,note,formats['note'])
    sheet.set_row(1,44);sheet.write_row(3,0,headers,formats['header']);sheet.freeze_panes(4,2)
    sheet.set_column(0,len(headers)-1,24);sheet.set_column(1,1,38);sheet.set_column(8,8,38)
    sheet.set_column(15,17,46)
    def price(card):
        value=card.get('price',{})
        if value.get('min') is None or value.get('max') is None:
            return '区间未取得；当前标价：'+str(value.get('observed_price') or '未取得')
        return f"{value['min']}–{value['max']}；{value.get('basis','')}"
    def sales(card):
        value=card.get('sales',{});return f"{value.get('text') or '未取得'}；{value.get('basis','')}"
    def reviews(card):
        return '\n'.join(f"{r['rating']}星：{r['text']}\n{r['source_url']}" for r in card.get('negative_reviews',[])) or '尚未采集真实差评'
    products=snapshot.get('products',[])
    for index,item in enumerate(products,4):
        amazon,temu=item['amazon'],item['temu']
        values=[item['asin'],amazon.get('title'),amazon.get('image'),amazon.get('url'),price(amazon),sales(amazon),amazon.get('rating'),
                item['same_count'],temu.get('title'),temu.get('image'),temu.get('url'),price(temu),sales(temu),temu.get('rating'),
                f"{temu.get('listed_at') or '未取得'}；{temu.get('listed_at_source') or '来源未取得'}",reviews(amazon),reviews(temu),item.get('reason')]
        for col,value in enumerate(values):
            image=(thumbnails or {}).get(value) if col in (2,9) else None
            if image:
                from PIL import Image
                with Image.open(image) as source: width,height=source.size
                scale=min(94/width,94/height)
                sheet.insert_image(index,col,str(image),{'x_scale':scale,'y_scale':scale,'object_position':1,
                                                       'url':amazon.get('url') if col==2 else temu.get('url')})
            elif col in (2,3,9,10) and isinstance(value,str) and value.startswith('https://'):
                _write_link(sheet,index,col,value,formats['link'],'查看图片' if col in (2,9) else '查看商品')
            else:sheet.write(index,col,value[:32760] if isinstance(value,str) else value,formats['text'])
        sheet.set_row(index,82)
    sheet.autofilter(3,0,max(4,len(products)+3),len(headers)-1)


def export_stage_book(destination, headers, rows, summary, category_name, *, third_layer=False, low_asins=None, report=None, conclusions=None, thumbnails=None, same_only=False, operator_snapshot=None):
    """Generic-category rows in the same readable layout as the proven layer workbooks."""
    display_headers = ['序号', 'Amazon主图', 'Amazon ASIN', 'Amazon标题', 'TEMU主图', 'TEMU商品ID',
                       'TEMU标题', '第一层文字事实', '图片向量相似度', '第二层状态',
                       'Amazon商品链接', 'TEMU商品链接']
    if third_layer:
        display_headers += headers[11:]
    title = f'{category_name}｜Amazon–TEMU 持续同款核验（第一、二、三层）' if third_layer else f'{category_name}｜Amazon–TEMU 持续同款核验（第一、二层）'
    main_name = '全量综合结果' if third_layer else '第一二层全量核验'
    binary = summary.get('视觉结论口径') == 'visible_core_binary_v1'
    same_only = same_only or binary

    def build(book):
        f = _formats(book)
        def insert_picture(sheet, row, col, url, product_url=None):
            from PIL import Image
            image = (thumbnails or {}).get(url)
            if not image:
                sheet.write(row, col, '图片缺失，待补采', f['missing'])
                return
            with Image.open(image) as source: width, height = source.size
            scale = min(104/width, 104/height)
            sheet.insert_image(row, col, str(image), {'x_scale': scale, 'y_scale': scale,
                'x_offset': 5, 'y_offset': 5, 'object_position': 1, 'url': product_url or url})
            sheet.write_blank(row, col, None, f['text'])
        if conclusions is not None:
            conclusion_products = [item for item in conclusions if item['selected']] if same_only else conclusions
            conclusion = book.add_worksheet('Amazon最终结论'); conclusion.hide_gridlines(2)
            cols = ['Amazon ASIN', 'Amazon标题', 'Amazon主图', '确认同款数', '待复核对数',
                    '不同款对数', '未分析或未入围对数', '商品对总数', '本任务同款上限',
                    '纳入同款及待复核专表', '最终结论', '结论依据', 'Amazon商品链接']
            if binary:
                cols[4] = '分析失败或待运行对数'
            if same_only:
                cols[9] = '纳入同款专表'
            conclusion.merge_range(0, 0, 0, len(cols)-1, f'{category_name}｜按Amazon商品汇总的最终结论', f['title'])
            conclusion.merge_range(1, 0, 1, len(cols)-1, '同款上限沿用采集前设置；按主体款式和核心结构二分类，仅汇总当前采样；未入围、待运行或分析失败不是商品结论。' if binary else '同款上限沿用采集前的任务设置；仅汇总当前采样。待复核、待补证、未分析及未入围不当作已确认不同款。', f['note'])
            conclusion.set_row(0, 28); conclusion.set_row(1, 42); conclusion.set_row(3, 48)
            conclusion.write_row(3, 0, cols, f['header']); conclusion.freeze_panes(4, 1)
            limit = summary.get('采集前设置的同款上限', (report or {}).get('policy', {}).get('same_limit', 20))
            for index, item in enumerate(conclusion_products, 4):
                values = [item['asin'], item['title'], item['url'], item['same'], item['pending'] if binary else item['review'],
                          item['different'], item['unanalysed'], item['total'], limit,
                          '是' if item['selected'] else '否', item['conclusion'], item['reason'], item['url']]
                for col, value in enumerate(values):
                    if col == 2: insert_picture(conclusion, index, col, item['image'], item['url'])
                    else: _write_value(conclusion, index, col, value, f['text'], f, cols[col])
                conclusion.set_row(index, 96)
            for col in range(len(cols)): conclusion.set_column(col, col, 52 if col in (1, 11) else 36 if col == 10 else 22)
            conclusion.autofilter(3, 0, max(4, len(conclusion_products)+3), len(cols)-1)
        audit = book.add_worksheet('汇总与规则' if third_layer else '汇总'); audit.hide_gridlines(2)
        verdicts = verdict_formats(book) if third_layer else {}
        final_index = headers.index('最终判断') if third_layer else None
        row_formats = {kind: {**f, 'link': formats['link']} for kind, formats in verdicts.items()}
        audit.merge_range('A1:B1', title, f['title'])
        audit.set_column('A:A', 38); audit.set_column('B:B', 90)
        for index, (key, value) in enumerate(summary.items(), 1):
            audit.write(index, 0, key, f['header'])
            audit.write(index, 1, json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value, f['text'])
        audit.freeze_panes(1, 0)

        def make_sheet(name, note=None):
            sheet = book.add_worksheet(name); sheet.hide_gridlines(2)
            sheet.merge_range(0, 0, 0, len(display_headers)-1, title, f['title'])
            sheet.merge_range(1, 0, 1, len(display_headers)-1,
                              note or '第一层核对文字硬冲突；第二层图片向量用于筛选候选，最终同款结论以第三层视觉证据为准。'+(LEGEND if third_layer else ''), f['note'])
            sheet.set_row(0, 28); sheet.set_row(1, 34); sheet.set_row(3, 32)
            sheet.write_row(3, 0, display_headers, f['header'])
            widths = [8, 16, 16, 38, 16, 20, 42, 46, 18, 34, 18, 18]
            widths += [20] * max(0, len(display_headers)-len(widths))
            for col, width in enumerate(widths): sheet.set_column(col, col, width)
            sheet.set_default_row(42); sheet.freeze_panes(4, 0)
            return sheet

        def write_one(sheet, index, record, with_images=False):
            status = str(record[10] or '')
            style = f['hard'] if '排除' in status else f['missing'] if '待补充' in status or '失败' in status else f['candidate'] if '进入第三层' in status else f['pending']
            transformed = [index-3, record[3], record[0], record[1], record[7], record[4], record[5],
                           record[8], record[9], record[10], record[2], record[6]]
            if third_layer: transformed += record[11:]
            kind = verdict_kind(record[final_index]) if third_layer else None
            for col, value in enumerate(transformed):
                if with_images and col in (1, 4):
                    insert_picture(sheet, index, col, value, record[2] if col == 1 else record[6])
                    continue
                fmt = (verdicts[kind]['number'] if col == 8 else verdicts[kind]['text']) if third_layer else f['score'] if col == 8 else style
                _write_value(sheet, index, col, value, fmt, row_formats[kind] if third_layer else f, display_headers[col], native_links=with_images)
            if with_images: sheet.set_row(index, 96)

        maximum = 1_048_572
        main_sheets = [make_sheet(main_name)]
        limit = summary.get('采集前设置的同款上限', (report or {}).get('policy', {}).get('same_limit', 20))
        low_name = f'同款不超过{limit}_同款及待复核'
        low_note = f'只列当前确认同款数≤{limit}的Amazon商品；仅保留同款、待复核及待补证行，不含不同款、未分析或第二层未入围行。待复核仍可能改变同款数量。'+LEGEND
        if same_only:
            low_name = f'同款不超过{limit}_同款明细'
            low_note = f'只列当前采样确认同款数≤{limit}的Amazon商品及其已确认同款行；沿用本任务已有判断，待复核及未入围行不纳入本表。'
            if binary:
                low_note = f'只列当前采样同款数≤{limit}的Amazon商品及其同款行；颜色标识及无关未展示部位不阻断；未入围和任务失败不计为不同款。'
        if summary.get('同款数量比较方式')=='lt':
            low_name=low_name.replace('不超过','少于')
            low_note=low_note.replace('≤','<')
        low_sheets = [make_sheet(low_name, low_note)] if low_asins is not None else []
        count = low_count = 0
        for record in rows:
            if count and count % maximum == 0:
                main_sheets.append(make_sheet(f'{main_name}_{len(main_sheets)+1}'))
            write_one(main_sheets[-1], count % maximum + 4, record)
            count += 1
            if low_asins is not None and str(record[0]) in low_asins and verdict_kind(record[final_index]) in (('same',) if same_only else ('same', 'review')):
                if low_count and low_count % maximum == 0:
                    low_sheets.append(make_sheet(f'{low_name}_{len(low_sheets)+1}', low_note))
                write_one(low_sheets[-1], low_count % maximum + 4, record, with_images=thumbnails is not None)
                low_count += 1
        for index, sheet in enumerate(main_sheets):
            amount = min(maximum, max(0, count-index*maximum))
            sheet.autofilter(3, 0, max(4, amount+3), len(display_headers)-1)
        for index, sheet in enumerate(low_sheets):
            amount = min(maximum, max(0, low_count-index*maximum))
            sheet.autofilter(3, 0, max(4, amount+3), len(display_headers)-1)
        if '商品对总数' in summary and count != summary['商品对总数']:
            raise ValueError('商品对行数与快照不一致')
        add_report(book, report)
        if third_layer and operator_snapshot is not None:
            add_operator_sheet(book, operator_snapshot, f, thumbnails)
        return [sheet.name for sheet in main_sheets], 4, count
    return _publish_formatted(destination, build)


def export(category_path, kind):
    cat=load_category(category_path); key=cat['category_key']; name=cat['category_name_zh']; base=ROOT/'runs'/key
    from selection_automation.visual_decision import binary_mode
    binary = binary_mode(cat)
    from selection_automation.commercial import current_report
    report = None if cat.get('execution_mode')=='collection' else current_report(ROOT, cat)
    if kind in ('layer12','layer123'):
        from export_selection_result import export as export_selection
        if report is None:
            raise ValueError('当前任务尚未接入同款分析')
        return export_selection(ROOT,cat,report)
    if kind=='amazon':
        source=json.loads((base/'amazon'/'amazon_frontend_latest.json').read_text(encoding='utf-8'))
        return export_amazon_book(base/'excel'/f'Amazon_US_{name}_Best_Sellers.xlsx', source, name, cat['amazon']['target_count'], report)
    if kind=='temu':
        union_summary_path = base/'temu'/'union'/'summary.json'
        return export_temu_book(base/'excel'/f'TEMU_US_{name}_关键词采集明细.xlsx',
                                list(jsonl(base/'temu'/'union'/'products.jsonl')), name,
                                cat['temu']['max_products_per_keyword'],
                                json.loads(union_summary_path.read_text(encoding='utf-8')) if union_summary_path.exists() else {}, report,
                                samples=[row for path in sorted((base/'temu/workers').glob('*/*/products.jsonl')) for row in jsonl(path)])
    raise ValueError('未知导出类型')


def export_snapshot(category_path, kind):
    """Export an existing historical analysis snapshot without changing current delivery."""
    if kind not in ('layer12', 'layer123'):
        raise ValueError('历史快照仅支持第一二层或三层导出')
    cat=load_category(category_path); key=cat['category_key']; name=cat['category_name_zh']; base=ROOT/'runs'/key
    from selection_automation.visual_decision import binary_mode
    from selection_automation.commercial import current_report
    binary=binary_mode(cat)
    report=None if cat.get('execution_mode')=='collection' else current_report(ROOT,cat)
    if kind=='layer123' and cat.get('stage3',{}).get('scoring_policy')=='shower_curtain_v3_1':
        from export_shower_layer123_excel import export as export_shower_layer123
        return export_shower_layer123(category_path)
    frozen=base/'stage123'/'export_frozen'
    summary=json.loads((frozen/f'layer12_summary_{key}.json').read_text(encoding='utf-8'))
    if summary['actual_pairs']!=summary['amazon_products']*summary['temu_products']: raise ValueError('pair count mismatch')
    expected={}
    image_sources = {}
    queue_counts=Counter()
    vision_context = VisionRevisionContext(cat)
    for row in jsonl(frozen/f'layer3_candidates_{key}.jsonl'):
        queue_counts['paid' if admitted(cat,row) else 'review'] += 1
        for side in ('amazon', 'temu'):
            if row.get(f'{side}_image_url') and row.get(f'{side}_image_path'):
                image_sources[row[f'{side}_image_url']] = row[f'{side}_image_path']
        try: expected[(row['asin'],row['goods_id'])]=set(compatible_vision_keys(cat,row,context=vision_context))
        except OSError: pass
    results={}; failures={}; log_warnings=[]
    for row in jsonl(base/'stage123'/'results.jsonl', allow_incomplete_tail=True, warnings=log_warnings):
        identity=(row.get('asin'),row.get('goods_id'))
        if row.get('vision_revision_key') in expected.get(identity,set()):
            if row.get('api_status')=='success': results[identity]=row
            elif row.get('api_status')=='failed': failures[identity]=row
    modules=cat['stage3']['modules']; labels=cat['stage3'].get('module_labels',{})
    from selection_automation.codex_visual_review import load_reviews, checked_facts, supersedes_api
    codex_reviews = load_reviews(base, expected)
    for pair, review in codex_reviews.items():
        if pair in results and supersedes_api(cat, review):
            results.pop(pair)
    codex_reviews = {pair: row for pair, row in codex_reviews.items() if pair not in results}
    codex_scores = {}
    invalid_codex = set()
    if binary:
        for pair, review in codex_reviews.items():
            try:
                checked_facts(cat, review['visual_facts'], review['goods_id'])
                codex_scores[pair] = review.get('module_result') or score(cat,review['visual_facts'])
            except (KeyError, TypeError, ValueError):
                invalid_codex.add(pair)
    valid_codex = {pair: row for pair, row in codex_reviews.items() if pair not in invalid_codex}
    pair_path=frozen/f'layer12_rows_{key}.jsonl'
    asins={r['asin'] for r in jsonl(pair_path)}
    headers=['ASIN','Amazon标题','Amazon商品链接','Amazon主图链接','TEMU商品ID','TEMU标题','TEMU商品链接','TEMU主图链接','文字判断依据','向量分','第二层状态']
    if kind=='layer123': headers+= [f'{labels.get(k,k)}/{v}' for k,v in modules.items()]+['第三层总分','最终判断','依据','视觉证据','模型','token']
    def records():
        for r in jsonl(pair_path):
            status=r['pair_status']
            from export_incremental_layer123_excel import second_status, fallback_outcome
            values=[r.get(k) for k in ['asin','amazon_title','amazon_product_url','amazon_image_url','goods_id','temu_title','temu_product_url','temu_image_url','stage1_text_fact','vector_score']]+[second_status(status)]
            if kind=='layer123':
                paid_candidate=status=='CANDIDATE' and admitted(cat,r)
                if status=='CANDIDATE' and not paid_candidate:
                    values[-1]='第二层未入围第三层'
                result=results.get((r['asin'],r['goods_id'])) if status=='CANDIDATE' else None
                if result:
                    scored=score(cat,result['visual_facts'])
                    evidence='；'.join(f"{labels.get(k,k)}：Amazon {m.get('amazon_evidence','')} / TEMU {m.get('temu_evidence','')}" for k,m in scored['modules'].items())
                    from selection_automation.review_evidence import visual_evidence_text, review_explanation
                    evidence = '；'.join(filter(None, [visual_evidence_text(cat,result['visual_facts']), evidence]))
                    values += [scored['modules'][k]['score'] if scored['modules'][k]['score'] is not None else '无法判断' for k in modules]+[scored['total_score'] if scored['total_score'] is not None else '未汇总（证据不全）',scored['state'],review_explanation(cat,result['visual_facts'],scored),evidence,result.get('model'),result.get('usage',{}).get('total_tokens')]
                elif (r['asin'],r['goods_id']) in codex_reviews and status == 'CANDIDATE':
                    review = codex_reviews[(r['asin'],r['goods_id'])]
                    state = {'SAME': '确认同款（Codex看图复核）','SIMILAR':'相似款（Agent看图核验）', 'DIFFERENT': '不同款（Codex看图复核）',
                             'UNKNOWN': '待补充证据（Codex看图复核）'}[review['review_status']]
                    evidence = f"Amazon：{review['amazon_evidence']}；TEMU：{review['temu_evidence']}"
                    reason = '；'.join(filter(None, [review['reason'], review.get('validation_note')]))
                    scored = review.get('module_result')
                    if not binary and scored is None and review.get('visual_facts') is not None:
                        scored = score(cat, review['visual_facts'])
                    if binary:
                        scored = codex_scores.get((r['asin'], r['goods_id']))
                        if scored is not None:
                            state = scored['state'] + '（Codex看图复核）'
                            reason = scored['reason']
                        else:
                            state = '分析失败（历史事实不完整）'
                            reason = '缺少完整原图事实，需要重新看图；旧结果不计为已完成'
                    if review['review_status'] == 'SAME' and scored and scored['state'] == '同款（存在规格变体）':
                        state = '同款（存在规格变体）（Codex看图复核）'
                    module_scores = ([scored['modules'][k]['score'] if scored['modules'][k]['score'] is not None else '无法判断' for k in modules]
                                     if scored else ['未评分（Codex看图复核）']*len(modules))
                    total_score = (scored['total_score'] if scored and scored['total_score'] is not None
                                   else '未汇总（证据不全）' if scored else '未评分（Codex看图复核）')
                    values += module_scores + [total_score,state,
                              reason,evidence,f"Codex看图复核｜{review['model']}｜{review['reviewer']}",None]
                else:
                    _,state,reason=fallback_outcome(status)
                    failed=failures.get((r['asin'],r['goods_id']))
                    if status=='CANDIDATE' and not paid_candidate:
                        state='第二层未入围第三层'
                        reason=f"向量分未达到第三层 {admission_threshold(cat):g} 分门槛，不发送付费请求"
                    elif status=='CANDIDATE' and failed:
                        state='待补充证据（视觉API失败）'
                        reason=f"{failed.get('failure_reason','API失败')}；已尝试{failed.get('attempt',0)}轮"
                    if paid_candidate and failed and failed.get('billing_state')=='NO_CHARGE' and str(failed.get('failure_reason','')).startswith(('IMAGE_DECODE_FAILED','IMAGE_FORMAT_UNSUPPORTED')):
                        state='待补充证据（图片输入无效，未调用API）'
                        reason=f"{failed['failure_reason']}；请检查或重新采集图片，未发送模型请求"
                    if status=='BELOW_VECTOR_THRESHOLD': reason='图片向量低于当前品类候选门槛'
                    if binary:
                        if status == 'CANDIDATE' and not paid_candidate:
                            state = '第二层未入围第三层'
                        elif '复核' in state or '补充' in state:
                            state = '分析失败（请求或输入失败）' if failed else '第三层待运行' if paid_candidate else '第二层未入围第三层'
                    values += [None]*len(modules)+[None,state,reason,'','',None]
            yield values
    ledger=list(jsonl(base/'stage123'/'requests.jsonl', allow_incomplete_tail=True, warnings=log_warnings))
    intents=list(jsonl(base/'stage123'/'request_intents.jsonl', allow_incomplete_tail=True, warnings=log_warnings))
    started={r.get('request_id') for r in intents if r.get('status')=='started' and r.get('request_id')}
    settled={r.get('request_id') for r in intents if r.get('status')=='settled' and r.get('request_id')}
    meta={'Amazon商品数':summary['amazon_products'],'TEMU去重数':summary['temu_products'],'商品对总数':summary['actual_pairs'],
          '状态计数':summary.get('states'),'第三层有效结果':len(results)+len(valid_codex),'规则版本':cat['stage3']['rule_version'],
          '阈值状态':cat['stage3'].get('threshold_status','待校准'),'快照时间':summary.get('snapshot_at'),
          '第三层入围门槛':admission_threshold(cat),'第三层入围候选数':queue_counts['paid'],
          '第二层保留复核未纳入第三层':queue_counts['review'],
          '视觉API请求数':len(started)+sum(r.get('status')!='local_error' and r.get('request_id') not in started for r in ledger),
          '视觉API返回token总计（含拆批和复判）':sum(r.get('usage',{}).get('total_tokens',0) for r in ledger)}
    meta['Codex看图复核完成数'] = len(valid_codex)
    if invalid_codex: meta['Codex历史事实不完整对数'] = len(invalid_codex)
    meta['视觉API有效结果'] = len(results)
    meta['Codex看图复核状态'] = dict(Counter(row['review_status'] for row in valid_codex.values()))
    meta['第三层综合已分析商品对'] = len(results) + len(valid_codex)
    if started-settled:
        meta['未结算视觉请求数'] = len(started-settled)
        meta['未结算请求提示'] = '未结算请求费用未知；请求被拒绝但费用仍待核对，不代表零费用。已阻止自动付费重发，请先核对供应商用量。'
    if log_warnings:
        meta['日志完整性提示'] = ('结果或请求日志有未完成的末行，仅展示完整记录；'
                                  '缺失判断仍待补充证据，请求数和费用记录可能不完整。')
        meta['未完成日志末行'] = log_warnings
    conclusions = None
    low_asins = None
    thumbnails = None
    if kind == 'layer123':
        from selection_automation.commercial import settings
        policy=settings(cat.get('commercial'))
        limit = policy['same_limit']
        conclusions = amazon_conclusions(headers, records(), limit, binary=binary, same_only=True,operator=policy['same_operator'])
        if binary: meta['视觉结论口径'] = 'visible_core_binary_v1'
        low_asins = {item['asin'] for item in conclusions if item['selected']}
        from concurrent.futures import ThreadPoolExecutor
        image_urls = {item['image'] for item in conclusions if item['selected'] and item['image']}
        for item in conclusions:
            if item['selected']: image_urls.update(item['review_images'])
        def thumbnail(url):
            return url, cached_thumbnail({'main_image_url': url, 'local_image_path': image_sources.get(url)},
                                        base/'excel'/'review_image_cache')
        with ThreadPoolExecutor(max_workers=8) as pool: thumbnails = dict(pool.map(thumbnail, sorted(image_urls)))
        meta['新表图片展示'] = 'Amazon最终结论及同款明细使用单元格内锚定图片，全量综合表保留图片链接'
        meta['新表缺失图片数'] = sum(image is None for image in thumbnails.values())
        meta['采集前设置的同款上限'] = limit
        meta['同款数量比较方式'] = policy['same_operator']
        meta['最终结论筛选'] = '仅保留确认同款数不超过采集前上限的Amazon商品'
        meta['最终结论Amazon数'] = len(low_asins)
        meta['同款及待复核专表Amazon数'] = len(low_asins)
        meta['同款及待复核专表行数'] = sum(item['same'] for item in conclusions if item['selected'])
        meta['专表结论筛选'] = '仅同款'
        meta['专表筛选规则'] = '按当前采样确认同款数不超过本任务上限选Amazon，仅保留已确认同款行；沿用本任务已有判断'
        if binary:
            meta['专表筛选规则'] = '按当前采样同款数不超过本任务上限选Amazon，仅保留同款行；主体款式和核心结构对应，颜色标识和无关未展示部位不阻断'
    operator_snapshot = None
    if kind == 'layer123':
        from selection_automation.operations import build_shortlist, read_json
        subscriptions=read_json(ROOT/'runtime/operations/subscriptions.json',{'subscriptions':[]})['subscriptions']
        subscription=next((r for r in subscriptions if r['category_key']==key),None)
        options={k:subscription[k] for k in ('same_limit','min_sales','new_days','require_new')} if subscription else {}
        operator_snapshot=build_shortlist(ROOT,cat,**options)
    return export_stage_book(base/'excel'/f'{name}_{"第一二三层" if kind=="layer123" else "第一二层"}.xlsx',
                             headers, records(), meta, name, third_layer=kind=='layer123',
                             low_asins=low_asins, report=report, conclusions=conclusions, thumbnails=thumbnails,
                             same_only=kind=='layer123', operator_snapshot=operator_snapshot)


def register_delivery(base, kind, delivered):
    """Record a successful standalone delivery; caller holds publish.lock."""
    if not (delivered.get('固定查看副本已更新') or delivered.get('时间戳查看快照已生成')):
        raise RuntimeError(f'{kind}: 工作簿尚未交付查看副本或快照，不登记为已发布')
    state_path = Path(base)/'publish_state.json'
    state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
    deliveries = dict(state.get('deliveries', {}))
    pending = dict(state.get('pending_views', {}))
    deliveries[kind] = delivered
    if delivered.get('固定查看副本已更新'):
        pending.pop(kind, None)
    else:
        pending[kind] = {name: delivered[name] for name in ('数据源工作簿', '固定查看副本')}
    state.update(deliveries=deliveries, pending_views=pending)
    from selection_automation.merge_temu_snapshots import atomic_write
    atomic_write(state_path,json.dumps(state,ensure_ascii=False,indent=2))


def main(argv=None):
    p=argparse.ArgumentParser(); p.add_argument('--category',type=Path,required=True); p.add_argument('--kind',choices=['amazon','temu','layer12','layer123'],required=True)
    a=p.parse_args(argv)
    cat=load_category(a.category)
    base=ROOT/'runs'/cat['category_key']
    from selection_automation.process_lock import ProcessLock
    with ProcessLock(base/'publish.lock'):
        delivered=export(a.category,a.kind)
        register_delivery(base,a.kind,delivered)
    print(json.dumps(delivered,ensure_ascii=False))


if __name__=='__main__':
    main()
