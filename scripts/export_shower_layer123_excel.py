"""Export the fixed 29-column shower-curtain layer-1/2/3 workbook."""
from __future__ import annotations

from collections import Counter
import json
import os
from pathlib import Path
import sys
import uuid
from urllib.parse import urlparse

import xlsxwriter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from selection_automation.category_config import load_category
from selection_automation.rules import score
from selection_workflow.excel_verdict_style import LEGEND, verdict_formats, verdict_kind
from selection_automation.vision import compatible_vision_keys
from selection_workflow.excel_delivery import data_workbook_path, publish_view_copies
from selection_workflow.commercial_excel import add_report


HEADERS = [
    "序号", "Amazon主图", "ASIN", "Amazon标题", "TEMU主图", "TEMU商品ID", "TEMU标题",
    "第一层文字事实", "图片向量相似度", "第二层状态", "Amazon商品链接", "TEMU商品链接",
    "第三层状态", "同一基础产品事实", "视觉证据", "视觉置信度", "主体款式图案版式/32",
    "核心结构安装/18", "材质透明度厚度/15", "功能子类型/12", "表面纹理工艺/10",
    "套装配件/7", "规格变体/6", "第三层总分", "最终判断", "判断依据", "总token", "耗时秒", "模型",
]
MODULES = [
    "design_identity", "core_structure", "material_opacity", "function_subtype",
    "surface_details", "bundle_accessories", "spec_variants",
]
LABELS = {
    "design_identity": "主体款式图案版式", "core_structure": "核心结构安装",
    "material_opacity": "材质透明度厚度", "function_subtype": "功能子类型",
    "surface_details": "表面纹理工艺", "bundle_accessories": "套装配件",
    "spec_variants": "规格变体",
}


from selection_automation.jsonl_records import read_jsonl as jsonl


def second_status(status: str) -> str:
    return {
        "STAGE1_EXCLUDED": "第一层排除：文字硬冲突",
        "BELOW_VECTOR_THRESHOLD": "第二层排除：图片向量低于50分",
        "CANDIDATE": "进入第三层：图片向量达到50分",
        "WAIT_IMAGE_SOURCE": "待补充图片来源",
        "FAILED_BOUNDED": "图片处理失败，已达重试上限",
        "FAILED_RETRYABLE": "图片处理失败，等待重试",
        "PENDING": "第一二层待处理",
    }.get(status, status or "第一二层待处理")


def evidence_text(value) -> str:
    if isinstance(value, list):
        return "；".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").strip()


def https(value) -> bool:
    parsed = urlparse(str(value or ""))
    return parsed.scheme == "https" and bool(parsed.netloc)


def hyperlink_formula(url: str, label: str) -> str:
    escaped_url = str(url).replace('"', '""')
    escaped_label = str(label).replace('"', '""')
    return f'=HYPERLINK("{escaped_url}","{escaped_label}")'


def export(category_path: Path) -> dict:
    category = load_category(category_path)
    key = category["category_key"]
    base = ROOT / "runs" / key
    frozen = base / "stage123" / "export_frozen"
    summary = json.loads((frozen / f"layer12_summary_{key}.json").read_text(encoding="utf-8"))
    rows_path = frozen / f"layer12_rows_{key}.jsonl"
    candidate_path = frozen / f"layer3_candidates_{key}.jsonl"
    expected = {}
    for row in jsonl(candidate_path):
        expected[(str(row["asin"]), str(row["goods_id"]))] = set(compatible_vision_keys(category, row))
    results, failures = {}, {}
    log_warnings = []
    for row in jsonl(base / "stage123" / "results.jsonl", allow_incomplete_tail=True, warnings=log_warnings):
        identity = (str(row.get("asin") or ""), str(row.get("goods_id") or ""))
        if row.get("vision_revision_key") not in expected.get(identity, set()):
            continue
        if row.get("api_status") == "success":
            results[identity] = row
        elif identity not in results:
            failures[identity] = row

    destination = base / "excel" / f"{category['category_name_zh']}_第一二三层.xlsx"
    data_path = data_workbook_path(destination)
    data_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = data_path.with_name(f".{data_path.stem}.{os.getpid()}.{uuid.uuid4().hex}.tmp.xlsx")
    book = xlsxwriter.Workbook(temporary, {"constant_memory": True, "strings_to_urls": False, "strings_to_formulas": False})
    from selection_automation.commercial import read_json, build_report
    navy, blue, pale = "#1F4E78", "#4472C4", "#D9EAF7"
    title = book.add_format({"font_name":"Microsoft YaHei","bold":True,"font_size":15,"font_color":"white","bg_color":navy,"align":"center","valign":"vcenter"})
    note = book.add_format({"font_name":"Microsoft YaHei","font_size":10,"bg_color":pale,"text_wrap":True,"valign":"vcenter"})
    header = book.add_format({"font_name":"Microsoft YaHei","font_size":10,"bold":True,"font_color":"white","bg_color":blue,"align":"center","valign":"vcenter","text_wrap":True})
    text = book.add_format({"font_name":"Microsoft YaHei","font_size":9,"text_wrap":True,"valign":"top"})
    center = book.add_format({"font_name":"Microsoft YaHei","font_size":9,"text_wrap":True,"align":"center","valign":"top"})
    number = book.add_format({"font_name":"Microsoft YaHei","font_size":9,"num_format":"0.00","align":"center","valign":"top"})
    link = book.add_format({"font_name":"Microsoft YaHei","font_size":9,"font_color":"blue","underline":1,"align":"center"})
    verdicts = verdict_formats(book, font_name='Microsoft YaHei', font_size=9)

    widths = [8,12,15,38,12,20,42,42,16,28,15,15,17,18,58,14,22,22,22,22,22,22,22,14,22,48,12,12,18]
    ws = book.add_worksheet("全量综合结果")
    ws.hide_gridlines(2)
    ws.merge_range(0, 0, 0, 28, "Amazon–TEMU 浴帘第一二三层严格同款核验", title)
    ws.merge_range(1, 0, 1, 28, "第一层排除文字硬冲突；第二层向量只筛候选；第三层先核验整体基础款，模块分数仅作辅助，明确身份差异不能被高分抵消。"+LEGEND, note)
    for col, value in enumerate(HEADERS): ws.write(3, col, value, header)
    for col, width in enumerate(widths): ws.set_column(col, col, width)
    ws.freeze_panes(4, 0)

    counts = Counter()
    row_count = 0
    for pair in jsonl(rows_path):
        row_count += 1
        status = str(pair.get("pair_status") or "")
        identity = (str(pair.get("asin") or ""), str(pair.get("goods_id") or ""))
        result = results.get(identity) if status == "CANDIDATE" else None
        third_status, base_fact, final, reason = "不适用", "", "待补充证据", "第一二层尚未完成"
        visual_evidence, confidence, module_values, total, tokens, elapsed, model = "", None, [None] * 7, None, None, None, ""
        if status == "STAGE1_EXCLUDED":
            final, reason = "不同款（文字硬冲突）", "第一层存在双边明确身份级文字硬冲突"
        elif status == "BELOW_VECTOR_THRESHOLD":
            final, reason = "不同款（图片向量门槛）", "图片向量相似度低于50分，不进入第三层"
        elif status in {"WAIT_IMAGE_SOURCE", "FAILED_BOUNDED", "FAILED_RETRYABLE"}:
            third_status, final, reason = "等待补充图片证据", "待补充证据", "图片来源缺失或本地处理失败"
        elif status == "CANDIDATE" and result:
            facts = result.get("visual_facts") or {}
            scored = score(category, facts)
            third_status, final, reason = "第三层已完成", scored["state"], scored["reason"]
            base_fact = "是" if final in {"确认同款", "同款（存在规格变体）"} else "否" if final == "不同款" else "不确定"
            confidence = scored.get("effective_confidence", facts.get("overall_confidence"))
            if isinstance(confidence, (int, float)) and confidence <= 1: confidence = round(confidence * 100, 2)
            evidence_parts = []
            module_values = []
            for name in MODULES:
                item = scored["modules"][name]
                left, right = evidence_text(item.get("amazon_evidence")), evidence_text(item.get("temu_evidence"))
                evidence_parts.append(f"{LABELS[name]}：Amazon {left or '未知'}；TEMU {right or '未知'}")
                module_values.append(f"{item['relation']}｜{item['score']}/{item['maximum']}｜A:{left or '未知'}｜T:{right or '未知'}")
            from selection_automation.identity import evidence_text as identity_evidence
            visual_evidence = "；".join(filter(None, [identity_evidence(facts), *evidence_parts]))
            total = scored.get("total_score")
            corrections = list(scored.get("normalization_notes") or [])
            if corrections:
                third_status = "第三层已完成（本地一致性纠偏）"
                reason += "；" + "；".join(corrections)
            tokens = (result.get("usage") or {}).get("total_tokens")
            elapsed, model = result.get("elapsed_seconds"), result.get("model") or ""
        elif status == "CANDIDATE":
            failed = failures.get(identity)
            if failed and int(failed.get("attempt", 0)) >= 2:
                third_status, final = "第三层请求失败", "待补充证据"
                reason = f"{failed.get('failure_reason') or '视觉API失败'}；已尝试{failed.get('attempt')}轮"
            else:
                third_status, final, reason = "第三层待运行", "第三层待运行", "达到50分候选门槛，尚无成功第三层结果"
        counts[final] += 1
        values = [
            row_count, None, identity[0], pair.get("amazon_title"), None, identity[1], pair.get("temu_title"),
            pair.get("stage1_text_fact"), pair.get("vector_score"), second_status(status), None, None,
            third_status, base_fact, visual_evidence, confidence, *module_values, total, final, reason,
            tokens, elapsed, model,
        ]
        excel_row = row_count + 3
        formats = verdicts[verdict_kind(final)]
        row_text, row_center, row_number, row_link = (formats[name] for name in ('text', 'center', 'number', 'link'))
        for col, value in enumerate(values):
            if col in {1, 4, 10, 11}: continue
            fmt = row_text if col in {3, 6, 7, 14, 16, 17, 18, 19, 20, 21, 22, 25} else row_number if isinstance(value, (int, float)) and not isinstance(value, bool) else row_center
            ws.write(excel_row, col, value, fmt)
        links = (
            (1, pair.get("amazon_image_url"), "查看主图"),
            (4, pair.get("temu_image_url"), "查看主图"),
            (10, f"https://www.amazon.com/dp/{identity[0]}", "查看商品"),
            (11, f"https://www.temu.com/goods.html?goods_id={identity[1]}", "查看商品"),
        )
        for col, url, label in links:
            if https(url):
                ws.write_formula(excel_row, col, hyperlink_formula(str(url), label), row_link, label)
            else:
                ws.write(excel_row, col, "链接缺失", row_center)
        ws.set_row(excel_row, 48)
    ws.autofilter(3, 0, row_count + 3, 28)

    sm = book.add_worksheet("汇总与规则")
    sm.hide_gridlines(2); sm.set_column("A:A", 34); sm.set_column("B:B", 100)
    sm.merge_range("A1:B1", "浴帘第一二三层汇总与规则", title)
    states = summary.get("states") or {}
    ledger = list(jsonl(base / "stage123" / "requests.jsonl", allow_incomplete_tail=True, warnings=log_warnings))
    intents = list(jsonl(base / 'stage123' / 'request_intents.jsonl', allow_incomplete_tail=True, warnings=log_warnings))
    started = {r.get('request_id') for r in intents if r.get('status') == 'started' and r.get('request_id')}
    settled = {r.get('request_id') for r in intents if r.get('status') == 'settled' and r.get('request_id')}
    entries = [
        ("Amazon商品数", summary.get("amazon_products")), ("TEMU去重商品数", summary.get("temu_products")),
        ("商品对总数", row_count), ("第一层文字硬冲突排除", states.get("STAGE1_EXCLUDED", 0)),
        ("第二层低于50分", states.get("BELOW_VECTOR_THRESHOLD", 0)), ("第三层候选", len(expected)),
        ("第三层有效成功结果", len(results)), ("第三层尚未成功", len(expected) - len(results)),
        ("结果状态计数", json.dumps(dict(counts), ensure_ascii=False)),
        ("第三层API请求批次", len(started) + sum(item.get("status") != "local_error" and item.get("request_id") not in started for item in ledger)),
        ("API返回token总计", sum(int((item.get("usage") or {}).get("total_tokens", 0) or 0) for item in ledger)),
        ("第一层规则", "只在双边文字明确证明身份级冲突时排除；文字不足继续进入图片层。"),
        ("第二层规则", "ResNet50 ImageNet V2主图向量余弦相似度×100，达到50分进入第三层；向量不计入最终总分。"),
        ("第三层协议", "整体基础款判断、具体双边共同细节、身份差异和未知项；七模块为辅助证据。"),
        ("第三层评分", "辅助分数保留供查看，不决定同款。整体一致须有具体共同细节支持，关键模块不能未知、仅相似或以身份变体解释。"),
        ("判定边界", "不自动提高零置信度、不调换两侧证据、不用向量高分改写第三层判断；旧协议视觉结果须重新分析。"),
        ("输出限制", "当前每侧使用一张主图；缺图、失败和未知证据不记为0分。全部明细集中在一个工作表，链接使用可点击公式以避开Excel超链接对象数量限制。"),
    ]
    if log_warnings:
        entries += [('日志完整性提示', '仅展示完整记录；日志末行未完成，缺失判断待补充证据，用量可能不完整。'),
                    ('未完成日志末行', json.dumps(log_warnings, ensure_ascii=False))]
    if started-settled:
        entries += [('未结算视觉请求数', len(started-settled)),
                    ('未结算请求提示', '费用未知；请求被拒绝但费用仍待核对，不代表零费用，已阻止自动付费重发，请先核对供应商用量。')]

    for index, (name, value) in enumerate(entries, 1):
        sm.write(index, 0, name, header); sm.write(index, 1, value, text); sm.set_row(index, 28)
    sm.freeze_panes(1, 0)
    add_report(book, read_json(base / 'commercial_report.json') or build_report(ROOT, category))
    book.close()
    try:
        check = load_workbook(temporary, read_only=True, data_only=False)
        try:
            expected_names = ["全量综合结果", "汇总与规则"]
            if check.sheetnames[:2] != expected_names: raise RuntimeError("工作表顺序错误")
            detail = check["全量综合结果"]
            checked_rows = detail.max_row - 4
            if [detail.cell(4, col).value for col in range(1, 30)] != HEADERS: raise RuntimeError("29列表头错误")
            if checked_rows != row_count: raise RuntimeError("明细行数错误")
        finally:
            check.close()
        os.replace(temporary, data_path)
    finally:
        temporary.unlink(missing_ok=True)
    delivered = publish_view_copies(data_path, destination, keep_snapshots=3)
    return {**delivered, "商品对总数": row_count, "第三层有效结果": len(results), "列数": 29, "明细分表数": 1}


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(); parser.add_argument("--category", type=Path, required=True)
    print(json.dumps(export(parser.parse_args().category), ensure_ascii=False))
